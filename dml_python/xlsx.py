"""
DML → XLSX — نفس نسق الطباعة في ملف إكسل.

``html_to_workbook(html, title)`` parses the printable preview HTML
(``render_preview_html`` output: header boxes, sections, tableviews,
footer) into an RTL workbook preserving order:
title → header text → sections (text + tables) → footer.

``table_to_workbook(columns, rows, title)`` builds the same-styled sheet
directly when no DML document is selected.
"""
from __future__ import annotations
import html as _html_mod
from html.parser import HTMLParser
from typing import Any, Dict, List, Optional

TITLE_FILL = "1F4E78"
TITLE_FONT_COLOR = "FFFFFF"
HEAD_FILL = "D9E1F2"
GROUP_FILL = "BDD7EE"
BORDER_COLOR = "94A3B8"
EVEN_FILL = "F2F2F2"


class _Cell:
    __slots__ = ("text", "bold", "header", "colspan", "rowspan")

    def __init__(self, text="", bold=False, header=False, colspan=1, rowspan=1):
        self.text = text
        self.bold = bold
        self.header = header
        self.colspan = colspan
        self.rowspan = rowspan


class _DocParser(HTMLParser):
    """Flatten print HTML into blocks: ('text', txt, bold) | ('table', rows)."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocks: List[Any] = []
        self._table_depth = 0
        self._rows: List[List[_Cell]] | None = None
        self._cur_row: List[_Cell] | None = None
        self._cur_cell: _Cell | None = None
        self._cell_text: List[str] = []
        self._in_th = False
        self._in_thead = False
        self._bold_depth = 0
        self._text_buf: List[str] = []
        self._text_bold = False
        self._skip = 0  # inside style/script

    # ── helpers ──
    def _flush_text(self):
        txt = "".join(self._text_buf).strip()
        self._text_buf = []
        if txt:
            self.blocks.append(("text", txt, self._text_bold))
        self._text_bold = False

    # ── parser hooks ──
    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        if tag in ("style", "script", "head"):
            self._skip += 1
            return
        if self._skip:
            return
        attrs = {k.lower(): (v or "") for k, v in attrs}
        if tag == "table":
            self._flush_text()
            self._table_depth += 1
            if self._table_depth == 1:
                self._rows = []
            return
        if self._table_depth:
            if tag == "thead":
                self._in_thead = True
            elif tag == "tr":
                self._cur_row = []
            elif tag in ("th", "td"):
                self._in_th = (tag == "th")
                self._cur_cell = _Cell(
                    header=(tag == "th") or self._in_thead,
                    colspan=max(1, int(attrs.get("colspan", "1") or 1)),
                    rowspan=max(1, int(attrs.get("rowspan", "1") or 1)),
                )
                self._cell_text = []
            elif tag == "br" and self._cur_cell is not None:
                self._cell_text.append("\n")
            return
        if tag in ("h1", "h2", "h3", "h4", "b", "strong"):
            self._bold_depth += 1
            self._text_bold = True
        elif tag == "br":
            self._text_buf.append("\n")
        elif tag == "p" and "".join(self._text_buf).strip():
            self._flush_text()

    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag in ("style", "script", "head"):
            self._skip = max(0, self._skip - 1)
            return
        if self._skip:
            return
        if tag == "table":
            self._table_depth = max(0, self._table_depth - 1)
            if self._table_depth == 0 and self._rows is not None:
                self.blocks.append(("table", self._rows))
                self._rows = None
            return
        if self._table_depth:
            if tag == "thead":
                self._in_thead = False
            elif tag in ("th", "td"):
                if self._cur_cell is not None and self._cur_row is not None:
                    self._cur_cell.text = "".join(self._cell_text).strip()
                    self._cur_row.append(self._cur_cell)
                self._cur_cell = None
                self._cell_text = []
                self._in_th = False
            elif tag == "tr":
                if self._cur_row is not None and self._rows is not None:
                    self._rows.append(self._cur_row)
                self._cur_row = None
            return
        if tag in ("h1", "h2", "h3", "h4", "b", "strong"):
            self._bold_depth = max(0, self._bold_depth - 1)
            if not self._bold_depth:
                self._flush_text()
        elif tag in ("p", "div"):
            self._flush_text()

    def handle_data(self, data):
        if self._skip:
            return
        if self._table_depth and self._cur_cell is not None:
            self._cell_text.append(data)
        elif not self._table_depth:
            self._text_buf.append(data)


def _style_sheet(ws, ncols: int):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    ws.sheet_view.rightToLeft = True
    ws.sheet_format.defaultRowHeight = 18
    thin = Side(style="thin", color=BORDER_COLOR)
    return {
        "title_font": Font(name="Arial", size=14, bold=True, color=TITLE_FONT_COLOR),
        "title_fill": PatternFill("solid", fgColor=TITLE_FILL),
        "title_align": Alignment(horizontal="center", vertical="center"),
        "head_font": Font(name="Arial", size=11, bold=True),
        "head_fill": PatternFill("solid", fgColor=HEAD_FILL),
        "head_align": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "grp_fill": PatternFill("solid", fgColor=GROUP_FILL),
        "cell_font": Font(name="Arial", size=11),
        "cell_align": Alignment(horizontal="right", vertical="center", wrap_text=True),
        "center": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
        "even_fill": PatternFill("solid", fgColor=EVEN_FILL),
    }


def _auto_widths(ws, widths: List[int], ncols: int):
    for i in range(ncols):
        w = widths[i] if i < len(widths) else 10
        try:
            ws.column_dimensions[openpyxl_col(i)].width = min(max(w + 2, 12), 45)
        except Exception:
            pass


def openpyxl_col(i: int) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(i + 1)


def _emit_text(ws, st, row: int, text: str, bold: bool, ncols: int, widths: List[int]):
    from openpyxl.styles import Alignment, Font
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(ncols, 1))
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name="Arial", size=12 if bold else 11, bold=bold)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths[0] = max(widths[0] if widths else 0, min(len(text), 60))
    return row + 1


def _emit_table(ws, st, row: int, rows: List[List[_Cell]], widths: List[int]) -> int:
    from copy import copy
    ncols = 1
    for r in rows:
        ncols = max(ncols, sum(c.colspan for c in r))
    while len(widths) < ncols:
        widths.append(10)
    # pending rowspans: col -> rows left
    pending: Dict[int, int] = {}
    for r in rows:
        col = 0
        r0 = row
        for cell in r:
            while pending.get(col, 0) > 0:
                col += 1
            for _ in range(cell.colspan):
                pending[col] = max(pending.get(col, 0), cell.rowspan)
                col += 1
            c0 = col - cell.colspan
            c = ws.cell(row=r0, column=c0 + 1, value=cell.text)
            c.font = st["head_font"] if (cell.header or cell.bold) else st["cell_font"]
            c.alignment = st["head_align"] if cell.header else st["center"]
            c.border = st["border"]
            if cell.header:
                c.fill = st["grp_fill"] if cell.colspan > 1 else st["head_fill"]
            if "\n" in (cell.text or ""):
                from openpyxl.styles import Alignment as _A
                a = copy(c.alignment)
                a.wrap_text = True
                c.alignment = a
            if cell.colspan > 1 or cell.rowspan > 1:
                ws.merge_cells(start_row=r0, start_column=c0 + 1,
                               end_row=r0 + cell.rowspan - 1, end_column=c0 + cell.colspan)
                # borders on merged range edges
                for rr in range(r0, r0 + cell.rowspan):
                    for cc in range(c0 + 1, c0 + cell.colspan + 1):
                        if rr == r0 and cc == c0 + 1:
                            continue
                        ws.cell(row=rr, column=cc).border = st["border"]
            for k, ln in enumerate(cell.text.split("\n")):
                widths[c0] = max(widths[c0], min(len(ln), 50))
        # tick pending
        for k in list(pending.keys()):
            pending[k] -= 1
            if pending[k] <= 0:
                del pending[k]
        row += 1
    # zebra for body rows is print-driven; skip (print CSS handles it)
    return row


def html_to_workbook(html: str, title: str = ""):
    """Parse printable DML HTML → openpyxl Workbook (RTL, styled)."""
    from openpyxl import Workbook
    p = _DocParser()
    p.feed(html or "")
    p.close()
    wb = Workbook()
    ws = wb.active
    ws.title = (title or "تقرير")[:31]
    st = _style_sheet(ws, 1)
    widths: List[int] = [10]
    # discover width early: max table columns
    ncols = 1
    for b in p.blocks:
        if b[0] == "table":
            ncols = max(ncols, max((sum(c.colspan for c in r) for r in b[1]), default=1))
    row = 1
    if title:
        from openpyxl.styles import Alignment, Font
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
        c = ws.cell(row=row, column=1, value=title)
        c.font = st["title_font"]
        c.fill = st["title_fill"]
        c.alignment = st["title_align"]
        ws.row_dimensions[row].height = 28
        row += 2
    for b in p.blocks:
        if b[0] == "text":
            row = _emit_text(ws, st, row, b[1], b[2], ncols, widths)
        else:
            row = _emit_table(ws, st, row, b[1], widths)
            row += 1  # blank separator between blocks
    _auto_widths(ws, widths, ncols)
    try:
        ws.sheet_properties.pageSetUpPr = openpyxl_pagesetup()
    except Exception:
        pass
    return wb


def openpyxl_pagesetup():
    from openpyxl.worksheet.properties import PageSetupProperties
    return PageSetupProperties(fitToPage=True)


def table_to_workbook(columns: List[Dict[str, Any]], rows: List[Dict[str, Any]], title: str = ""):
    """Plain report table → styled workbook (same grid look, no DML doc)."""
    # transpose: header row
    head = [_Cell(str(c.get("alias") or c.get("name") or ""), header=True) for c in (columns or [])]
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = (title or "تقرير")[:31]
    st = _style_sheet(ws, 1)
    widths: List[int] = [10]
    ncols = max(len(head), 1)
    row = 1
    if title:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
        c = ws.cell(row=row, column=1, value=title)
        c.font = st["title_font"]
        c.fill = st["title_fill"]
        c.alignment = st["title_align"]
        ws.row_dimensions[row].height = 28
        row += 2
    if head:
        row = _emit_table(ws, st, row, [head], widths)
        # بث مباشر صفاً بصف (بلا قائمة body كاملة) + تخطي الزبرا فوق 20 ألف صف
        total = len(rows or [])
        zebra = total <= 20000
        for i, r in enumerate(rows or []):
            line = []
            for c in (columns or []):
                key = c.get("alias") or c.get("name") or ""
                v = r.get(key, r.get(c.get("name") or "", ""))
                if isinstance(v, dict):
                    v = ""
                line.append(_Cell("" if v is None else str(v)))
            # zebra
            r0 = row
            row = _emit_table(ws, st, row, [line], widths)
            if zebra and i % 2 == 1:
                for cc in range(1, ncols + 1):
                    try:
                        ws.cell(row=r0, column=cc).fill = st["even_fill"]
                    except Exception:
                        pass
    _auto_widths(ws, widths, ncols)
    return wb


def workbook_to_bytes(wb) -> bytes:
    from io import BytesIO
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
