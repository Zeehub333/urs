"""Workspace login / first-superuser tests (no live DB — fake PG cursor)."""
import json
import os
from types import SimpleNamespace
from unittest import mock

from django.test import RequestFactory, SimpleTestCase, TestCase

from urs import views as _v


# sys_users shape as discovered on main_hq_2026 (name, type, nullable, default)
SYS_USERS_COLS = [
    ("id", "integer", "NO", "nextval('sys_users_id_seq'::regclass)"),
    ("username", "character varying", "NO", None),
    ("full_name", "character varying", "NO", None),
    ("email", "character varying", "NO", None),
    ("password", "character varying", "NO", None),
    ("role", "character varying", "NO", None),
    ("phone", "character varying", "YES", None),
    ("company_code", "character varying", "YES", None),
    ("branch_code", "character varying", "YES", None),
    ("is_active", "boolean", "NO", "true"),
    ("created_at", "timestamp without time zone", "YES", "CURRENT_TIMESTAMP"),
]


class _FakeCur:
    """Answers the information_schema / count / insert queries views.py issues."""

    def __init__(self, state):
        self.state = state
        self.executed = []
        self._rows = []

    def execute(self, sql, params=None):
        self.executed.append((sql, params))
        if "information_schema.tables" in sql:
            self._rows = [(self.state["schema"],)]
        elif "is_nullable" in sql:
            self._rows = list(self.state["fullcols"])
        elif "lower(column_name)" in sql:
            self._rows = [(n.lower(), t) for (n, t, _nu, _df) in self.state["fullcols"]]
        elif "SELECT column_name FROM" in sql:
            self._rows = [(n,) for (n, _t, _nu, _df) in self.state["fullcols"]]
        elif "count(*)" in sql:
            self._rows = [(self.state["count"],)]
        elif sql.strip().upper().startswith("SELECT 1"):
            self._rows = [(1,)] if self.state.get("exists") else []
        elif sql.strip().upper().startswith("INSERT"):
            self.state["insert"] = (sql, params)
            self._rows = []
        elif sql.strip().upper().startswith("UPDATE"):
            self.state["update"] = (sql, params)
            self._rows = []
        elif "LIMIT 1" in sql and "WHERE" in sql:
            _ur = self.state.get("userrow")
            self._rows = [_ur] if _ur is not None else []
        else:
            self._rows = []

    def fetchall(self):
        return list(self._rows)

    def fetchone(self):
        return self._rows[0] if self._rows else None

    def close(self):
        pass


class _FakeConn:
    def __init__(self, state):
        self.state = state
        self.cur = _FakeCur(state)

    def cursor(self):
        return self.cur

    def commit(self):
        pass

    def close(self):
        pass


def _fresh_state(**kw):
    st = {"schema": "main_hq_2026", "fullcols": list(SYS_USERS_COLS),
          "count": 0, "exists": False, "insert": None,
          "userrow": None, "update": None}
    st.update(kw)
    return st


def _patched(state):
    """Patch all DB-touching seams of the workspace endpoints."""
    return mock.patch.multiple(
        _v,
        _ws_master_ok=mock.Mock(return_value=True),
        _ws_login_connection=mock.Mock(return_value=({"fake": True}, "sys_users", "")),
        _ws_pg_connect=mock.Mock(side_effect=lambda _obj: _FakeConn(state)),
    )


class RequiredFieldsUnitTests(SimpleTestCase):
    def test_detects_email_and_role_only(self):
        st = _fresh_state()
        cur = _FakeCur(st)
        req = _v._ws_required_user_fields(
            cur, "main_hq_2026", "sys_users",
            skip={"username", "password", "full_name", "is_active"})
        self.assertEqual(sorted(c["name"] for c in req), ["email", "role"])

    def test_nullable_and_defaulted_are_not_required(self):
        st = _fresh_state(fullcols=[
            ("a", "character varying", "YES", None),
            ("b", "integer", "NO", "0"),
            ("c", "text", "NO", None),
        ])
        req = _v._ws_required_user_fields(_FakeCur(st), "s", "t", skip=[])
        self.assertEqual([c["name"] for c in req], ["c"])

    def test_real_columns_map(self):
        st = _fresh_state(fullcols=[("Email", "character varying", "NO", None)])
        self.assertEqual(_v._ws_real_columns(_FakeCur(st), "s", "t"), {"email": "Email"})

    def test_never_raises_on_broken_cursor(self):
        class _Bad:
            def execute(self, *a, **k):
                raise RuntimeError("down")
        self.assertEqual(_v._ws_required_user_fields(_Bad(), "s", "t"), [])
        self.assertEqual(_v._ws_real_columns(_Bad(), "s", "t"), {})

    def test_fields_schema_types(self):
        st = _fresh_state(fullcols=[
            ("email", "character varying", "NO", None),
            ("is_active", "boolean", "NO", "true"),
            ("phone", "character varying", "YES", None),
            ("birth_date", "date", "YES", None),
            ("created_at", "timestamp without time zone", "YES", None),
            ("login_time", "time without time zone", "YES", None),
            ("notes", "text", "YES", None),
            ("age", "integer", "YES", None),
        ])
        fields = _v._ws_user_fields_schema(_FakeCur(st), "s", "t", skip=["created_at"])
        by_name = {f["name"]: f for f in fields}
        self.assertNotIn("created_at", by_name)
        self.assertEqual(by_name["email"]["input_type"], "email")
        self.assertEqual(by_name["is_active"]["input_type"], "checkbox")
        self.assertEqual(by_name["phone"]["input_type"], "tel")
        self.assertEqual(by_name["birth_date"]["input_type"], "date")
        self.assertEqual(by_name["login_time"]["input_type"], "time")
        self.assertEqual(by_name["notes"]["input_type"], "textarea")
        self.assertEqual(by_name["age"]["input_type"], "number")



class SuperuserEndpointTests(SimpleTestCase):
    def setUp(self):
        self.rf = RequestFactory()
        self.dirs = mock.patch("urs.workspace.workspace_dirs",
                               return_value=[SimpleNamespace(name="workspace_1")])
        self.dirs.start()
        self.addCleanup(self.dirs.stop)

    def _post(self, state, payload):
        with _patched(state):
            req = self.rf.post("/api/workspaces/workspace_1/superuser/",
                               data=json.dumps(payload),
                               content_type="application/json")
            return _v.api_workspace_superuser(req, "workspace_1"), state

    def test_users_count_reports_required_when_empty(self):
        st = _fresh_state()
        with _patched(st):
            req = self.rf.get("/api/workspaces/workspace_1/users-count/")
            res = _v.api_workspace_users_count(req, "workspace_1")
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200)
        self.assertEqual(body["count"], 0)
        self.assertEqual(sorted(c["name"] for c in body["required"]), ["email", "role"])
        self.assertTrue(len(body.get("fields", [])) > 0)
        field_names = [f["name"] for f in body["fields"]]
        self.assertIn("email", field_names)
        self.assertIn("phone", field_names)

    def test_superuser_type_casting(self):
        st = _fresh_state()
        res, _ = self._post(st, {
            "username": "admin", "password": "secret123", "full_name": "المدير",
            "extra": {"email": "admin", "role": "admin", "is_active": "true", "phone": ""},
        })
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        sql, params = st["insert"]
        # Check that is_active was cast to True (bool)
        self.assertIn(True, params)
        # Check that empty nullable phone was cast to None (SQL NULL)
        self.assertIn(None, params)


    def test_users_count_no_required_when_not_empty(self):
        st = _fresh_state(count=3)
        with _patched(st):
            req = self.rf.get("/api/workspaces/workspace_1/users-count/")
            res = _v.api_workspace_users_count(req, "workspace_1")
        body = json.loads(res.content.decode())
        self.assertEqual(body["count"], 3)
        self.assertEqual(body["required"], [])

    def test_accepts_extra_and_inserts_only_known_columns(self):
        st = _fresh_state()
        res, _ = self._post(st, {
            "username": "admin", "password": "secret123", "full_name": "المدير",
            "extra": {"email": "a@x.co", "role": "admin",
                      "username": "hacker", "nope'); DROP TABLE x;--": "z"},
        })
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        self.assertTrue(body.get("ok"))
        sql, params = st["insert"]
        self.assertIn('"email"', sql)
        self.assertIn('"role"', sql)
        self.assertNotIn("DROP", sql)
        self.assertIn("a@x.co", params)
        self.assertIn("admin", params)
        # password must be hashed, never plaintext
        self.assertNotIn("secret123", params)

    def test_refuses_when_users_exist(self):
        st = _fresh_state(count=2)
        res, _ = self._post(st, {"username": "admin", "password": "secret123",
                                 "extra": {"email": "a@x.co", "role": "admin"}})
        self.assertEqual(res.status_code, 400)
        self.assertIsNone(st["insert"])

    def test_rejects_missing_required_fields(self):
        st = _fresh_state()
        res, _ = self._post(st, {"username": "admin", "password": "secret123"})
        body = json.loads(res.content.decode())
        _req = body.get("required", [])
        _missing = body.get("error", "")
        # When email is auto-composed (username@domain), it should NOT be in missing
        # Only role should be missing since it's explicitly required but not provided
        self.assertEqual(res.status_code, 400)
        # role must be missing (not provided, no auto-compose)
        self.assertIn("role", _missing)
        # email should NOT be missing (auto-composed from username + domain)
        self.assertNotIn("email", _missing)
        self.assertEqual(sorted(c["name"] for c in _req), ["email", "role"])
        self.assertIsNone(st["insert"])


class PasswordHashNamespaceTests(SimpleTestCase):
    """Django-format hashes (pbkdf2_sha256$) must verify via Django's parser;
    std-format hashes (pbkdf2_sha256_std$) via the stdlib branch. A past
    regression aliased the two namespaces and broke all Django-hash logins."""

    def test_django_format_verifies(self):
        from django.contrib.auth.hashers import make_password
        from fmlk_engine.engine import verify_secret, is_hashed_secret
        h = make_password("h123456*")
        self.assertTrue(h.startswith("pbkdf2_sha256$"))
        self.assertTrue(is_hashed_secret(h))
        self.assertTrue(verify_secret("h123456*", h))
        self.assertFalse(verify_secret("wrong", h))
        self.assertTrue(_v._verify_login_password("h123456*", h, username="admin"))
        self.assertFalse(_v._verify_login_password("wrong", h, username="admin"))

    def test_std_format_verifies(self):
        import hashlib as _hl
        import os as _os
        from fmlk_engine.engine import verify_secret, _STD_HASH_PREFIX, _STD_HASH_ITERS
        salt = _os.urandom(16).hex()
        dk = _hl.pbkdf2_hmac("sha256", "h123456*".encode(), bytes.fromhex(salt), _STD_HASH_ITERS)
        h = "%s%d$%s$%s" % (_STD_HASH_PREFIX, _STD_HASH_ITERS, salt, dk.hex())
        self.assertTrue(verify_secret("h123456*", h))
        self.assertFalse(verify_secret("wrong", h))

    def test_namespaces_stay_distinct(self):
        from fmlk_engine.engine import _STD_HASH_PREFIX
        self.assertNotEqual(_STD_HASH_PREFIX, "pbkdf2_sha256$")

    def test_reference_fastapi_dialect_verifies(self):
        # Reference service: salt = token_hex used as UTF-8 bytes (NOT hex-decoded).
        import hashlib as _hl
        from fmlk_engine.engine import verify_secret
        salt, iters = "c9f978d5bdbedd43ca78757148a29ff8", 600000
        dk = _hl.pbkdf2_hmac("sha256", "h123456*".encode(), salt.encode(), iters).hex()
        h = "pbkdf2_sha256$%d$%s$%s" % (iters, salt, dk)
        self.assertTrue(verify_secret("h123456*", h))
        self.assertFalse(verify_secret("wrong", h))
        self.assertTrue(_v._verify_login_password("h123456*", h, username="admin"))

    def test_db_hash_shape_with_known_password(self):
        # Same shape as the production row (100k iters, hex salt/digest):
        # a hash MADE from a known password must verify end-to-end.
        import hashlib as _hl
        from fmlk_engine.engine import verify_secret
        salt, iters = "ab12cd34ef56ab12cd34ef56ab12cd34", 100000
        dk = _hl.pbkdf2_hmac("sha256", "known-pw-1".encode(), salt.encode(), iters).hex()
        h = "pbkdf2_sha256$%d$%s$%s" % (iters, salt, dk)
        self.assertTrue(verify_secret("known-pw-1", h))
        self.assertFalse(verify_secret("h123456*", h))

    def test_blank_padded_hash_verifies(self):
        # char(n) columns blank-pad — padded Django hash must still verify
        from django.contrib.auth.hashers import make_password
        h = make_password("h123456*") + "   "
        self.assertTrue(_v._verify_login_password("h123456*", h, username="admin"))
        self.assertFalse(_v._verify_login_password("wrong", h, username="admin"))

    def test_hash_itself_is_not_a_password(self):
        # pass-the-hash: entering the stored hash must NOT authenticate
        from django.contrib.auth.hashers import make_password
        import hashlib as _hl
        dh = make_password("h123456*")
        self.assertFalse(_v._verify_login_password(dh, dh, username="admin"))
        self.assertTrue(_v._verify_login_password("h123456*", dh, username="admin"))
        hx = _hl.sha256("h123456*".encode()).hexdigest()
        self.assertFalse(_v._verify_login_password(hx, hx, username="admin"))
        self.assertTrue(_v._verify_login_password("h123456*", hx, username="admin"))
        # legacy plaintext still works
        self.assertTrue(_v._verify_login_password("plain123", "plain123", username="admin"))


class PasswordHashColumnTests(SimpleTestCase):
    """password_hash (and siblings) must be treated as secrets everywhere:
    login columns, engine hashing/masking, designer type aliases."""

    def test_login_columns_include_password_hash(self):
        self.assertIn("password_hash", _v._WS_PASS_COLS)

    def test_engine_secret_names_include_password_hash(self):
        from fmlk_engine.engine import FMLKFormEngine
        self.assertIn("password_hash", FMLKFormEngine.SECRET_NAMES)

    def test_designer_alias_maps_password_hash(self):
        from fmlk_engine.field_types import normalize_input_type, get_type
        self.assertEqual(normalize_input_type("password_hash"), "password")
        self.assertEqual(normalize_input_type("email"), "email")
        self.assertEqual(normalize_input_type("email_field"), "email")
        self.assertEqual(normalize_input_type("mail"), "email")
        em_type = get_type("email")
        self.assertIsNotNone(em_type)
        self.assertIn("email_domain", em_type.get("attrs", []))


class UsersColumnMappingTests(SimpleTestCase):
    def setUp(self):
        from pathlib import Path as _P
        self.rf = RequestFactory()
        self.dirs = mock.patch("urs.workspace.workspace_dirs",
                               return_value=[_P("workspace_1")])
        self.dirs.start()
        self.addCleanup(self.dirs.stop)

    def test_colmap_reads_and_validates(self):
        with mock.patch("urs.workspace.read_conf",
                        return_value={"USERS_USER_COLUMN": "Login_Name",
                                      "USERS_PASSWORD_COLUMN": "pwd_hash"}):
            self.assertEqual(_v._ws_users_colmap("workspace_1"), ("login_name", "pwd_hash"))
        with mock.patch("urs.workspace.read_conf", return_value={}):
            self.assertEqual(_v._ws_users_colmap("workspace_1"), ("", ""))
        with mock.patch("urs.workspace.read_conf",
                        return_value={"USERS_USER_COLUMN": "bad name!",
                                      "USERS_PASSWORD_COLUMN": "password"}):
            self.assertEqual(_v._ws_users_colmap("workspace_1"), ("", "password"))

    def test_handled_cols_honor_pins(self):
        cols = {"username": "varchar", "login": "varchar",
                "password": "varchar", "pwd": "varchar"}
        u, p, _f, _s, _a = _v._ws_user_handled_cols(cols, "login", "pwd")
        self.assertEqual((u, p), ("login", "pwd"))
        u, p, _f, _s, _a = _v._ws_user_handled_cols(cols)
        self.assertEqual((u, p), ("username", "password"))
        u, p, _f, _s, _a = _v._ws_user_handled_cols(cols, "nope", "")
        self.assertEqual((u, p), ("username", "password"))

    def test_write_conf_roundtrip(self):
        import tempfile as _tf
        from pathlib import Path as _P
        with _tf.TemporaryDirectory() as td:
            ws = _P(td)
            (ws / "workspace.conf").write_text(
                "# test\nPRIMARY_CONNECTION=urs_local\nUSERS_TABLE=s.t\n", encoding="utf-8")
            _v._write_ws_conf(ws, "urs_local", "s.t", "active", "login", "pwd_hash")
            from urs import workspace as _wsm
            conf = _wsm.read_conf(ws / "workspace.conf")
            self.assertEqual(conf.get("USERS_USER_COLUMN"), "login")
            self.assertEqual(conf.get("USERS_PASSWORD_COLUMN"), "pwd_hash")
            # clearing drops the keys
            _v._write_ws_conf(ws, "urs_local", "s.t", "active", "", "")
            conf = _wsm.read_conf(ws / "workspace.conf")
            self.assertNotIn("USERS_USER_COLUMN", conf)
            self.assertNotIn("USERS_PASSWORD_COLUMN", conf)

    def test_users_columns_endpoint(self):
        st = _fresh_state()
        with _patched(st), \
                mock.patch.object(_v, "_ws_users_colmap", return_value=("username", "")):
            req = self.rf.get("/api/workspaces/workspace_1/users-columns/")
            res = _v.api_workspace_users_columns(req, "workspace_1")
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200)
        names = {c["name"] for c in body["columns"]}
        self.assertIn("username", names)
        self.assertIn("email", names)
        self.assertEqual(body["user_column"], "username")

    def test_verify_login_password_strict_algo(self):
        import hashlib as _hl
        pw = "S3cret!"
        md5v = _hl.md5(pw.encode()).hexdigest()
        sha1v = _hl.sha1(pw.encode()).hexdigest()
        sha256v = _hl.sha256(pw.encode()).hexdigest()
        mysqlv = "*" + _hl.sha1(_hl.sha1(pw.encode()).digest()).hexdigest().upper()
        pgmv = "md5" + _hl.md5((pw + "omar").encode()).hexdigest()
        self.assertTrue(_v._verify_login_password(pw, pw, algo="plain"))
        self.assertTrue(_v._verify_login_password(pw, md5v, algo="md5"))
        self.assertTrue(_v._verify_login_password(pw, sha1v, algo="sha1"))
        self.assertTrue(_v._verify_login_password(pw, sha256v, algo="sha256"))
        self.assertTrue(_v._verify_login_password(pw, mysqlv, algo="mysql_sha1"))
        self.assertTrue(_v._verify_login_password(pw, pgmv, username="omar", algo="pg_md5"))
        # strict: wrong scheme or wrong secret fails (no cascade)
        self.assertFalse(_v._verify_login_password(pw, sha256v, algo="md5"))
        self.assertFalse(_v._verify_login_password(pw, md5v, algo="sha256"))
        self.assertFalse(_v._verify_login_password("nope", md5v, algo="md5"))
        self.assertFalse(_v._verify_login_password(pw, md5v, algo="bogus"))
        # auto cascade untouched
        self.assertTrue(_v._verify_login_password(pw, md5v))
        self.assertTrue(_v._verify_login_password(pw, md5v, algo="auto"))

    def test_verify_login_password_classical_algos(self):
        from fmlk_engine.crypto import shift_unicode, atbash_cipher, rot13_cipher
        import base64 as _b64
        pw = "S3cret!"
        # Caesar raw with explicit shift + default shift + prefixed form
        self.assertTrue(_v._verify_login_password(pw, shift_unicode(pw, 5), algo="caesar", shift=5))
        self.assertFalse(_v._verify_login_password(pw, shift_unicode(pw, 5), algo="caesar", shift=3))
        self.assertTrue(_v._verify_login_password(pw, shift_unicode(pw, 3), algo="caesar"))
        _pref = "caesar$7$" + _b64.b64encode(shift_unicode(pw, 7).encode()).decode()
        self.assertTrue(_v._verify_login_password(pw, _pref, algo="caesar", shift=3))
        self.assertFalse(_v._verify_login_password("nope", _pref, algo="caesar"))
        # Atbash / ROT-13 raw + prefixed
        self.assertTrue(_v._verify_login_password(pw, atbash_cipher(pw), algo="atbash"))
        self.assertTrue(_v._verify_login_password(pw, "atbash$" + atbash_cipher(pw), algo="atbash"))
        self.assertTrue(_v._verify_login_password(pw, rot13_cipher(pw), algo="rot13"))
        self.assertTrue(_v._verify_login_password(pw, "rot13$" + rot13_cipher(pw), algo="rot13"))
        self.assertFalse(_v._verify_login_password(pw, atbash_cipher(pw), algo="rot13"))
        self.assertFalse(_v._verify_login_password(pw, rot13_cipher(pw), algo="md5"))

    def test_auth_cfg_reads_and_validates(self):
        with mock.patch("urs.workspace.read_conf",
                        return_value={"AUTH_CONN": "ORCL", "AUTH_USER_COLUMN": "Login_Name",
                                      "AUTH_PASSWORD_COLUMN": "pwd_hash", "AUTH_FULL_COLUMN": "Full Name!",
                                      "AUTH_HASH": "MD5", "AUTH_SHIFT": "7"}):
            cfg = _v._ws_auth_cfg("workspace_1")
            self.assertEqual(cfg["conn"], "ORCL")
            self.assertEqual(cfg["user_col"], "login_name")
            self.assertEqual(cfg["pass_col"], "pwd_hash")
            self.assertEqual(cfg["full_col"], "")
            self.assertEqual(cfg["hash"], "md5")
            self.assertEqual(cfg["shift"], 7)
        with mock.patch("urs.workspace.read_conf", return_value={}):
            cfg = _v._ws_auth_cfg("workspace_1")
            self.assertEqual(cfg, {"conn": "", "user_col": "", "pass_col": "",
                                   "full_col": "", "hash": "auto", "shift": 3})
        with mock.patch("urs.workspace.read_conf",
                        return_value={"AUTH_HASH": "rot13", "AUTH_SHIFT": "abc"}):
            cfg = _v._ws_auth_cfg("workspace_1")
            self.assertEqual(cfg["hash"], "rot13")
            self.assertEqual(cfg["shift"], 3)
        with mock.patch("urs.workspace.read_conf",
                        return_value={"AUTH_HASH": "vigenere"}):
            self.assertEqual(_v._ws_auth_cfg("workspace_1")["hash"], "auto")

    def test_write_conf_auth_roundtrip(self):
        import tempfile as _tf
        from pathlib import Path as _P
        with _tf.TemporaryDirectory() as td:
            ws = _P(td)
            (ws / "workspace.conf").write_text("# test\n", encoding="utf-8")
            _v._write_ws_conf(ws, "urs_local", "s.t", "active", "", "",
                              auth_conn="ORCL", auth_user_column="login",
                              auth_password_column="pwd", auth_full_column="full_name",
                              auth_hash="sha256")
            from urs import workspace as _wsm
            conf = _wsm.read_conf(ws / "workspace.conf")
            self.assertEqual(conf.get("AUTH_CONN"), "ORCL")
            self.assertEqual(conf.get("AUTH_USER_COLUMN"), "login")
            self.assertEqual(conf.get("AUTH_PASSWORD_COLUMN"), "pwd")
            self.assertEqual(conf.get("AUTH_FULL_COLUMN"), "full_name")
            self.assertEqual(conf.get("AUTH_HASH"), "sha256")
            # Caesar shift round-trips; default shift is not persisted
            _v._write_ws_conf(ws, "urs_local", "s.t", "active", "", "",
                              auth_conn="ORCL", auth_hash="caesar", auth_shift="7")
            conf = _wsm.read_conf(ws / "workspace.conf")
            self.assertEqual(conf.get("AUTH_HASH"), "caesar")
            self.assertEqual(conf.get("AUTH_SHIFT"), "7")
            _v._write_ws_conf(ws, "urs_local", "s.t", "active", "", "",
                              auth_conn="ORCL", auth_hash="caesar", auth_shift="3")
            conf = _wsm.read_conf(ws / "workspace.conf")
            self.assertNotIn("AUTH_SHIFT", conf)
            # clearing drops the keys (auto hash is not persisted)
            _v._write_ws_conf(ws, "urs_local", "s.t", "active", "", "")
            conf = _wsm.read_conf(ws / "workspace.conf")
            for k in ("AUTH_CONN", "AUTH_USER_COLUMN", "AUTH_PASSWORD_COLUMN",
                      "AUTH_FULL_COLUMN", "AUTH_HASH"):
                self.assertNotIn(k, conf)

    def test_modal_has_mapping_selects(self):
        import os as _os
        path = _os.path.join(_os.path.dirname(__file__), "templates", "home.html")
        with open(path, encoding="utf-8") as f:
            src = f.read()
        self.assertIn('id="wsF_user_column"', src)
        self.assertIn('id="wsF_password_column"', src)
        self.assertIn("wsColsFetch", src)
        self.assertIn("users_user_column", src)

    def test_modal_has_external_auth_selects(self):
        import os as _os
        path = _os.path.join(_os.path.dirname(__file__), "templates", "home.html")
        with open(path, encoding="utf-8") as f:
            src = f.read()
        for needle in ('id="wsF_auth_conn"', 'id="wsF_auth_hash"',
                        'id="wsF_auth_user_column"', 'id="wsF_auth_password_column"',
                        'id="wsF_auth_full_column"', 'id="wsF_auth_shift"',
                        "wsAuthConnsFetch", "wsAuthColsFetch", "wsAuthShiftToggle",
                        "auth_conn", "auth_hash", "auth_shift",
                        'value="mysql_sha1"', 'value="pg_md5"',
                        'value="caesar"', 'value="atbash"', 'value="rot13"'):
            self.assertIn(needle, src)


class BranchPasswordContractTests(SimpleTestCase):
    """insert (hashed) / verify (plain) must hold on branch/member saves too:
    password-named columns hash even with a non-password inputType."""

    def test_name_based_detection(self):
        cols = [{"name": "username", "input_type": "text"},
                {"name": "password_hash", "input_type": "text"},
                {"name": "pwd", "input_type": "text"},
                {"name": "notes", "input_type": "textarea"}]
        self.assertEqual(_v._branch_password_cols(None, cols), {"password_hash", "pwd"})

    def test_prep_hashes_plain_and_keeps_hash(self):
        from fmlk_engine.engine import hash_secret as _hs
        already = _hs("h123456*")
        rows = _v._prep_branch_secrets(
            [{"username": "a", "password_hash": "h123456*"},
             {"username": "b", "password_hash": already},
             {"username": "c", "password_hash": "**********"}],
            {"password_hash"})
        self.assertTrue(rows[0]["password_hash"].startswith("pbkdf2_sha256"))
        self.assertNotIn("h123456*", rows[0]["password_hash"])
        self.assertEqual(rows[1]["password_hash"], already)
        self.assertNotIn("password_hash", rows[2])


class _SchemaFakeCur:
    """Schema-aware fake PG cursor: {(schema, table): {cols, rows}}."""

    def __init__(self, tables):
        self.tables = tables
        self.executed = []
        self._rows = []

    def execute(self, sql, params=None):
        import re as _re
        self.executed.append((sql, params))
        if "information_schema.columns" in sql and "table_schema" in sql:
            sch, tbl = (params or ["", ""])[:2]
            t = self.tables.get((sch, tbl))
            self._rows = [(c,) for c in (t["cols"] if t else [])]
        elif "information_schema.tables" in sql:
            tbl = (params or [""])[0]
            self._rows = [(s,) for (s, t) in self.tables if t == tbl]
        elif "count(*)" in sql:
            self._rows = [(0,)]
        elif "LIMIT 1" in sql and "WHERE" in sql:
            m = _re.search(r'FROM "([^"]+)"\."([^"]+)"', sql)
            sch, tbl = (m.group(1), m.group(2)) if m else ("", "")
            u = (params or [""])[0]
            t = self.tables.get((sch, tbl))
            r = (t["rows"].get(u) if t else None)
            self._rows = [r] if r is not None else []
        else:
            self._rows = []

    def fetchall(self):
        return list(self._rows)

    def fetchone(self):
        return self._rows[0] if self._rows else None

    def close(self):
        pass


class _SchemaFakeConn:
    def __init__(self, tables):
        self.cur = _SchemaFakeCur(tables)

    def cursor(self):
        return self.cur

    def commit(self):
        pass

    def close(self):
        pass


class LoginSchemaTests(SimpleTestCase):
    """Users accepted from external + local regardless of local schema."""

    def setUp(self):
        from pathlib import Path as _P
        self.rf = RequestFactory()
        self.dirs = mock.patch("urs.workspace.workspace_dirs",
                               return_value=[_P("workspace_1")])
        self.dirs.start()
        self.addCleanup(self.dirs.stop)

    def _login(self, tables, utable, obj_schema, payload, auth_conn=""):
        from types import SimpleNamespace as _NS
        obj = _NS(schema=obj_schema)
        auth = {"conn": auth_conn, "user_col": "", "pass_col": "",
                "full_col": "", "hash": "auto", "shift": 3}
        with mock.patch.multiple(
                _v,
                _ws_login_connection=mock.Mock(return_value=(obj, utable, "")),
                _ws_pg_connect=mock.Mock(side_effect=lambda _o: _SchemaFakeConn(tables)),
                _ws_auth_cfg=mock.Mock(return_value=dict(auth)),
                _ws_login_row_external=mock.Mock(side_effect=AssertionError("no external expected"))):
            req = self.rf.post("/x/", data=json.dumps(payload), content_type="application/json")
            res = _v.api_workspace_login(req, "workspace_1")
        return res

    def _cols(self):
        return ["username", "password", "is_active"]

    def test_candidates_order(self):
        from types import SimpleNamespace as _NS
        self.assertEqual(_v._ws_login_local_candidates(_NS(schema=""), "IAS20246.USER_R", "fy", False),
                         [("IAS20246", "USER_R")])
        self.assertEqual(_v._ws_login_local_candidates(_NS(schema="app_base"), "app_users", "fy_2026", False),
                         [("fy_2026", "app_users"), ("app_base", "app_users")])
        self.assertEqual(_v._ws_login_local_candidates(_NS(schema="fy_2026"), "app_users", "fy_2026", False),
                         [("fy_2026", "app_users")])
        self.assertEqual(_v._ws_login_local_candidates(_NS(schema="app_base"), "app_users", "fy_2026", True),
                         [("fy_2026", "app_users"), ("fy_2026", "users"),
                          ("app_base", "app_users"), ("app_base", "users")])

    def test_local_falls_back_to_conn_schema(self):
        tables = {("app_base", "app_users"):
                  {"cols": self._cols(), "rows": {"schema_u1": ("schema_u1", "pw1", True)}}}
        res = self._login(tables, "app_users", "app_base",
                          {"schema": "fy_2026", "username": "schema_u1", "password": "pw1"})
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        self.assertEqual(body["user"]["username"], "schema_u1")

    def test_local_prefers_fiscal(self):
        tables = {("fy_2026", "app_users"):
                  {"cols": self._cols(), "rows": {"both_u": ("both_u", "FPW", True)}},
                  ("app_base", "app_users"):
                  {"cols": self._cols(), "rows": {"both_u": ("both_u", "BPW", True)}}}
        res = self._login(tables, "app_users", "app_base",
                          {"schema": "fy_2026", "username": "both_u", "password": "BPW"})
        self.assertEqual(res.status_code, 403, json.loads(res.content.decode()))
        res = self._login(tables, "app_users", "app_base",
                          {"schema": "fy_2026", "username": "both_u", "password": "FPW"})
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)

    def test_dotted_pins_schema(self):
        tables = {("IAS20246", "USER_R"):
                  {"cols": self._cols(), "rows": {"pin_u": ("pin_u", "pw9", True)}}}
        from types import SimpleNamespace as _NS
        seen = {}

        _orig = _SchemaFakeConn

        class _RecConn(_orig):
            def cursor(self):
                return self.cur

        conn = _RecConn(tables)
        obj = _NS(schema="other_sch")
        with mock.patch.multiple(
                _v,
                _ws_login_connection=mock.Mock(return_value=(obj, "IAS20246.USER_R", "")),
                _ws_pg_connect=mock.Mock(return_value=conn),
                _ws_auth_cfg=mock.Mock(return_value={"conn": "", "user_col": "", "pass_col": "",
                                                     "full_col": "", "hash": "auto", "shift": 3}),
                _ws_login_row_external=mock.Mock(side_effect=AssertionError("no external expected"))):
            req = self.rf.post("/x/", data=json.dumps(
                {"schema": "fy_2026", "username": "pin_u", "password": "pw9"}),
                content_type="application/json")
            res = _v.api_workspace_login(req, "workspace_1")
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        for sql, _p in conn.cur.executed:
            self.assertNotIn('"fy_2026"', sql)
            self.assertNotIn('"other_sch"', sql)

    def test_missing_lists_tried_sources(self):
        res = self._login({}, "app_users", "app_base",
                          {"schema": "fy_2026", "username": "ghost", "password": "x"})
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 400, body)
        self.assertIn("بحث أيضاً في", body.get("error", ""))
        self.assertIn("app_base.app_users", body.get("error", ""))

    def test_external_prefers_conn_default_over_fiscal(self):
        calls = []

        class _Cur:
            def __init__(self, probe):
                self.probe = probe

            @property
            def description(self):
                return [("u_id",), ("password",)]

            def fetchone(self):
                return None if self.probe else ("10", "pw10")

            def close(self):
                pass

        class _Shim:
            conn = None

            def _exec(self, sql, _p):
                calls.append(sql)
                return _Cur("WHERE 1=0" in sql)

        auth = {"user_col": "u_id", "pass_col": "password", "full_col": "",
                "hash": "auto", "shift": 3}
        with mock.patch.multiple(
                _v,
                _fmlk_external_exec=mock.Mock(return_value=(_Shim(), "oracle")),
                _fmlk_conn_row=mock.Mock(return_value=object()),
                _default_schema_for=mock.Mock(return_value="IAS20246")):
            row, *_ = _v._ws_login_row_external("ORCL", "USER_R", "fy_2026", "10", dict(auth))
        self.assertEqual(row, ("10", "pw10"))
        self.assertTrue(any('"IAS20246"."USER_R"' in s for s in calls), calls)
        self.assertFalse(any("fy_2026" in s.lower() for s in calls), calls)

    def test_external_dotted_still_pins(self):
        calls = []

        class _Cur:
            def __init__(self, probe):
                self.probe = probe

            @property
            def description(self):
                return [("u_id",), ("password",)]

            def fetchone(self):
                return None if self.probe else ("10", "pw10")

            def close(self):
                pass

        class _Shim:
            conn = None

            def _exec(self, sql, _p):
                calls.append(sql)
                return _Cur("WHERE 1=0" in sql)

        auth = {"user_col": "u_id", "pass_col": "password", "full_col": "",
                "hash": "auto", "shift": 3}
        with mock.patch.multiple(
                _v,
                _fmlk_external_exec=mock.Mock(return_value=(_Shim(), "oracle")),
                _fmlk_conn_row=mock.Mock(return_value=object()),
                _default_schema_for=mock.Mock(return_value="OTHER")):
            row, *_ = _v._ws_login_row_external("ORCL", "IAS20246.USER_R", "fy_2026", "10", dict(auth))
        self.assertEqual(row, ("10", "pw10"))
        self.assertTrue(any('"IAS20246"."USER_R"' in s for s in calls), calls)


class LoginUnlockTests(SimpleTestCase):
    def setUp(self):
        from pathlib import Path as _P
        self.rf = RequestFactory()
        self.dirs = mock.patch("urs.workspace.workspace_dirs",
                               return_value=[_P("workspace_1")])
        self.dirs.start()
        self.addCleanup(self.dirs.stop)

    def _call(self, fn, payload, master=True, method="post"):
        with mock.patch.multiple(
                _v,
                _ws_master_ok=mock.Mock(return_value=master)):
            if method == "get":
                req = self.rf.get("/x/")
            else:
                req = self.rf.post("/x/", data=json.dumps(payload), content_type="application/json")
            return fn(req, "workspace_1")

    def _lock(self, user, ip="9.9.9.9"):
        from django.core.cache import cache
        key = _v._login_throttle_key("workspace_1", user)
        for _ in range(5):
            _v._login_note_fail(ip, key)
        self.addCleanup(cache.delete, key)
        self.addCleanup(cache.delete, "wslogin:ip:" + ip)
        return key

    def test_requires_master(self):
        res = self._call(_v.api_workspace_login_unlock, {"username": "u"}, master=False)
        self.assertEqual(res.status_code, 403)
        self.assertTrue(json.loads(res.content.decode()).get("master_required"))

    def test_method_and_ws(self):
        res = self._call(_v.api_workspace_login_unlock, {}, method="get")
        self.assertEqual(res.status_code, 405)
        with mock.patch.multiple(_v, _ws_master_ok=mock.Mock(return_value=True)):
            req = self.rf.post("/x/", data=json.dumps({"username": "u"}), content_type="application/json")
            res = _v.api_workspace_login_unlock(req, "nope_ws")
        self.assertEqual(res.status_code, 404)

    def test_retry_is_honest(self):
        key = self._lock("unlock_u1")
        blocked, retry = _v._login_locked("9.9.9.9", key)
        self.assertTrue(blocked)
        self.assertGreater(retry, 0)
        self.assertLessEqual(retry, 600)

    def test_unlock_clears_lock(self):
        key = self._lock("unlock_u2")
        blocked, _ = _v._login_locked("9.9.9.9", key)
        self.assertTrue(blocked)
        with mock.patch.multiple(_v, _ws_master_ok=mock.Mock(return_value=True)):
            req = self.rf.post("/x/", data=json.dumps({"username": "unlock_u2"}),
                               content_type="application/json", REMOTE_ADDR="9.9.9.9")
            res = _v.api_workspace_login_unlock(req, "workspace_1")
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        self.assertTrue(body.get("ok"))
        self.assertIn("user", body.get("cleared", []))
        self.assertIn("ip", body.get("cleared", []))
        blocked, _ = _v._login_locked("9.9.9.9", key)
        self.assertFalse(blocked)

    def test_unlock_legacy_int_counter(self):
        from django.core.cache import cache
        key = _v._login_throttle_key("workspace_1", "unlock_legacy")
        cache.set(key, 5, 600)
        self.addCleanup(cache.delete, key)
        blocked, retry = _v._login_locked("9.9.9.8", key)
        self.assertTrue(blocked)
        self.assertEqual(retry, 600)
        res = self._call(_v.api_workspace_login_unlock, {"username": "unlock_legacy"})
        self.assertEqual(res.status_code, 200)

    def test_enter_has_unlock_ui(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "templates", "enter.html"), encoding="utf-8") as fh:
            html = fh.read()
        for needle in ("liUnlock", "wsLoginUnlock", "login/unlock/", "master/auth/",
                       "liMaster", "d.locked", "فك القفل"):
            self.assertIn(needle, html)


class LoginDiagResetTests(SimpleTestCase):
    def setUp(self):
        from pathlib import Path as _P
        self.rf = RequestFactory()
        self.dirs = mock.patch("urs.workspace.workspace_dirs",
                               return_value=[_P("workspace_1")])
        self.dirs.start()
        self.addCleanup(self.dirs.stop)

    def _call(self, state, fn, payload, master=True):
        with mock.patch.multiple(
                _v,
                _ws_master_ok=mock.Mock(return_value=master),
                _ws_login_connection=mock.Mock(return_value=({"fake": True}, "sys_users", "")),
                _ws_pg_connect=mock.Mock(side_effect=lambda _obj: _FakeConn(state))):
            req = self.rf.post("/x/", data=json.dumps(payload), content_type="application/json")
            return fn(req, "workspace_1"), state

    def test_hash_kind(self):
        import hashlib as _hl
        from django.contrib.auth.hashers import make_password
        k = _v._ws_hash_kind
        self.assertEqual(k(""), "empty")
        self.assertEqual(k(make_password("x")), "django-hash")
        import os as _os
        from fmlk_engine.engine import _STD_HASH_PREFIX, _STD_HASH_ITERS
        salt = _os.urandom(16).hex()
        dk = _hl.pbkdf2_hmac("sha256", b"x", bytes.fromhex(salt), _STD_HASH_ITERS).hex()
        self.assertEqual(k("%s%d$%s$%s" % (_STD_HASH_PREFIX, _STD_HASH_ITERS, salt, dk)), "std-pbkdf2")
        self.assertEqual(k("pbkdf2_sha256$100000$" + "ab" * 16 + "$" + "cd" * 32), "foreign-hex-pbkdf2")
        self.assertEqual(k("ab" * 32), "hex-digest")
        self.assertEqual(k("plain123"), "plaintext-or-unknown")

    def test_diag_unknown_user(self):
        st = _fresh_state()
        res, _ = self._call(st, _v.api_workspace_login_diag, {"username": "ghost"})
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200)
        self.assertFalse(body["found"])
        self.assertEqual(body["verdict"], "unknown-user")

    def test_diag_bad_password_shape(self):
        from django.contrib.auth.hashers import make_password
        st = _fresh_state(userrow=("admin", make_password("real-pw-1"), True))
        res, _ = self._call(st, _v.api_workspace_login_diag, {"username": "admin"})
        body = json.loads(res.content.decode())
        self.assertTrue(body["found"])
        self.assertEqual(body["verdict"], "bad-password")
        self.assertEqual(body["hash_format"], "django-hash")
        self.assertTrue(body["code_selftest"])

    def test_diag_duplicates(self):
        from django.contrib.auth.hashers import make_password
        st = _fresh_state(count=2, userrow=("admin", make_password("x"), True))
        res, _ = self._call(st, _v.api_workspace_login_diag, {"username": "admin"})
        body = json.loads(res.content.decode())
        self.assertEqual(body["verdict"], "duplicates")

    def test_diag_inactive(self):
        from django.contrib.auth.hashers import make_password
        st = _fresh_state(userrow=("admin", make_password("x"), False))
        res, _ = self._call(st, _v.api_workspace_login_diag, {"username": "admin"})
        body = json.loads(res.content.decode())
        self.assertEqual(body["verdict"], "inactive")

    def test_diag_requires_master(self):
        st = _fresh_state()
        res, _ = self._call(st, _v.api_workspace_login_diag, {"username": "admin"}, master=False)
        self.assertEqual(res.status_code, 403)

    def test_reset_ok(self):
        from fmlk_engine.engine import verify_secret as _vs
        st = _fresh_state(exists=True)
        res, _ = self._call(st, _v.api_workspace_login_reset,
                            {"username": "admin", "new_password": "h123456*"})
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        sql, params = st["update"]
        self.assertIn("UPDATE", sql)
        self.assertTrue(_vs("h123456*", params[0]))
        self.assertNotIn("h123456*", params)

    def test_reset_short_and_unknown(self):
        st = _fresh_state(exists=True)
        res, _ = self._call(st, _v.api_workspace_login_reset,
                            {"username": "admin", "new_password": "123"})
        self.assertEqual(res.status_code, 400)
        self.assertIsNone(st["update"])
        st = _fresh_state(exists=False)
        res, _ = self._call(st, _v.api_workspace_login_reset,
                            {"username": "ghost", "new_password": "h123456*"})
        self.assertEqual(res.status_code, 404)

    def test_modal_has_diag_ui(self):
        import os as _os
        path = _os.path.join(_os.path.dirname(__file__), "templates", "home.html")
        with open(path, encoding="utf-8") as f:
            src = f.read()
        self.assertIn('id="wsDg_user"', src)
        self.assertIn("wsLoginDiag", src)
        self.assertIn("wsLoginReset", src)
        self.assertIn("login-diag", src)


class LoginDualSourceTests(SimpleTestCase):
    """Login verifies BOTH connections: external auth source first, then local primary."""

    def setUp(self):
        from pathlib import Path as _P
        self.rf = RequestFactory()
        self.dirs = mock.patch("urs.workspace.workspace_dirs",
                               return_value=[_P("workspace_1")])
        self.dirs.start()
        self.addCleanup(self.dirs.stop)

    def _login(self, state, ext_row=None, ext_exc=None,
               auth_cfg=None, payload=None):
        _auth = {"conn": "ORCL", "user_col": "", "pass_col": "", "full_col": "",
                 "hash": "auto", "shift": 3}
        if auth_cfg is not None:
            _auth.update(auth_cfg)

        def _fake_ext(*a, **kw):
            if ext_exc is not None:
                raise ext_exc
            return ext_row

        _pg = mock.Mock(side_effect=lambda _obj: _FakeConn(state))
        with mock.patch.multiple(
                _v,
                _ws_master_ok=mock.Mock(return_value=True),
                _ws_login_connection=mock.Mock(return_value=({"fake": True}, "sys_users", "")),
                _ws_pg_connect=_pg,
                _ws_auth_cfg=mock.Mock(return_value=dict(_auth)),
                _ws_login_row_external=mock.Mock(side_effect=_fake_ext)):
            req = self.rf.post("/x/", data=json.dumps(payload or
                                {"schema": "main_hq_2026", "username": "admin", "password": "pw123"}),
                               content_type="application/json")
            res = _v.api_workspace_login(req, "workspace_1")
        return res, _pg

    def test_external_hit_wins_without_touching_local(self):
        st = _fresh_state()
        res, _pg = self._login(
            st, ext_row=(("extuser", "extpw", True), "is_active", None, None, "IAS", "LOGIN", "PWD"),
            payload={"schema": "main_hq_2026", "username": "extuser", "password": "extpw"})
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        self.assertTrue(body.get("ok"))
        self.assertEqual(body["user"]["username"], "extuser")
        _pg.assert_not_called()

    def test_external_miss_falls_back_to_local(self):
        st = _fresh_state(userrow=("admin", "pw123", True))
        res, _pg = self._login(st, ext_row=None)
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        self.assertEqual(body["user"]["username"], "admin")
        self.assertTrue(_pg.called)

    def test_external_unreachable_falls_back_to_local(self):
        st = _fresh_state(userrow=("admin", "pw123", True))
        res, _pg = self._login(st, ext_row=None, ext_exc=RuntimeError("down"))
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        self.assertEqual(body["user"]["username"], "admin")

    def test_external_bad_password_is_final(self):
        st = _fresh_state(userrow=("admin", "pw123", True))
        res, _pg = self._login(
            st, ext_row=(("admin", "extpw", True), "is_active", None, None, "IAS", "LOGIN", "PWD"),
            payload={"schema": "main_hq_2026", "username": "admin", "password": "pw123"})
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 403, body)
        _pg.assert_not_called()

    def test_external_config_error_is_final(self):
        st = _fresh_state(userrow=("admin", "pw123", True))
        res, _pg = self._login(st, ext_row=None, ext_exc=ValueError("bad column"))
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 400, body)
        _pg.assert_not_called()

    def test_no_auth_conn_uses_local_only(self):
        st = _fresh_state(userrow=("admin", "pw123", True))
        with mock.patch.multiple(
                _v,
                _ws_master_ok=mock.Mock(return_value=True),
                _ws_login_connection=mock.Mock(return_value=({"fake": True}, "sys_users", "")),
                _ws_pg_connect=mock.Mock(side_effect=lambda _obj: _FakeConn(st)),
                _ws_auth_cfg=mock.Mock(return_value={"conn": "", "user_col": "", "pass_col": "",
                                                     "full_col": "", "hash": "auto", "shift": 3}),
                _ws_login_row_external=mock.Mock(side_effect=AssertionError("must not query external"))):
            req = self.rf.post("/x/", data=json.dumps(
                {"schema": "main_hq_2026", "username": "admin", "password": "pw123"}),
                content_type="application/json")
            res = _v.api_workspace_login(req, "workspace_1")
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)

    def test_enter_page_marks_external_auth(self):
        import os as _os
        path = _os.path.join(_os.path.dirname(__file__), "templates", "enter.html")
        with open(path, encoding="utf-8") as f:
            src = f.read()
        self.assertIn("has_external_auth", src)
        self.assertIn("مصدر التخويل", src)


class LoginDiagExternalTests(SimpleTestCase):
    """login-diag against the external auth source (saved conf or draft)."""

    def setUp(self):
        from pathlib import Path as _P
        self.rf = RequestFactory()
        self.dirs = mock.patch("urs.workspace.workspace_dirs",
                               return_value=[_P("workspace_1")])
        self.dirs.start()
        self.addCleanup(self.dirs.stop)

    def _diag(self, payload, auth_cfg=None, ext_row=None, dupes=0):
        _auth = {"conn": "ORCL", "user_col": "", "pass_col": "", "full_col": "",
                 "hash": "md5", "shift": 3}
        if auth_cfg is not None:
            _auth.update(auth_cfg)
        with mock.patch.multiple(
                _v,
                _ws_master_ok=mock.Mock(return_value=True),
                _ws_auth_cfg=mock.Mock(return_value=dict(_auth)),
                _ws_login_row_external=mock.Mock(
                    return_value=(ext_row, "is_active", None, None, "IAS", "LOGIN", "PWD")
                    if ext_row is not None else (None, "is_active", None, None, "IAS", "LOGIN", "PWD")),
                _ws_login_count_external=mock.Mock(return_value=dupes)):
            req = self.rf.post("/x/", data=json.dumps(payload),
                               content_type="application/json")
            return _v.api_workspace_login_diag(req, "workspace_1")

    def test_diag_external_ready(self):
        import hashlib as _hl
        stored = _hl.md5("pw123".encode()).hexdigest()
        res = self._diag({"username": "omar", "table": "IAS.USERS"},
                         ext_row=("omar", stored, True))
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        self.assertEqual(body["source"], "external")
        self.assertEqual(body["auth_conn"], "ORCL")
        self.assertTrue(body["found"])
        self.assertTrue(body["shape_ok"])
        self.assertEqual(body["verdict"], "ready")

    def test_diag_external_unknown_and_mismatch(self):
        res = self._diag({"username": "ghost", "table": "IAS.USERS"}, ext_row=None)
        body = json.loads(res.content.decode())
        self.assertEqual(body["verdict"], "unknown-user")
        res = self._diag({"username": "omar", "table": "IAS.USERS"},
                         ext_row=("omar", "plaintext-pw", True))
        body = json.loads(res.content.decode())
        self.assertEqual(body["verdict"], "hash-mismatch")
        self.assertFalse(body["shape_ok"])

    def test_diag_external_draft_conn(self):
        import hashlib as _hl
        stored = _hl.md5("pw123".encode()).hexdigest()
        res = self._diag({"username": "omar", "table": "IAS.USERS", "conn": "ORCL2",
                          "hash": "md5"},
                         auth_cfg={"conn": "", "user_col": "", "pass_col": "",
                                   "full_col": "", "hash": "auto", "shift": 3},
                         ext_row=("omar", stored, True))
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        self.assertEqual(body["auth_conn"], "ORCL2")

    def test_diag_external_missing_table(self):
        with mock.patch("urs.workspace.read_conf", return_value={}):
            res = self._diag({"username": "omar"})
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 400, body)

    def test_auth_shape_ok(self):
        self.assertTrue(_v._ws_auth_shape_ok("auto", "anything"))
        self.assertTrue(_v._ws_auth_shape_ok("md5", "d" * 32))
        self.assertFalse(_v._ws_auth_shape_ok("md5", "short"))
        self.assertTrue(_v._ws_auth_shape_ok("sha256", "e" * 64))
        self.assertTrue(_v._ws_auth_shape_ok("mysql_sha1", "*" + "A" * 40))
        self.assertFalse(_v._ws_auth_shape_ok("mysql_sha1", "A" * 40))
        self.assertTrue(_v._ws_auth_shape_ok("pg_md5", "md5" + "b" * 32))
        self.assertTrue(_v._ws_auth_shape_ok("caesar", "anything-raw"))
        self.assertTrue(_v._ws_auth_shape_ok("caesar", "caesar$5$abc"))
        self.assertFalse(_v._ws_auth_shape_ok("caesar", "caesar$"))
        self.assertTrue(_v._ws_auth_shape_ok("rot13", "rot13$abc"))
        self.assertFalse(_v._ws_auth_shape_ok("rot13", "rot13$"))

    def test_external_row_quotes_real_case(self):
        # Oracle "u_id" != U_ID — the fetch must quote the TRUE probed case
        seen = []

        class _Cur:
            description = [("U_ID",), ("PWD_HASH",)]

            def fetchone(self):
                return None

            def fetchall(self):
                return []

            def close(self):
                pass

        class _Shim:
            conn = None

            def _exec(self, sql, params):
                seen.append(sql)
                return _Cur()

        with mock.patch.object(_v, "_fmlk_external_exec", return_value=(_Shim(), "oracle")):
            row = _v._ws_login_row_external(
                "ORCL", "IAS.USER_R", "", "omar",
                {"user_col": "u_id", "pass_col": "pwd_hash", "full_col": "",
                 "hash": "auto", "shift": 3})
        self.assertIsNone(row[0])
        _sels = [s for s in seen if "WHERE 1=0" not in s]
        self.assertTrue(_sels)
        self.assertIn('"U_ID"', _sels[0])
        self.assertIn('"PWD_HASH"', _sels[0])
        self.assertNotIn('"u_id"', _sels[0])

    def test_modal_has_auth_diag_ui(self):
        import os as _os
        path = _os.path.join(_os.path.dirname(__file__), "templates", "home.html")
        with open(path, encoding="utf-8") as f:
            src = f.read()
        for needle in ("wsAuthDiag", "wsAuthDg_user", "wsAuthDgMsg", "فحص المطابقة"):
            self.assertIn(needle, src)


class FormulasMigrateTests(SimpleTestCase):
    def test_generate_shapes(self):
        from rml_python import formulas_migrate as _fm
        g = _fm.generate_ddl("public")
        self.assertTrue(g["total"] > 50)
        self.assertTrue(all(f["sql"].startswith("CREATE OR REPLACE FUNCTION") for f in g["functions"]))
        names = {f["name"] for f in g["functions"]}
        self.assertIn("UPPER", names)
        self.assertIn("IF", names)
        skipped = {s["name"] for s in g["skipped"]}
        self.assertIn("SUM", skipped)
        self.assertIn("XLOOKUP", skipped)
        self.assertEqual(_fm.pg_name("Upper"), "fn_upper")
        one = _fm.function_ddl("s1", "UPPER")
        self.assertIn('"s1"."fn_upper"', one["sql"])
        with self.assertRaises(KeyError):
            _fm.function_ddl("public", "XLOOKUP")

    def test_translate_and_expr(self):
        from rml_python import formulas_migrate as _fm
        self.assertEqual(_fm.translate_formula("=UPPER([name])"), 'UPPER("name")')
        self.assertEqual(_fm.translate_formula("=IF([a], [b], 'x')"), 'IF("a", "b", \'x\')')
        with self.assertRaises(ValueError):
            _fm.translate_formula("=XLOOKUP([a], [b])")
        with self.assertRaises(ValueError):
            _fm.translate_formula("=UPPER([a]); DROP TABLE t")
        self.assertEqual(_fm.expr_to_sql("price * 1.15"), "price * 1.15")
        self.assertEqual(_fm.expr_to_sql("=UPPER([n])"), 'UPPER("n")')
        self.assertEqual(_fm.expr_to_sql("XSQL: SELECT id FROM t"), "( SELECT id FROM t )")
        with self.assertRaises(ValueError):
            _fm.expr_to_sql("XSQL: SELECT XLOOKUP(a) FROM t")
        with self.assertRaises(ValueError):
            _fm.expr_to_sql("")

    def test_validate_call_and_ddl(self):
        from rml_python import formulas_migrate as _fm
        self.assertEqual(_fm.validate_call("=myschema.fn_upper(name)"), "myschema.fn_upper(name)")
        with self.assertRaises(ValueError):
            _fm.validate_call("not a call")
        with self.assertRaises(ValueError):
            _fm.validate_call("fn_x(a); DROP TABLE t")
        meta = _fm.validate_ddl("CREATE OR REPLACE FUNCTION public.fn_a(TEXT) RETURNS TEXT LANGUAGE sql AS $$ SELECT $1 $$;")
        self.assertEqual((meta["schema"], meta["name"]), ("public", "fn_a"))
        with self.assertRaises(ValueError):
            _fm.validate_ddl("DROP TABLE t;")
        with self.assertRaises(ValueError):
            _fm.validate_ddl("CREATE OR REPLACE FUNCTION public.fn_a() RETURNS TEXT AS $$ SELECT 1 $$; DELETE FROM t;")

    def test_rulevars_expression_kinds(self):
        from rml_python import rulevars as _rv
        self.assertEqual(_rv._value_sql("expression", "=UPPER([n])"), 'UPPER("n")')
        self.assertEqual(_rv._value_sql("expression", "price * 2"), "price * 2")
        self.assertEqual(_rv._value_sql("custom_function", "fn_upper(name)"), "fn_upper(name)")
        with self.assertRaises(ValueError):
            _rv._value_sql("custom_function", "hello world")
        with self.assertRaises(ValueError):
            _rv._value_sql("expression", "=XLOOKUP([a])")

    def test_formulas_endpoints(self):
        req = RequestFactory().get("/api/formulas/")
        res = _v.api_formulas_list(req)
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200)
        self.assertTrue(body["total"] > 50)
        req = RequestFactory().post("/api/formulas/migrate/",
                                    data=json.dumps({"schema": "public", "functions": ["UPPER", "NOPE"]}),
                                    content_type="application/json")
        res = _v.api_formulas_migrate(req)
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200)
        self.assertFalse(body["executed"])
        self.assertEqual(len(body["ddl"]), 1)
        self.assertEqual(body["skipped"][0]["name"], "NOPE")
        # execute without master → 403, no DB touched
        req = RequestFactory().post("/api/formulas/migrate/",
                                    data=json.dumps({"connection": "x", "execute": True}),
                                    content_type="application/json")
        res = _v.api_formulas_migrate(req)
        self.assertEqual(res.status_code, 403)
        req = RequestFactory().post("/api/formulas/ddl/",
                                    data=json.dumps({"connection": "x", "ddl": "DROP TABLE t"}),
                                    content_type="application/json")
        res = _v.api_formulas_ddl(req)
        self.assertEqual(res.status_code, 403)

    def test_autocomplete_endpoint(self):
        from rml_python import formulas_migrate as _fm
        defs = _fm.autocomplete_defaults()
        self.assertTrue(len(defs) > 60)
        self.assertTrue(all({"fn", "tpl", "hint", "kind"} <= set(d) for d in defs))
        kinds = {d["kind"] for d in defs}
        self.assertEqual(kinds, {"formula", "xsql"})
        req = RequestFactory().get("/api/formulas/autocomplete/")
        res = _v.api_formulas_autocomplete(req)
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200)
        self.assertTrue(len(body["defaults"]) > 60)
        self.assertEqual(body["migrated"], [])
        self.assertEqual(body["custom"], [])
        req = RequestFactory().get("/api/formulas/autocomplete/", {"connection": "nope", "schema": "public"})
        res = _v.api_formulas_autocomplete(req)
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200)
        self.assertIn("note", body)

    def test_bare_aliases(self):
        from rml_python import formulas_migrate as _fm
        a = _fm.alias_ddl("public", "IF")
        self.assertEqual(a["name"], "if")
        self.assertEqual(a["alias_of"], "fn_if")
        self.assertIn('"public"."fn_if"($1, $2, $3)', a["sql"])
        self.assertIsNone(_fm.alias_ddl("public", "UPPER"))
        g = _fm.generate_ddl("public", ["IF", "UPPER"])
        self.assertEqual([f["name"] for f in g["functions"]], ["IF", "UPPER"])
        self.assertEqual([a["name"] for a in g["aliases"]], ["if"])
        g2 = _fm.generate_ddl("public", ["IF"], alias=False)
        self.assertEqual(g2["aliases"], [])
        # bare calls pass the expression translator untouched
        self.assertEqual(_fm.translate_formula("=myfunc([x])"), 'myfunc("x")')
        self.assertEqual(_fm.validate_call("myfunc(a, b)"), "myfunc(a, b)")

    def test_migrate_preview_has_aliases(self):
        req = RequestFactory().post("/api/formulas/migrate/",
                                    data=json.dumps({"schema": "public", "functions": ["IF"]}),
                                    content_type="application/json")
        res = _v.api_formulas_migrate(req)
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200)
        self.assertEqual(body["aliases"], ["if"])
        self.assertEqual(len(body["ddl"]), 2)

    def test_autocomplete_all_sources(self):
        from rml_python import formulas_migrate as _fm
        natives = _fm.native_entries()
        self.assertTrue(len(natives) > 40)
        self.assertTrue(all({"fn", "tpl", "hint", "kind"} <= set(d) for d in natives))
        self.assertTrue(all(d["kind"] == "native" for d in natives))
        fns = {d["fn"] for d in natives}
        self.assertTrue({"UPPER", "COALESCE", "STRING_AGG", "ROW_NUMBER"} <= fns)
        req = RequestFactory().get("/api/formulas/autocomplete/", {"all": "1"})
        res = _v.api_formulas_autocomplete(req)
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200)
        self.assertTrue(len(body["natives"]) > 40)
        self.assertTrue(len(body["defaults"]) > 60)
        self.assertEqual(body["migrated"], [])
        self.assertEqual(body["custom"], [])

    def test_wizard_autocomplete_fetch(self):
        import os as _os
        path = "rml_python/shared/rml_wizard_script.html"
        with open(path, encoding="utf-8") as f:
            src = f.read()
        for token in ("fxEqMatches", "fxFnsFetch", "fxFnPool", "fxFnFilter",
                      "FX_FN_CACHE", "/api/formulas/autocomplete/?all=1"):
            self.assertIn(token, src)

    def test_fx_modal_select_all(self):
        import os as _os
        base = _os.path.dirname(__file__)
        with open(_os.path.join(base, "templates", "report_designer.html"), encoding="utf-8") as f:
            src = f.read()
        for token in ("fxToggleAll", "fxSyncAll", 'id="fxAll"', "تحديد الكل"):
            self.assertIn(token, src)

    def test_designer_has_fx_ui(self):
        import os as _os
        base = _os.path.dirname(__file__)
        with open(_os.path.join(base, "templates", "report_designer.html"), encoding="utf-8") as f:
            src = f.read()
        for token in ("fxOpen", "fxMigrate", "fxDdlRun", "/api/formulas/migrate/", "دالة مخصصة"):
            self.assertIn(token, src)

    def test_designer_autosave_markup(self):
        import os as _os
        base = _os.path.dirname(__file__)
        with open(_os.path.join(base, "templates", "report_designer.html"), encoding="utf-8") as f:
            tpl = f.read()
        for token in ("autosaveToggle", "autosaveStatus", "wbAutosaveToggle", "تلقائي"):
            self.assertIn(token, tpl)
        with open(_os.path.join(base, "..", "rml_python", "shared",
                                "rml_wizard_script.html"), encoding="utf-8") as f:
            js = f.read()
        for token in ("wbFillWizardFromData", "wbAutosaveCollect", "wbAutosaveTick",
                      "wbAutosaveToggle", "wbAutosaveClear", "wbAutosaveRestore",
                      "wbAutosaveDiscard", "wbAutosaveBoot", "autosaveBanner",
                      "استعادة المسودة", "rml_autosave_v1_", "rml_autosave_on",
                      "report-designer/"):
            self.assertIn(token, js)


class RmlSortDedupeTests(SimpleTestCase):
    """Sorting must survive duplicated SELECT aliases (MSSQL 209): ORDER BY
    uses the 1-based ordinal of the first matching column; exact-duplicate
    select items are dropped."""

    def _cols(self):
        from types import SimpleNamespace as _NS
        return [_NS(alias="رقم الاكسبرس", name="PublicNumber", expr="T0.PublicNumber"),
                _NS(alias="تاريخ الحوالة", name="TheDate", expr="T0.TheDate"),
                _NS(alias="تاريخ الحوالة", name="TheDate", expr="T0.TheDate"),
                _NS(alias="المحافظة", name="ProvinceName", expr="T5.ProvinceName")]

    def test_order_by_ordinal_first_occurrence(self):
        from rml_python.engine import _build_order_by
        cols = self._cols()
        self.assertEqual(_build_order_by({"column": "تاريخ الحوالة", "direction": "desc"}, cols),
                         " ORDER BY 2 DESC")
        self.assertEqual(_build_order_by([{"column": "المحافظة"}], cols), " ORDER BY 4 ASC")

    def test_order_by_fallback_alias(self):
        from rml_python.engine import _build_order_by
        self.assertEqual(_build_order_by({"column": "ghost", "direction": "xx"}, []),
                         ' ORDER BY "ghost" ASC')

    def test_order_by_formula_matches_stored_expr(self):
        # ORA-00972 regression: player sends the resolved formula; backend
        # must resolve it to the SELECT ordinal, never quote it as identifier.
        from types import SimpleNamespace as _NS
        from rml_python.engine import _build_order_by
        formula = "=if(@priceMethod=o1,[16.IAS_ITEM_PRICE.I_PRICE],[PRIMARY_COST])"
        cols = [_NS(alias="رمز الصنف", name="I_CODE", expr="T0.I_CODE"),
                _NS(alias="التسعير", name="price", expr=formula)]
        self.assertEqual(_build_order_by({"column": formula, "direction": "desc"}, cols),
                         " ORDER BY 2 DESC")

    def test_order_by_formula_eq_insensitive(self):
        from types import SimpleNamespace as _NS
        from rml_python.engine import _build_order_by
        cols = [_NS(alias="التسعير", name="price", expr="if(@x=1,[t.c],[d])")]
        self.assertEqual(_build_order_by({"column": "=IF(@x=1,[t.c],[d])"}, cols),
                         " ORDER BY 1 ASC")

    def test_order_by_unresolvable_formula_skipped(self):
        from types import SimpleNamespace as _NS
        from rml_python.engine import _build_order_by
        cols = [_NS(alias="A", name="a", expr="t.a")]
        self.assertEqual(_build_order_by({"column": "=if(@p=1,[9.t.c],[d])", "direction": "desc"}, cols), "")
        self.assertEqual(_build_order_by([{"column": "=sum([t.x])"}, {"column": "A", "direction": "desc"}], cols),
                         " ORDER BY 1 DESC")
        # plain (non-formula) unknowns keep the old quoted fallback
        self.assertEqual(_build_order_by({"column": "ghost"}, cols), ' ORDER BY "ghost" ASC')

    def test_dedupe_select_items(self):
        from rml_python.engine import _dedupe_select_items
        items = ['"T0"."A" AS "x"', '"T0"."B" AS "y"', '"T0"."A" AS "x"',
                 '"T0"."C" AS "x"']
        self.assertEqual(_dedupe_select_items(items),
                         ['"T0"."A" AS "x"', '"T0"."B" AS "y"', '"T0"."C" AS "x"'])
        self.assertEqual(_dedupe_select_items([]), [])


class WhereParamGuardTests(SimpleTestCase):
    """ORA-00904 regression: doc-param refs (@LEV) and raw formulas must
    never reach SQL WHERE as quoted identifiers."""

    def test_unused_tables_never_enter_plan(self):
        # importing a table's fields must not route/JOIN it anywhere until
        # some expression actually references it (perf invariant).
        from types import SimpleNamespace as _NS
        from rml_python.engine import RMLReportEngine
        eng = RMLReportEngine.__new__(RMLReportEngine)
        cols = [_NS(alias="A", name="a", expr="[MST.A]")]
        field_table = {"a": "MST", "f0": "BIG"}
        used = eng._used_tables(
            cols, [{"field": "A", "op": "equals", "value": "1"}],
            None, None, field_table, "")
        self.assertEqual(used, {"MST"})

    def test_param_ref_and_formula_skipped(self):
        from rml_python.engine import _build_where
        sql, params = _build_where([
            {"field": "@LEV", "op": "equals", "value": "L1"},
            {"field": "=if(@x=1,[t.c],[d])", "op": "equals", "value": "v"},
            {"field": "code", "op": "equals", "value": "x"},
        ])
        self.assertNotIn("@LEV", sql)
        self.assertNotIn("=if(", sql)
        self.assertIn('"code"', sql)
        self.assertEqual(len(params), 1)
        self.assertIn("x", list(params.values()))

    def test_drop_keys_skip_expression_only_params(self):
        from rml_python.engine import _build_where
        sql, params = _build_where(
            [{"field": "lev_no", "op": "equals", "value": "L1", "refname": "@LEV"},
             {"field": "code", "op": "equals", "value": "x"}],
            drop_keys={"@lev", "lev_no"})
        self.assertNotIn("lev_no", sql.lower().replace('"code"', ''))
        self.assertIn('"code"', sql)
        self.assertEqual(len(params), 1)

    def test_normal_filters_untouched(self):
        from rml_python.engine import _build_where
        sql, params = _build_where([
            {"field": "code", "op": "equals", "value": "x"},
            {"field": "name", "op": "contains", "value": "y"},
        ])
        self.assertIn('"code"', sql)
        self.assertIn("LIKE", sql)
        self.assertEqual(len(params), 2)

    def test_unknown_type_plain_text_casts(self):
        # ORA-01722: 'L1' vs unknown-typed LEV_NO → text compare, never converts
        from rml_python.engine import _build_where
        sql, params = _build_where([{"field": "LEV_NO", "op": "equals", "value": "L1"}])
        self.assertIn('CAST("LEV_NO" AS VARCHAR(4000))=:p1', sql)
        self.assertEqual(params, {"p1": "L1"})

    def test_known_types_keep_legacy_paths(self):
        from types import SimpleNamespace as _NS
        from rml_python.engine import _build_where
        num_fields = [_NS(name="lev", data_type="INTEGER")]
        sql, _ = _build_where([{"field": "lev", "op": "equals", "value": "L1"}], fields=num_fields)
        self.assertEqual(sql.strip(), "WHERE 1=0")
        txt_fields = [_NS(name="code", data_type="VARCHAR")]
        sql, params = _build_where([{"field": "code", "op": "equals", "value": "L1"}], fields=txt_fields)
        self.assertIn('"code"=:p1', sql)
        self.assertNotIn("CAST", sql)
        sql, _ = _build_where([{"field": "LEV_NO", "op": "equals", "value": "5"}])
        self.assertIn('"LEV_NO"=:p1', sql)
        self.assertNotIn("CAST", sql)

    def test_between_and_in_cast(self):
        from rml_python.engine import _build_where
        sql, _ = _build_where([{"field": "LEV_NO", "op": "between", "valFrom": "L1", "valTo": "L9"}])
        self.assertIn("CAST(", sql)
        self.assertIn("BETWEEN", sql)
        sql, _ = _build_where([{"field": "LEV_NO", "op": "in", "value": ["L1", "L2"]}])
        self.assertIn("CAST(", sql)
        self.assertIn("IN (", sql)

    def test_placeholder_custom_column_skipped(self):
        from rml_python.engine import _build_where
        sql, params = _build_where([
            {"field": "v_custom_1", "op": "equals", "value": "L1"},
            {"field": "code", "op": "equals", "value": "x"},
        ])
        self.assertNotIn("v_custom", sql)
        self.assertIn('"code"', sql)
        self.assertEqual(len(params), 1)

    def test_no_filter_keys_require_explicit_flag(self):
        from types import SimpleNamespace as _NS
        from rml_python.engine import RMLReportEngine
        eng = RMLReportEngine.__new__(RMLReportEngine)
        dp_linked_custom = _NS(column="LEV_NO", param_refname="@LEV", id="dp_1",
                               no_filter=False, is_custom=True)
        dp_expr_only = _NS(column="v_custom_2", param_refname="@PM", id="dp_2",
                           no_filter=True, is_custom=True)
        eng.compiler = _NS(doc_params=lambda: [dp_linked_custom, dp_expr_only])
        keys = eng._doc_no_filter_keys()
        # linked custom WITHOUT the flag must filter normally
        self.assertNotIn("lev_no", keys)
        self.assertNotIn("@lev", keys)
        # explicit expression-only flag still excludes (incl. its placeholder)
        self.assertIn("v_custom_2", keys)
        self.assertIn("@pm", keys)
        # ...and the WHERE builder honors it end to end
        from rml_python.engine import _build_where
        sql, _ = _build_where(
            [{"field": "LEV_NO", "op": "equals", "value": "1"}],
            drop_keys=keys)
        self.assertIn("LEV_NO", sql)


class XlookupTranspileTests(SimpleTestCase):
    """ORA-00904 regression: same-scope XLOOKUP() in RML formulas must
    compile to CASE (never reach Oracle SQL verbatim)."""

    def test_basic_three_args(self):
        from rml_python.namespaces import _transpile_xlookup_calls
        out = _transpile_xlookup_calls("XLOOKUP('L1', [LEV_NO], [I_PRICE])")
        self.assertEqual(out, "CASE WHEN CAST([LEV_NO] AS VARCHAR(4000)) = 'L1' THEN [I_PRICE] ELSE NULL END")

    def test_four_args_default(self):
        from rml_python.namespaces import _transpile_xlookup_calls
        out = _transpile_xlookup_calls("XLOOKUP('L1', [LEV_NO], [I_PRICE], 0)")
        self.assertEqual(out, "CASE WHEN CAST([LEV_NO] AS VARCHAR(4000)) = 'L1' THEN [I_PRICE] ELSE 0 END")

    def test_numeric_literal_stays_direct(self):
        from rml_python.namespaces import _transpile_xlookup_calls
        out = _transpile_xlookup_calls("XLOOKUP(1, [LEV_NO], [I_PRICE])")
        self.assertEqual(out, "CASE WHEN [LEV_NO] = 1 THEN [I_PRICE] ELSE NULL END")

    def test_case_insensitive_and_nested(self):
        from rml_python.namespaces import _transpile_xlookup_calls
        out = _transpile_xlookup_calls(
            "xlookup('L1', [LEV_NO], IF(@priceMethod=o1, [I_PRICE], [PRIMARY_COST]))")
        self.assertIn("CASE WHEN", out)
        self.assertNotIn("XLOOKUP", out.upper().replace("CASE", ""))
        self.assertIn("'L1'", out)

    def test_string_literal_with_name_inside_untouched(self):
        from rml_python.namespaces import _transpile_xlookup_calls
        s = "SELECT 'XLOOKUP(' AS c"
        self.assertEqual(_transpile_xlookup_calls(s), s)

    def test_wrong_arity_raises_arabic(self):
        from rml_python.namespaces import _transpile_xlookup_calls
        with self.assertRaises(ValueError) as cm:
            _transpile_xlookup_calls("XLOOKUP('L1', [LEV_NO])")
        self.assertIn("XLOOKUP", str(cm.exception))

    def test_non_formula_passthrough(self):
        # hand-written SQL without leading '=' keeps XLOOKUP verbatim
        from rml_python.namespaces import _resolve_expression
        self.assertIn("XLOOKUP", _resolve_expression("XLOOKUP(a, b, c)", [], {}, set()))

    def test_user_formula_end_to_end_shape(self):
        from rml_python.namespaces import _transpile_xlookup_calls
        out = _transpile_xlookup_calls(
            "XLOOKUP('L1', [LEV_NO], CASE WHEN 'o1'='o1' THEN [I_PRICE] ELSE [PRIMARY_COST] END)")
        self.assertTrue(out.startswith("CASE WHEN "))
        self.assertIn("CAST([LEV_NO] AS VARCHAR(4000)) = 'L1'", out)
        self.assertIn("ELSE NULL END", out)


class RmlRefnameTests(SimpleTestCase):
    """col_refname: @ref resolves to the final row VALUE (derived tables),
    never to a merged expression."""

    def _col(self, **kw):
        from types import SimpleNamespace as _NS
        d = {"id": kw.get("name", ""), "name": "", "alias": "", "expr": "",
             "col_refname": None, "where_clause": None}
        d.update(kw)
        return _NS(**d)

    def _eng(self):
        from rml_python.engine import RMLReportEngine as _E
        return object.__new__(_E)

    def test_refname_plan_and_wrap(self):
        eng = self._eng()
        a = self._col(name="net", alias="الصافي", expr='"T0"."AMOUNT" - "T0"."DISC"', col_refname="net_total")
        b = self._col(name="tax", alias="الضريبة", expr="@net_total * 0.15", col_refname="tax_v")
        c = self._col(name="gross", alias="الإجمالي", expr="@net_total + @tax_v")
        cols = [a, b, c]
        plan = eng._plan_value_refs(cols)
        self.assertIsNotNone(plan)
        self.assertEqual([[x.name for x in lvl] for lvl in plan["levels"]], [["tax"], ["gross"]])
        self.assertEqual([x.name for x in plan["base"]], ["net"])
        sql = eng._wrap_value_levels(
            'SELECT "T0"."AMOUNT" - "T0"."DISC" AS "الصافي"', '"sch"."t" "T0"', "", "",
            " ORDER BY 2 DESC", "", ["الصافي", "الضريبة", "الإجمالي"],
            plan["levels"], plan["resolved"])
        self.assertIn('"الصافي"', sql)
        self.assertIn("ORDER BY 2 DESC", sql)
        self.assertNotIn("@net_total", sql)
        # unknown @tokens → no plan (old behavior preserved)
        d = self._col(name="x", alias="X", expr="@ghost + 1")
        self.assertIsNone(eng._plan_value_refs([d]))

    def test_refname_cycles_and_python(self):
        eng = self._eng()
        a = self._col(name="a", alias="A", expr="@b_r", col_refname="a_r")
        b = self._col(name="b", alias="B", expr="@a_r", col_refname="b_r")
        with self.assertRaises(ValueError):
            eng._plan_value_refs([a, b])
        s = self._col(name="s", alias="S", expr="@py_r + 1")
        p = self._col(name="p", alias="P", expr="XLOOKUP([a], [b])", col_refname="py_r")
        with self.assertRaises(ValueError):
            eng._plan_value_refs([s, p])
        me = self._col(name="m", alias="M", expr="@m_r + 1", col_refname="m_r")
        with self.assertRaises(ValueError):
            eng._plan_value_refs([me])

    def test_refname_with_get_and_bracket(self):
        eng = self._eng()
        p = self._col(name="salary", alias="الراتب", expr="=get(employees.basic_salary)", col_refname="base_sal")
        s = self._col(name="bonus", alias="البونص", expr="[base_sal] * 0.1")
        plan = eng._plan_value_refs([p, s])
        self.assertIsNotNone(plan)
        self.assertEqual([x.name for x in plan["base"]], ["salary"])
        self.assertEqual([[x.name for x in lvl] for lvl in plan["levels"]], [["bonus"]])

    def test_compiler_parses_refname(self):
        from rml_python.compiler import RMLColumn
        import dataclasses
        c = RMLColumn(id="1", name="n", alias="a")
        self.assertIsNone(c.col_refname)
        c2 = dataclasses.replace(c, col_refname="net_total")
        self.assertEqual(c2.col_refname, "net_total")
        self.assertEqual(c2.to_dict()["colRefname"], "net_total")


class RmlPaletteTests(SimpleTestCase):
    def test_xml_roundtrip_visible_width_refname(self):
        import tempfile as _tf
        from pathlib import Path as _P
        xml = ('<rml><rpt_metadata name="t" displayName="T" category="X" schema="S"/>'
               '<columns>'
               '<column id="1" name="a" alias="A" expr="a" col_refname="a_v"/>'
               '<column id="2" name="b" alias="B" expr="b" visible="0" width="220"/>'
               '<column id="3" name="c" alias="C" expr="c"/>'
               '</columns></rml>')
        with _tf.TemporaryDirectory() as td:
            p = _P(td) / "t.rml"
            p.write_text(xml, encoding="utf-8")
            from rml_python.compiler import RMLReportCompiler
            cols = {c.name: c for c in RMLReportCompiler(path=p).columns()}
        self.assertEqual(cols["a"].col_refname, "a_v")
        self.assertFalse(cols["b"].visible)
        self.assertEqual(cols["b"].width, 220)
        self.assertTrue(cols["c"].visible)
        self.assertIsNone(cols["c"].width)

    def test_render_xml_persists_new_attrs(self):
        cols = [{"id": "1", "name": "a", "alias": "A", "expr": "a", "col_refname": "a_v"},
                {"id": "2", "name": "b", "alias": "B", "expr": "b", "visible": False, "width": 220},
                {"id": "3", "name": "c", "alias": "C", "expr": "c"}]
        out = _v._render_rml_xml("t", "T", "fa-x", "X", "S", "", None, [], [], cols,
                                 [], [], "master", None, [], "", [], False, [], 1, "", {})
        self.assertIn('col_refname="a_v"', out)
        self.assertIn('visible="0"', out)
        self.assertIn('width="220"', out)
        self.assertNotIn('visible="0" verbo', out)

    def test_empty_column_and_render_xml_no_db_column_fallback(self):
        from rml_python.compiler import RMLColumn, RMLField
        from rml_python.engine import _build_select
        col = RMLColumn(id="25", name="col_25", alias="عمود 25", expr="")
        fld = RMLField(id="1", name="salary", table_source="emp", connection_id="1")
        sel = _build_select([col], fields=[fld])
        self.assertIn('NULL AS "عمود 25"', sel)
        self.assertNotIn('"col_25"', sel)

        cols = [{"id": "25", "name": "col_25", "alias": "عمود 25", "expr": ""}]
        out = _v._render_rml_xml("t", "T", "fa-x", "X", "S", "", None, [], [], cols,
                                 [], [], "master", None, [], "", [], False, [], 1, "", {})
        self.assertIn('name="col_25"', out)
        self.assertIn('expr=""', out)

    def test_design_preview_validates(self):
        req = RequestFactory().post("/api/apps/settings/rml/design-preview/",
                                    data=json.dumps({}), content_type="application/json")
        res = _v.api_rml_design_preview(req, "settings")
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 400)
        self.assertIn("اتصال", body.get("error", ""))

    def test_wizard_palette_markup(self):
        with open("rml_python/shared/rml_wizard_script.html", encoding="utf-8") as f:
            src = f.read()
        for token in ("rpt-flex", "rptSideColsHTML", "rptPreviewFetch", "design-preview",
                      "rptDrag", "rptDrop", "rptPropsRender", "col_refname", "rptDetListHTML",
                      "قيمة ✓"):
            self.assertIn(token, src)

    def test_player_width_visibility(self):
        import os as _os
        path = _os.path.join(_os.path.dirname(__file__), "templates", "report_player.html")
        with open(path, encoding="utf-8") as f:
            src = f.read()
        for token in ("colWidthInput", "setColWidth", "colWidthStyle", "colWidths",
                      "rml_widths_", "visible === false"):
            self.assertIn(token, src)


class SuperuserModalMarkupTests(SimpleTestCase):
    def test_home_modal_supports_dynamic_extra_fields(self):
        import os as _os
        path = _os.path.join(_os.path.dirname(__file__), "templates", "home.html")
        with open(path, encoding="utf-8") as f:
            src = f.read()
        self.assertIn("data-ws-su-extra", src)
        self.assertIn("extra", src)


class JsonSourceStoreTests(SimpleTestCase):
    """Pure json_source tests (temp dir as base — no repo writes, no DB)."""

    def setUp(self):
        import tempfile as _tf
        self._td = _tf.TemporaryDirectory()
        self.addCleanup(self._td.cleanup)
        self.base = self._td.name

    def _write(self, name, doc):
        import os as _os
        import json as _js
        p = _os.path.join(self.base, name)
        with open(p, "w", encoding="utf-8") as f:
            _js.dump(doc, f, ensure_ascii=False)
        return name

    def test_shapes_and_tables(self):
        from urs import json_source as _j
        self._write("a.json", {"users": [{"id": 1}], "orders": []})
        doc, mode = _j.load("a.json", self.base)
        self.assertEqual(mode, "file")
        self.assertEqual(_j.tables(doc), ["orders", "users"])
        self._write("b.json", [{"id": 1}])
        doc2, _ = _j.load("b.json", self.base)
        self.assertEqual(_j.tables(doc2), ["data"])
        with self.assertRaises(ValueError):
            _j.tables({"a": 1})
        with self.assertRaises(ValueError):
            _j.load("missing.json", self.base)

    def test_source_guards(self):
        from urs import json_source as _j
        self.assertTrue(_j.is_json_source("https://x/y.json"))
        self.assertTrue(_j.is_json_source("data/f.json"))
        self.assertFalse(_j.is_json_source("data/f.csv"))
        self.assertFalse(_j.is_json_source(""))
        with self.assertRaises(ValueError):
            _j.resolve_path("../evil.json", self.base)
        with self.assertRaises(ValueError):
            _j.resolve_path("data.csv", self.base)

    def test_columns_inference(self):
        from urs import json_source as _j
        self._write("c.json", {"t": [
            {"id": 1, "price": 2.5, "ok": True, "day": "2026-01-05", "name": "x"},
            {"id": 2, "price": 3, "ok": False, "day": "2026-01-06", "name": "y", "extra": None},
        ]})
        doc, _ = _j.load("c.json", self.base)
        cols = {c["name"]: c["type"] for c in _j.columns(doc, "t")}
        self.assertEqual(cols["id"], "INTEGER")
        self.assertEqual(cols["price"], "NUMERIC")
        self.assertEqual(cols["ok"], "BOOLEAN")
        self.assertEqual(cols["day"], "DATE")
        self.assertEqual(cols["name"], "TEXT")
        self.assertEqual(cols["extra"], "TEXT")

    def test_preview_filters(self):
        from urs import json_source as _j
        self._write("p.json", {"t": [
            {"id": 1, "name": "ahmad"}, {"id": 2, "name": "Sara"}, {"id": 3, "name": "ahmad x"}]})
        doc, _ = _j.load("p.json", self.base)
        cols, rows = _j.preview(doc, "t", 50, q="AHMAD")
        self.assertEqual(len(rows), 2)
        cols, rows = _j.preview(doc, "t", 50, filters=[{"col": "id", "op": "gt", "val": "1"}])
        self.assertEqual([r["id"] for r in rows], [2, 3])
        cols, rows = _j.preview(doc, "t", 1)
        self.assertEqual(len(rows), 1)

    def test_row_write_roundtrip(self):
        from urs import json_source as _j
        self._write("w.json", {"t": [{"id": 1, "name": "a"}]})
        doc, _ = _j.load("w.json", self.base)
        _j.insert(doc, "t", {"id": 2, "name": "b"})
        self.assertEqual(_j.update(doc, "t", {"id": 1}, {"name": "a2"}), 1)
        self.assertEqual(_j.update(doc, "t", {"id": 9}, {"name": "z"}), 0)
        self.assertEqual(_j.delete(doc, "t", {"id": 2}), 1)
        with self.assertRaises(ValueError):
            _j.update(doc, "t", {}, {"name": "z"})
        _j.save("w.json", doc, self.base)
        doc2, _ = _j.load("w.json", self.base)
        self.assertEqual(doc2, {"t": [{"id": 1, "name": "a2"}]})


class JsonConnectionEndpointTests(SimpleTestCase):
    """JSON connection endpoints with a mocked connection row (no DB)."""

    def setUp(self):
        import tempfile as _tf
        import json as _js
        import os as _os
        self.rf = RequestFactory()
        self._td = _tf.TemporaryDirectory()
        self.addCleanup(self._td.cleanup)
        self.doc = {"users": [{"id": 1, "name": "ahmad", "role": "admin"},
                              {"id": 2, "name": "sara", "role": "user"}]}
        with open(_os.path.join(self._td.name, "u.json"), "w", encoding="utf-8") as f:
            _js.dump(self.doc, f, ensure_ascii=False)
        self.fake = SimpleNamespace(id=99, name="json1", host="u.json", port=0,
                                    user="", password="", instance="", engine="json",
                                    conn_type="json", is_local=True, is_queryable=True,
                                    schema="", endpoint="att")
        self._eff = mock.patch.object(_v, "_effective_or_row", return_value=self.fake)
        self._eff.start()
        self.addCleanup(self._eff.stop)
        self._base = mock.patch.object(_v, "BASE_DIR", self._td.name)
        self._base.start()
        self.addCleanup(self._base.stop)

    def test_tables_columns_preview(self):
        req = self.rf.get("/api/connections/99/tables/")
        res = _v.api_connection_tables(req, 99)
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200)
        self.assertEqual(body["tables"], [{"name": "users", "schema": "", "full": "users"}])
        req = self.rf.get("/api/connections/99/tables/users/columns/")
        res = _v.api_connection_table_columns(req, 99, "users")
        body = json.loads(res.content.decode())
        names = {c["name"] for c in body["columns"]}
        self.assertEqual(names, {"id", "name", "role"})
        req = self.rf.get("/api/connections/99/tables/users/preview/", {"limit": "50", "q": "sara"})
        res = _v.api_connection_table_preview(req, 99, "users")
        body = json.loads(res.content.decode())
        self.assertEqual(body["columns"], ["id", "name", "role"])
        self.assertEqual(body["rows"], [[2, "sara", "user"]])

    def test_row_insert_update_delete(self):
        req = self.rf.post("/api/connections/99/tables/users/rows/insert/",
                           data=json.dumps({"row": {"id": 3, "name": "omar"}}),
                           content_type="application/json")
        res = _v.api_json_row_insert(req, 99, "users")
        self.assertEqual(res.status_code, 200)
        req = self.rf.post("/api/connections/99/tables/users/rows/update/",
                           data=json.dumps({"match": {"id": 3}, "patch": {"role": "user"}}),
                           content_type="application/json")
        res = _v.api_json_row_update(req, 99, "users")
        self.assertEqual(json.loads(res.content.decode())["updated"], 1)
        req = self.rf.post("/api/connections/99/tables/users/rows/delete/",
                           data=json.dumps({"match": {"id": 3}}),
                           content_type="application/json")
        res = _v.api_json_row_delete(req, 99, "users")
        self.assertEqual(json.loads(res.content.decode())["deleted"], 1)
        # persisted back to the file
        from urs import json_source as _j
        doc, _ = _j.load("u.json", self._td.name)
        self.assertEqual(len(doc["users"]), 2)

    def test_rejects_non_json_connection(self):
        self.fake.engine = "postgres"
        req = self.rf.post("/api/connections/99/tables/users/rows/insert/",
                           data=json.dumps({"row": {"id": 1}}),
                           content_type="application/json")
        res = _v.api_json_row_insert(req, 99, "users")
        self.assertEqual(res.status_code, 400)

    def test_test_connection_ok_and_missing(self):
        ok, payload, status = _v._test_connection_obj(self.fake)
        self.assertTrue(ok)
        self.assertEqual(status, 200)
        self.assertIn("2", payload["note"])
        bad = SimpleNamespace(**{**vars(self.fake), "host": "nope.json"})
        ok, payload, status = _v._test_connection_obj(bad)
        self.assertFalse(ok)
        self.assertEqual(status, 400)

    def test_player_and_wizard_know_json(self):
        import os as _os
        base = _os.path.dirname(__file__)
        with open(_os.path.join(base, "templates", "forms_player.html"), encoding="utf-8") as f:
            self.assertIn("CONN_JSON_ENGINES", f.read())
        with open("rml_python/shared/rml_wizard_script.html", encoding="utf-8") as f:
            self.assertIn("json", f.read())


class UserProfileAndHeaderTests(SimpleTestCase):
    def setUp(self):
        self.rf = RequestFactory()

    def test_change_password_endpoint(self):
        from fmlk_engine.engine import hash_secret
        st = _fresh_state(
            fullcols=[("username", "text", "NO", None),
                      ("password", "text", "NO", None),
                      ("full_name", "text", "YES", None)],
            userrow=("admin", hash_secret("secret123"), "Admin Full Name")
        )
        req = self.rf.post("/api/user/change-password/",
                           data=json.dumps({"current_password": "secret123",
                                            "new_password": "newpassword123",
                                            "confirm_password": "newpassword123"}),
                           content_type="application/json")
        req.session = {
            "ws_user": {"username": "admin", "full_name": "Admin Full Name"},
            "workspace": "workspace_1",
            "fiscal_schema": "main_hq_2026",
        }
        with mock.patch.multiple(_v,
                                 _ws_login_connection=mock.Mock(return_value=({"fake": True}, "sys_users", "")),
                                 _ws_pg_connect=mock.Mock(side_effect=lambda _obj: _FakeConn(st))):
            res = _v.api_user_change_password(req)
            self.assertEqual(res.status_code, 200)
            data = json.loads(res.content)
            self.assertTrue(data.get("ok"))
            self.assertIsNotNone(st.get("update"))
            sql, params = st["update"]
            self.assertIn("UPDATE", sql)
            self.assertNotEqual(params[0], "newpassword123")
            self.assertTrue(params[0].startswith("pbkdf2_"))

    def test_profile_update_endpoint(self):
        st = _fresh_state(
            fullcols=[("username", "text", "NO", None),
                      ("full_name", "text", "YES", None),
                      ("email", "text", "YES", None)],
            userrow=("admin", "Admin Old", "admin@old.com")
        )
        req = self.rf.post("/api/user/profile/",
                           data=json.dumps({"full_name": "Ahmed Mohamed", "email": "ahmed@example.com"}),
                           content_type="application/json")
        req.session = {
            "ws_user": {"username": "admin", "full_name": "Admin Old"},
            "workspace": "workspace_1",
            "fiscal_schema": "main_hq_2026",
        }
        with mock.patch.multiple(_v,
                                 _ws_login_connection=mock.Mock(return_value=({"fake": True}, "sys_users", "")),
                                 _ws_pg_connect=mock.Mock(side_effect=lambda _obj: _FakeConn(st))):
            res = _v.api_user_profile(req)
            self.assertEqual(res.status_code, 200)
            data = json.loads(res.content)
            self.assertEqual(data.get("full_name"), "Ahmed Mohamed")
            self.assertEqual(req.session["ws_user"]["full_name"], "Ahmed Mohamed")


class OracleClientUtilTests(SimpleTestCase):
    def test_oracle_error_hints(self):
        from urs.oracle_util import oracle_error_hint
        h3015 = oracle_error_hint("DPY-3015: password verifier type 0x939 is not supported", user="rpt")
        self.assertIn("0x939", h3015)
        self.assertIn("rpt", h3015)

        h1017 = oracle_error_hint("ORA-01017: invalid username/password", user="rpt", password="pbkdf2_sha256$600000$test")
        self.assertIn("ORA-01017", h1017)
        self.assertIn("hash", h1017)

    def test_fmlk_engine_does_not_hash_connection_passwords(self):
        from fmlk_engine.engine import FMLKFormEngine
        eng = mock.Mock(spec=FMLKFormEngine)
        eng.table = "urs_connection"
        out = FMLKFormEngine._hash_secrets(eng, {"password": "plain_password_123"}, {"password"})
        self.assertEqual(out.get("password"), "plain_password_123")

    def test_connections_fmlk_has_custom_actions(self):
        from pathlib import Path
        from fmlk_engine.compiler import FMLKFormCompiler
        c = FMLKFormCompiler(Path("workspace_2/apps/settings/connections.fmlk"))
        action_names = [a.name for a in c.actions()]
        self.assertIn("test_connection", action_names)
        self.assertIn("test_connection_form", action_names)
        tc = next(a for a in c.actions() if a.name == "test_connection")
        self.assertEqual(tc.label, "اختبار الاتصال")
        self.assertEqual(tc.endpoint, "/api/connections/test/")
        self.assertEqual(tc.level, "record")

    def test_resolve_connection_record_prefers_name_over_conflicting_id(self):
        from urs.views import _resolve_connection_record
        from urs.models import Connection
        with mock.patch("urs.models.Connection.objects.filter") as m_filter:
            m_target = mock.Mock(spec=Connection)
            m_target.id = 16
            m_target.name = "اونكس تعافي"
            m_filter.return_value.first.return_value = m_target
            obj = _resolve_connection_record({"id": 2, "name": "اونكس تعافي"})
            self.assertEqual(obj.name, "اونكس تعافي")
            m_filter.assert_called_with(name="اونكس تعافي")

    def test_resolve_connection_record_with_arabic_aliases(self):
        from urs.views import _resolve_connection_record
        from urs.models import Connection
        with mock.patch("urs.models.Connection.objects.filter") as m_filter:
            m_target = mock.Mock(spec=Connection)
            m_target.name = "اونكس تعافي"
            m_filter.return_value.first.return_value = m_target
            obj = _resolve_connection_record({
                "اسم الاتصال": "اونكس تعافي",
                "نوع المحرك": "oracle",
                "المضيف / IP الجهاز": "172.16.10.100",
                "__pk_id": 2,
            })
            self.assertEqual(obj.name, "اونكس تعافي")
            m_filter.assert_called_with(name="اونكس تعافي")

    def test_password_encryption_algorithms_and_shift(self):
        from fmlk_engine.crypto import list_password_algorithms, encrypt_or_hash_password, verify_password_algorithm
        from fmlk_engine.field_types import get_type, CONFIG_ATTRS
        # 1. Config attributes & registry
        self.assertIn("hash_algo", CONFIG_ATTRS["password"])
        self.assertIn("shift", CONFIG_ATTRS["password"])
        pw_t = get_type("password")
        self.assertIn("hash_algo", pw_t.get("attrs", []))
        self.assertIn("shift", pw_t.get("attrs", []))

        # 2. 20 algorithms count & order
        algos = list_password_algorithms()
        self.assertEqual(len(algos), 20)
        self.assertEqual(algos[0]["key"], "caesar")
        self.assertEqual(algos[-1]["key"], "scrypt")

        # 3. Plaintext ("none" / "لا شيء")
        pw = "MyPass123_@#!"
        self.assertEqual(encrypt_or_hash_password(pw, "none"), pw)
        self.assertTrue(verify_password_algorithm(pw, pw, "none"))

        # 4. Caesar with positive & negative unicode shift
        enc_pos = encrypt_or_hash_password(pw, "caesar", shift=4)
        self.assertTrue(enc_pos.startswith("caesar$4$"))
        self.assertTrue(verify_password_algorithm(pw, enc_pos))
        self.assertFalse(verify_password_algorithm("wrong_password", enc_pos))

        enc_neg = encrypt_or_hash_password(pw, "caesar", shift=-5)
        self.assertTrue(enc_neg.startswith("caesar$-5$"))
        self.assertTrue(verify_password_algorithm(pw, enc_neg))

        # 5. All 20 algorithms encrypt and verify properly
        for a in algos:
            k = a["key"]
            enc = encrypt_or_hash_password(pw, k)
            self.assertTrue(verify_password_algorithm(pw, enc, k), f"Failed for algo {k}")

        # 6. Engine integration
        from fmlk_engine.engine import hash_secret, verify_secret
        eng_enc = hash_secret(pw, algorithm="caesar", shift=7)
        self.assertTrue(eng_enc.startswith("caesar$7$"))
        self.assertTrue(verify_secret(pw, eng_enc))

    def test_custom_actions_roundtrip_and_compiler(self):
        from fmlk_engine.compiler import FMLKFormCompiler
        xml = """<fml>
          <fml_metadata name="test" displayName="Test" table="t"/>
          <fields><field name="id" dataType="INTEGER" primary_key="true"/></fields>
          <custom_actions>
            <action name="test_connection" label="اختبار الاتصال" endpoint="/api/connections/test/" icon="fa-plug-circle-check" badge_color="#059669" level="record"/>
            <action name="conn_stats" label="إحصائيات الاتصال" endpoint="/api/connections/stats/" render="modals/conn_stats.html" icon="fa-chart-simple" badge_color="#4f46e5" level="record"/>
            <action name="export_data" label="تصدير" endpoint="/api/fmlk/export/" icon="fa-file-excel" badge_color="#16a34a" level="view"/>
          </custom_actions>
        </fml>"""
        comp = FMLKFormCompiler.from_string(xml)
        actions = comp.actions()
        self.assertEqual(len(actions), 3)
        self.assertEqual(actions[0].name, "test_connection")
        self.assertEqual(actions[0].label, "اختبار الاتصال")
        self.assertEqual(actions[0].badge_color, "#059669")
        self.assertEqual(actions[0].level, "record")
        self.assertEqual(actions[1].render, "modals/conn_stats.html")
        self.assertEqual(actions[2].level, "view")

        # Test dictionary export format
        dicts = [a.to_dict() for a in actions]
        self.assertEqual(dicts[0]["endpoint"], "/api/connections/test/")
        self.assertEqual(dicts[1]["render"], "modals/conn_stats.html")
        self.assertEqual(dicts[2]["badge_color"], "#16a34a")

    def test_resolve_connection_schema_isolates_external_connections(self):
        from unittest.mock import patch, MagicMock
        from urs.views import _resolve_connection_schema
        from django.test import RequestFactory
        rf = RequestFactory()

        class DummyConn:
            def __init__(self, id, name, schema):
                self.id = id
                self.name = name
                self.schema = schema

        primary = DummyConn(1, "rex_local", "default_pg")
        external = DummyConn(16, "اونكس تعافي", "IAS20264")

        # Mock workspace context: primary_conn="rex_local", fiscal_schema="ies202601"
        with patch("urs.views._ws_session_context", return_value=("ws1", "ies202601", "rex_local")):
            # 1. External connection MUST preserve its own schema and NOT get corrupted by ies202601
            req = rf.get("/api/connections/16/tables/")
            sch = _resolve_connection_schema(req, external)
            self.assertEqual(sch, "IAS20264")

            # 2. Primary connection receives the workspace fiscal_schema
            sch_p = _resolve_connection_schema(req, primary)
            self.assertEqual(sch_p, "ies202601")

            # 3. Explicit URL param ?schema= overrides
            req_custom = rf.get("/api/connections/16/tables/?schema=CUSTOM_SCH")
            sch_c = _resolve_connection_schema(req_custom, external)
            self.assertEqual(sch_c, "CUSTOM_SCH")

    def test_api_connection_table_values_validation(self):
        from unittest.mock import patch, MagicMock
        from urs.views import api_connection_table_values
        from django.test import RequestFactory
        import json

        rf = RequestFactory()

        # 1. Connection not found -> 404
        with patch("urs.views._effective_or_row", return_value=None):
            req = rf.get("/api/connections/999/tables/ACCOUNT/values/?column=A_CODE")
            resp = api_connection_table_values(req, 999, "ACCOUNT")
            self.assertEqual(resp.status_code, 404)

        # 2. Missing or invalid column name -> 400
        dummy = MagicMock()
        dummy.id = 1
        dummy.engine = "postgres"
        with patch("urs.views._effective_or_row", return_value=dummy):
            # Missing column
            req_no_col = rf.get("/api/connections/1/tables/ACCOUNT/values/")
            resp_no_col = api_connection_table_values(req_no_col, 1, "ACCOUNT")
            self.assertEqual(resp_no_col.status_code, 400)

            # Invalid column (SQL injection pattern)
            req_bad_col = rf.get("/api/connections/1/tables/ACCOUNT/values/?column=bad;drop")
            resp_bad_col = api_connection_table_values(req_bad_col, 1, "ACCOUNT")
            self.assertEqual(resp_bad_col.status_code, 400)

        # 3. Successful values fetch using mock _table_values_obj
        with patch("urs.views._effective_or_row", return_value=dummy), \
             patch("urs.views._resolve_connection_schema", return_value="public"), \
             patch("urs.views._table_values_obj", return_value=["val1", "val2", "val3"]):
            req_ok = rf.get("/api/connections/1/tables/ACCOUNT/values/?column=A_CODE&limit=10")
            resp_ok = api_connection_table_values(req_ok, 1, "ACCOUNT")
            self.assertEqual(resp_ok.status_code, 200)
            data = json.loads(resp_ok.content.decode("utf-8"))
            self.assertEqual(data["values"], ["val1", "val2", "val3"])
            self.assertEqual(data["total"], 3)
            self.assertEqual(data["column"], "A_CODE")

    def test_column_currency_symbol_persistence_and_parsing(self):
        from rml_python.compiler import RMLReportCompiler, Column
        import xml.etree.ElementTree as ET
        from urs.views import _write_detail_el

        # 1. Column dataclass and to_dict
        col = Column(id="1", name="salary", alias="الراتب", is_amount=True, currency_symbol="ر.س")
        d = col.to_dict()
        self.assertEqual(d["currency_symbol"], "ر.س")
        self.assertEqual(d["currencySymbol"], "ر.س")
        self.assertTrue(d["is_amount"])

        # 2. RML parsing from XML
        xml = """<report>
            <metadata name="test_currency"/>
            <columns>
                <column id="1" name="amount" alias="المبلغ" is_amount="1" currency_symbol="USD"/>
                <column id="2" name="tax" alias="الضريبة" is_amount="1" currency_symbol="ر.ي"/>
            </columns>
            <detail table="invoices" master="id" detail="customer_id">
                <column id="1" name="item_total" alias="إجمالي البند" is_amount="1" currency_symbol="SAR"/>
            </detail>
        </report>"""
        compiler = RMLReportCompiler.from_string(xml)
        cols = compiler.columns()
        self.assertEqual(len(cols), 2)
        self.assertEqual(cols[0].currency_symbol, "USD")
        self.assertTrue(cols[0].is_amount)
        self.assertEqual(cols[1].currency_symbol, "ر.ي")

        det = compiler.detail()
        self.assertIsNotNone(det)
        self.assertEqual(len(det.columns), 1)
        self.assertEqual(det.columns[0].currency_symbol, "SAR")
        self.assertTrue(det.columns[0].is_amount)

        # 3. Serialization in _write_detail_el
        root = ET.Element("report")
        detail_data = {
            "table": "inv_items",
            "master": "inv_id",
            "detail": "id",
            "columns": [
                {"id": "1", "name": "price", "alias": "السعر", "is_amount": True, "currency_symbol": "EUR"}
            ]
        }
        _write_detail_el(root, ET, detail_data)
        det_el = root.find("detail")
        self.assertIsNotNone(det_el)
        col_el = det_el.find("column")
        self.assertIsNotNone(col_el)
        self.assertEqual(col_el.get("currency_symbol"), "EUR")
        self.assertEqual(col_el.get("is_amount"), "1")


class PolicyExcelImportTests(SimpleTestCase):
    """Policy match-values: per-source Excel template + import merge."""

    def test_excel_template_download(self):
        from openpyxl import load_workbook
        import io as _io
        rf = RequestFactory()
        req = rf.get("/api/search/excel-template/", {"column": "BRANCH_CODE"})
        res = _v.api_search_excel_template(req)
        self.assertEqual(res.status_code, 200)
        self.assertIn("spreadsheetml", res["Content-Type"])
        self.assertIn("attachment", res["Content-Disposition"])
        wb = load_workbook(filename=_io.BytesIO(res.content), read_only=True, data_only=True)
        try:
            ws = wb.active
            self.assertEqual(ws["A1"].value, "BRANCH_CODE")
        finally:
            try:
                wb.close()
            except Exception:
                pass

    def test_excel_template_arabic_column(self):
        rf = RequestFactory()
        req = rf.get("/api/search/excel-template/", {"column": "رقم الحوالة"})
        res = _v.api_search_excel_template(req)
        self.assertEqual(res.status_code, 200)
        self.assertIn("filename*=UTF-8''", res["Content-Disposition"])

    def test_excel_template_requires_column(self):
        rf = RequestFactory()
        res = _v.api_search_excel_template(rf.get("/api/search/excel-template/"))
        self.assertEqual(res.status_code, 400)

    def test_policy_modal_excel_markup(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "templates", "app_detail.html"), encoding="utf-8") as f:
            html = f.read()
        for needle in ("rpDownloadTemplate", "rpImportExcel", "excel-template",
                       "import-excel", "rpCnt-", "rpXl-", "قالب", "Excel"):
            self.assertIn(needle, html)


class TableFilterRowsTests(SimpleTestCase):
    """Field-conditions modal: rows + ANY/ALL match persist across reload."""

    def test_write_table_opts_roundtrip(self):
        import xml.etree.ElementTree as ET
        rml = ET.Element("rml")
        _v._write_table_opts_el(rml, ET, [{
            "name": "IAS_BILL_MST", "conn": "2",
            "where": "[IAS_BILL_MST.CC_CODE] = '11' OR [IAS_BILL_MST.CC_CODE] = '12'",
            "where_rows": [
                {"col": "CC_CODE", "op": "=", "val": "11", "val2": ""},
                {"col": "CC_CODE", "op": "=", "val": "12", "val2": ""},
                {"op": "=", "val": "x"},
                {"col": "CC_CODE", "op": "hack'); DROP--", "val": "1"},
                "junk",
            ],
            "where_match": "any",
        }])
        t = rml.find("table_opts/table")
        self.assertIsNotNone(t)
        self.assertEqual(t.get("where_match"), "any")
        import json as _json
        rows = _json.loads(t.get("where_rows"))
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0], {"col": "CC_CODE", "op": "=", "val": "11", "val2": ""})

    def test_write_table_opts_match_all_and_empty(self):
        import xml.etree.ElementTree as ET
        rml = ET.Element("rml")
        _v._write_table_opts_el(rml, ET, [{"name": "T", "where": "x", "where_match": "bogus"}])
        t = rml.find("table_opts/table")
        self.assertIsNone(t.get("where_match"))
        rml2 = ET.Element("rml")
        _v._write_table_opts_el(rml2, ET, [{"name": "T", "where_match": "AND"}])
        self.assertEqual(rml2.find("table_opts/table").get("where_match"), "all")

    def test_compiler_table_opts_roundtrip(self):
        from rml_python.compiler import RMLReportCompiler
        comp = RMLReportCompiler(xml_text=(
            '<rml><table_opts>'
            '<table name="T" conn="2" where="x" '
            'where_rows=\'[{"col": "C", "op": "between", "val": "1", "val2": "9"}]\' '
            'where_match="OR"/>'
            '</table_opts></rml>'))
        opts = comp.table_opts()
        self.assertEqual(len(opts), 1)
        self.assertEqual(opts[0]["where_rows"],
                         [{"col": "C", "op": "between", "val": "1", "val2": "9"}])
        self.assertEqual(opts[0]["where_match"], "any")

    def test_designer_table_filter_markup(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "..", "rml_python", "shared",
                                "rml_wizard_script.html"), encoding="utf-8") as f:
            js = f.read()
        for needle in ("importFreeWhereState", "tfModalFreeWhere",
                        "onImportFreeWhereInput", "onTfModalFreeWhereInput",
                        "renderImportFreeWhere", "renderTfModalFreeWhere",
                        "generateImportTableWhere", "generateTableFilterModalWhere",
                        "tcond:import", "tcond:modal",
                        "wbTableCondPool", "wbTableCondScope", "wbTableCondUnscope",
                        "wizardTableFilterRows", "wizardTableFilterMatchMode",
                        "where_rows", "where_match"):
            self.assertIn(needle, js)
        with open(os.path.join(base, "..", "rml_python", "shared",
                                "rml_wizard_modal.html"), encoding="utf-8") as f:
            modal = f.read()
        self.assertIn('id="importFreeWhereWrap"', modal)
        self.assertIn('id="tfModalFreeWrap"', modal)
        self.assertIn('id="importFilterPreview"', modal)
        self.assertIn('id="tfModalPreview"', modal)
        with open(os.path.join(base, "..", "rml_python", "shared",
                                "UniCondition.js"), encoding="utf-8") as f:
            uc = f.read()
        for needle in ("UniCondition", "fieldHTML", "chipRow", "insertAt",
                        "colChips", "setValue", "CURRENT_DATE"):
            self.assertIn(needle, uc)

    def test_filter_current_date_markup(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "..", "rml_python", "shared",
                                "UniCondition.js"), encoding="utf-8") as f:
            uc = f.read()
        self.assertIn("CURRENT_DATE", uc)
        with open(os.path.join(base, "..", "rml_python", "shared",
                                "rml_wizard_script.html"), encoding="utf-8") as f:
            js = f.read()
        self.assertIn("CURRENT_DATE", js)

    def test_unicondition_included_before_wizard(self):
        base = os.path.dirname(__file__)
        for tpl in ("report_designer.html", "report_player.html", "app_detail.html"):
            with open(os.path.join(base, "templates", tpl), encoding="utf-8") as f:
                html = f.read()
            self.assertIn('{% include "UniCondition.js" %}', html)
            self.assertLess(html.index('{% include "UniCondition.js" %}'),
                            html.index('{% include "rml_wizard_script.html" %}'))


class GwPruneTests(SimpleTestCase):
    """Removing a table import drops its general-where conjunct (file + UI)."""

    def test_split_top_and(self):
        self.assertEqual(_v._split_top_and("([A.X] = 1) AND ([B.Y] = 2)"),
                         ["([A.X] = 1) ", " ([B.Y] = 2)"])
        self.assertEqual(_v._split_top_and("[N] = 'A AND B' AND [M] = 1"),
                         ["[N] = 'A AND B' ", " [M] = 1"])
        self.assertEqual(_v._split_top_and("([A.X] = 1 OR [A.Y] = 2)"),
                         ["([A.X] = 1 OR [A.Y] = 2)"])
        self.assertEqual(_v._split_top_and(""), [""])

    def test_prune_drops_removed_table_conjunct(self):
        kept, dropped = _v._prune_stale_general_where(
            "([T1.A] = 1) AND ([GONE.G] = 'x') AND ([T1.B] > 2)",
            {"T1"})
        self.assertEqual(kept, "([T1.A] = 1) AND ([T1.B] > 2)")
        self.assertEqual(dropped, ["([GONE.G] = 'x')"])

    def test_prune_keeps_bare_and_namespace_and_mixed(self):
        kept, dropped = _v._prune_stale_general_where(
            "[G_CODE] = 'x' AND ([T1.A] = [GONE.G]) AND ([NS1.M] = 1)",
            {"T1"}, namespaces={"ns1"})
        self.assertEqual(dropped, [])
        self.assertIn("[G_CODE] = 'x'", kept)
        self.assertIn("([T1.A] = [GONE.G])", kept)
        self.assertIn("([NS1.M] = 1)", kept)

    def test_prune_three_part_and_braces(self):
        kept, dropped = _v._prune_stale_general_where(
            "([2.GONE.G] = 1) AND ({T1.A} = 2)", {"T1"})
        self.assertEqual(kept, "({T1.A} = 2)")
        self.assertEqual(dropped, ["([2.GONE.G] = 1)"])

    def test_prune_empty(self):
        self.assertEqual(_v._prune_stale_general_where("", {"T1"}), ("", []))
        self.assertEqual(_v._prune_stale_general_where("   ", set()), ("", []))

    def test_normalize_prunes_and_records(self):
        data = {
            "fields": [{"name": "A", "table_source": "T1", "conn_id": "1"}],
            "columns": [{"id": "1", "name": "a", "alias": "a", "expr": "[T1.A]"}],
            "views": [{"name": "Sales"}],
            "general_where": "([T1.A] = 1) AND ([IAS_ITM_MST.G_CODE] = 'x')",
        }
        out, derr = _v._rml_normalize_payload(data)
        self.assertIsNone(derr)
        self.assertEqual(out["general_where"], "([T1.A] = 1)")
        self.assertEqual(out["_pruned_general_where"], ["([IAS_ITM_MST.G_CODE] = 'x')"])

    def test_drop_table_strips_gw_markup(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "..", "rml_python", "shared",
                                "rml_wizard_script.html"), encoding="utf-8") as f:
            js = f.read()
        for needle in ("wbStripGwConjuncts", "dropConds",
                       "أُزيل الجدول وشروطه العامة المرتبطة",
                       "j0.pruned", "أُسقطت بنود جداول مزالة"):
            self.assertIn(needle, js)

    def _gw_engine(self):
        from rml_python.engine import RMLReportEngine
        eng = RMLReportEngine.__new__(RMLReportEngine)
        eng.fields = [SimpleNamespace(name="PRICE_LVL", table_source="expire_items", conn_id="16"),
                      SimpleNamespace(name="EXP_DATE", table_source="expire_items", conn_id="16"),
                      SimpleNamespace(name="BRN_NO", table_source="S_BRN", conn_id="16")]
        eng.connections = []
        return eng

    def test_prune_gw_tables_drops_unused_only(self):
        eng = self._gw_engine()
        ft = eng._field_table_map()
        gw = ("([expire_items.PRICE_LVL] = '1') AND "
              "([16.S_BRN.BRN_NO] Like '10_2')")
        kept, dropped = eng._prune_gw_tables(gw, {"S_BRN"}, ft)
        self.assertIn("[expire_items.PRICE_LVL]", kept)
        self.assertNotIn("S_BRN", kept)
        self.assertEqual(dropped, ["([16.S_BRN.BRN_NO] Like '10_2')"])

    def test_prune_gw_tables_keeps_mixed_and_bare(self):
        eng = self._gw_engine()
        ft = eng._field_table_map()
        gw = ("[PRICE_LVL] = 'x' AND "
              "([expire_items.PRICE_LVL] = [S_BRN.BRN_NO]) AND "
              "([S_BRN.BRN_NO] = '2' OR [expire_items.PRICE_LVL] = '3')")
        kept, dropped = eng._prune_gw_tables(gw, {"S_BRN"}, ft)
        self.assertEqual(dropped, [])
        self.assertIn("[PRICE_LVL] = 'x'", kept)
        self.assertIn("([expire_items.PRICE_LVL] = [S_BRN.BRN_NO])", kept)
        self.assertIn("OR", kept)

    def test_prune_gw_tables_empty(self):
        eng = self._gw_engine()
        self.assertEqual(eng._prune_gw_tables("", {"S_BRN"}, {}), ("", []))
        self.assertEqual(eng._prune_gw_tables("[A] = 1", set(), {}), ("[A] = 1", []))


class GroupBysTests(SimpleTestCase):
    """Multi GROUP BY section + conditional CASE grouping + multi-table import."""

    def test_compiler_group_bys(self):
        from rml_python.compiler import RMLReportCompiler
        comp = RMLReportCompiler(xml_text=(
            '<rml><group_bys>'
            '<gb field="cat" order="1"/>'
            '<gb field="bucket" order="2" when_field="cat" op="=" value="1" '
            'then_field="sales" else_field="disc"/>'
            '<gb field="cat" order="3"/>'
            '</group_bys></rml>'))
        gbs = comp.group_bys()
        self.assertEqual(len(gbs), 2)
        self.assertEqual(gbs[0], {"field": "cat", "order": 1})
        self.assertEqual(gbs[1]["cond"]["when_field"], "cat")
        self.assertEqual(gbs[1]["cond"]["then_field"], "sales")

    def test_normalize_group_bys(self):
        from rml_python.engine import _normalize_group_bys
        self.assertEqual(_normalize_group_bys("a"), ["a"])
        self.assertEqual(_normalize_group_bys(["a", {"field": "b"}, "", "a", None]),
                         ["a", {"field": "b", "order": None}])
        self.assertEqual(_normalize_group_bys([]), [])
        kept = _normalize_group_bys([{"field": "x", "cond": {"when_field": "w"}}])
        self.assertEqual(len(kept), 1)
        self.assertIn("cond", kept[0])

    def test_extend_classic_group_by_helpers(self):
        from rml_python.engine import (_split_select_items, _strip_select_alias,
                                       _extend_classic_group_by, _count_over_groups)
        self.assertEqual(_split_select_items('a, SUM(b) AS "s", CASE WHEN x THEN 1 ELSE 0 END AS "c"'),
                          ['a', 'SUM(b) AS "s"', 'CASE WHEN x THEN 1 ELSE 0 END AS "c"'])
        # commas inside parens/quotes/IN-lists don't split
        self.assertEqual(len(_split_select_items("f(a, 'x,y'), g(CASE WHEN z IN (1,2) THEN 3 END)")), 2)
        self.assertEqual(_strip_select_alias('SUM(q) AS "مجموع"'), 'SUM(q)')
        self.assertEqual(_strip_select_alias('"CITY"'), '"CITY"')
        # bare CASE appended; aggregates and already-grouped keys untouched
        gc = _extend_classic_group_by(
            ' GROUP BY "CITY"',
            '"CITY" AS "c", SUM("AMT") AS "s", CASE WHEN "P" > 0 THEN "P" ELSE 0 END AS "p"')
        self.assertIn('"CITY"', gc)
        self.assertIn('CASE WHEN "P" > 0 THEN "P" ELSE 0 END', gc)
        self.assertNotIn('SUM("AMT")', gc)
        self.assertEqual(gc.count('"CITY"'), 1)
        # no keys + aggregates → GROUP BY auto-created (ORA-00937 proof)
        gc2 = _extend_classic_group_by(
            "", '"CITY" AS "c", SUM("AMT") AS "s"')
        self.assertTrue(gc2.strip().upper().startswith("GROUP BY"))
        self.assertIn('"CITY"', gc2)
        # no aggregates + no keys → untouched plain select
        self.assertEqual(_extend_classic_group_by("", '"A" AS "a", "B" AS "b"'), "")
        # grouped COUNT wraps (one row = number of groups)
        self.assertEqual(
            _count_over_groups("t", " WHERE x=1", ' GROUP BY "C"'),
            'SELECT COUNT(*) as cnt FROM (SELECT 1 FROM t WHERE x=1 GROUP BY "C") t')
        self.assertEqual(_count_over_groups("t", "", ""),
                          "SELECT COUNT(*) as cnt FROM t")

    _GB_CLASSIC_XML = """<rml><rpt_metadata name="t" displayName="t"/>
      <fields>
        <field id="1" name="CITY" type="VARCHAR" table_source="T"/>
        <field id="2" name="AMT" type="NUMERIC" table_source="T"/>
        <field id="3" name="PRC" type="NUMERIC" table_source="T"/>
      </fields>
      <columns>
        <column id="1" name="c" alias="المدينة" expr="=[CITY]"/>
        <column id="2" name="s" alias="المجموع" expr="=SUM([AMT])"/>
        <column id="3" name="p" alias="السعر" expr="=IF([PRC] &gt; 0, [PRC], 0)"/>
      </columns>
      <table_opts><table name="T" conn="1" is_default="1"/></table_opts>
      <group_bys><gb field="CITY" order="1"/></group_bys></rml>"""

    def _gb_classic_sql(self, group_by):
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        comp = RMLReportCompiler(xml_text=self._GB_CLASSIC_XML)
        eng = RMLReportEngine(comp, mock.Mock())
        plan = eng._plan_structure(None, [], None, group_by)
        sql, _ = eng._compile_sql2(plan, [], None, 1, 50, group_by)
        return sql

    def test_classic_sql_auto_groups_bare_case(self):
        # عارٍ + تجميعي + CASE عارٍ مع مفاتيح ناقصة → يُستكمل GROUP BY (ORA-00979)
        sql = self._gb_classic_sql([{"field": "CITY", "order": 1}])
        _up = sql.upper()
        self.assertIn("GROUP BY", _up)
        _gb = _up.split("GROUP BY", 1)[1]
        self.assertIn('"CITY"', _gb)
        self.assertIn("CASE WHEN", _gb)
        self.assertNotIn("SUM(", _gb)

    def test_classic_sql_auto_groups_without_keys(self):
        # بلا مفاتيح إطلاقاً مع مزج عارٍ/تجميعي → GROUP BY مُنشأة (ORA-00937)
        sql = self._gb_classic_sql(None)
        _up = sql.upper()
        self.assertIn("GROUP BY", _up)
        _gb = _up.split("GROUP BY", 1)[1]
        self.assertIn('"CITY"', _gb)
        self.assertIn("CASE WHEN", _gb)

    def test_at_alias_inlines_without_eq_prefix(self):
        # @اسم_عربي لعمود مجاور: التضمين بلا بادئة "=" المخزنة — وإلا (=... → ORA-00936
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        comp = RMLReportCompiler(xml_text="""<rml><rpt_metadata name="t" displayName="t"/>
          <fields>
            <field id="1" name="CITY" type="VARCHAR" table_source="T"/>
            <field id="2" name="AMT" type="NUMERIC" table_source="T"/>
            <field id="3" name="PRC" type="NUMERIC" table_source="T"/>
          </fields>
          <columns>
            <column id="1" name="c" alias="المدينة" expr="=[CITY]"/>
            <column id="2" name="s" alias="المجموع" expr="=SUM([AMT])"/>
            <column id="3" name="p" alias="التسعير" expr="=IF([PRC] &gt; 0, [PRC], 0)"/>
            <column id="4" name="v" alias="القيمة" expr="=ROUND([s] * @التسعير, 2)"/>
          </columns>
          <table_opts><table name="T" conn="1" is_default="1"/></table_opts>
          <group_bys><gb field="CITY" order="1"/></group_bys></rml>""")
        eng = RMLReportEngine(comp, mock.Mock())
        plan = eng._plan_structure(None, [], None, [{"field": "CITY", "order": 1}])
        sql, _ = eng._compile_sql2(plan, [], None, 1, 50, [{"field": "CITY", "order": 1}])
        self.assertNotIn("(=CASE", sql)
        self.assertNotIn("(=", sql)
        self.assertIn('(CASE WHEN "PRC" > 0 THEN "PRC" ELSE 0 END)', sql)
        self.assertIn("GROUP BY", sql.upper())

    def test_execute_grouped_count_counts_groups(self):
        # العدّ المجمّع صف واحد = عدد المجموعات (لا عدّ أول مجموعة)
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        comp = RMLReportCompiler(xml_text=self._GB_CLASSIC_XML)
        eng = RMLReportEngine(comp, mock.Mock())
        seen = []

        class _Cur:
            def __init__(self, sql):
                self._sql = sql
                self.description = []
            def fetchone(self):
                return (7,)
            def fetchall(self):
                return []
            def close(self):
                pass

        def _fake_exec(self, db, sql, params):
            seen.append(sql)
            return _Cur(sql)

        with mock.patch.object(RMLReportEngine, "_exec_on", _fake_exec):
            res = eng.execute({"page": 1, "pageSize": 20})
        counts = [s for s in seen if "COUNT(*)" in s.upper()]
        self.assertEqual(len(counts), 1)
        self.assertIn("(SELECT 1 FROM", counts[0].upper())
        self.assertIn("GROUP BY", counts[0].upper())
        self.assertEqual(res["total"], 7)
        mains = [s for s in seen if "COUNT(*)" not in s.upper()]
        self.assertEqual(len(mains), 1)
        self.assertIn("GROUP BY", mains[0].upper())
        self.assertIn("CASE WHEN", mains[0].upper().split("GROUP BY", 1)[1])

    def test_build_group_clause_multi_and_case(self):
        from rml_python.engine import _build_group_clause
        self.assertEqual(_build_group_clause(["cat", "yr"], None, [], {}, {}, [], None),
                         ' GROUP BY "cat", "yr"')
        self.assertEqual(_build_group_clause("cat", None, [], {}, {}, [], None),
                         ' GROUP BY "cat"')
        self.assertEqual(_build_group_clause([], None, [], {}, {}, [], None), "")
        cond = [{"field": "bucket", "cond": {"when_field": "cat", "op": "=",
                                             "value": "1", "then_field": "sales",
                                             "else_field": "disc"}}]
        self.assertEqual(_build_group_clause(cond, None, [], {}, {}, [], None),
                         ' GROUP BY CASE WHEN "cat" = \'1\' THEN "sales" ELSE "disc" END')

    def test_write_group_bys_roundtrip(self):
        import xml.etree.ElementTree as ET
        rml = ET.Element("rml")
        _v._write_group_bys_el(rml, ET, [
            {"field": "cat", "order": 1},
            {"field": "bucket", "order": 2,
             "cond": {"when_field": "cat", "op": "=", "value": "1",
                      "then_field": "sales", "else_field": ""}},
            {"field": "cat"},
            {"field": "  "},
        ])
        gbs = rml.findall("group_bys/gb")
        self.assertEqual(len(gbs), 2)
        self.assertEqual(gbs[1].get("when_field"), "cat")
        self.assertEqual(gbs[1].get("then_field"), "sales")
        from rml_python.compiler import RMLReportCompiler
        xml = ET.tostring(rml, encoding="unicode")
        comp = RMLReportCompiler(xml_text="<rml>" + xml + "</rml>")
        back = comp.group_bys()
        self.assertEqual(len(back), 2)
        self.assertEqual(back[1]["cond"]["value"], "1")

    def test_table_opts_display_source_roundtrip(self):
        import xml.etree.ElementTree as ET
        rml = ET.Element("rml")
        _v._write_table_opts_el(rml, ET, [
            {"name": "T", "conn": "1", "display": "المبيعات", "source": "select"},
            {"name": "U", "conn": "1", "display": "T", "source": "bogus"},
        ])
        tels = rml.findall("table_opts/table")
        self.assertEqual(tels[0].get("display"), "المبيعات")
        self.assertEqual(tels[0].get("source"), "select")
        self.assertEqual(tels[1].get("display"), "T")
        self.assertIsNone(tels[1].get("source"))
        from rml_python.compiler import RMLReportCompiler
        comp = RMLReportCompiler(
            xml_text="<rml>" + ET.tostring(rml, encoding="unicode") + "</rml>")
        opts = comp.table_opts()
        self.assertEqual(opts[0]["display"], "المبيعات")
        self.assertEqual(opts[0]["source"], "select")
        self.assertEqual(opts[1]["source"], "")

    def test_designer_groupby_import_markup(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "..", "rml_python", "shared",
                                "rml_wizard_script.html"), encoding="utf-8") as f:
            js = f.read()
        for needle in ("wizardGroupBys", "wbGroupBysHtml", "wbToggleGroupBy",
                       "wbOpenGroupByModal", "wbColPickHTML", "group_bys",
                       "importMulti", "importSelTables", "importMergeMode",
                       "importUnionSel", "importOneEntry", "UNION ALL",
                       "from select", "updateImportFromPreview",
                       "wbOpenTableColsModal", "wbFieldUsedBy", "applyTableColsModal",
                       "tableColsModal", "wbTableColsFilter"):
            self.assertIn(needle, js)
        with open(os.path.join(base, "..", "rml_python", "shared",
                                "rml_wizard_modal.html"), encoding="utf-8") as f:
            modal = f.read()
        for needle in ("max-w-6xl", "importTableSel", "importSelTables",
                       "importFromPreview", "importMergeMode", "importColSel"):
            self.assertIn(needle, modal)


class FmlkExternalTests(SimpleTestCase):
    """Forms: external (read-only) connections for select sources + computed + details."""

    _XML = ('<fml><fml_metadata name="t" displayName="t" table="m" connection="urs_local"/>'
            '<fields>'
            '<field id="1" name="it" alias="صنف" inputType="select" refTable="ITEMS" refFk="ID" refDisplay="NM" refConn="ORCL"/>'
            '<field id="2" name="s" alias="س" inputType="select" options="a,b,[S.T.C]">'
            '<options_source><source table="T2" schema="S2" column="C2" conn="ORCL2"/></options_source>'
            "</field></fields>"
            '<details><detail table="IAS.ITEMS" alias="ا" master="id" link_mode="junction" conn="ORCL" sub_key="ID"/>'
            '<detail table="LOC" alias="ب" master="id" link_mode="direct" detail="m_id"/>'
            "</details></fml>")

    def test_field_ref_conn_roundtrip(self):
        from fmlk_engine.compiler import FMLKFormCompiler
        f = FMLKFormCompiler(xml_text=self._XML).fields()
        self.assertEqual(f[0].ref_conn, "ORCL")
        d = f[0].to_dict()
        self.assertEqual(d["refConn"], "ORCL")
        self.assertEqual(d["ref_conn"], "ORCL")
        self.assertIsNone(f[1].ref_conn)

    def test_options_source_conn_roundtrip(self):
        from fmlk_engine.compiler import FMLKFormCompiler
        f = FMLKFormCompiler(xml_text=self._XML).fields()
        # legacy [S.T.C] ref stays conn-less
        _legacy = [s for s in f[1].options_source
                   if s.get("table") == "T" and s.get("schema") == "S" and s.get("column") == "C"]
        self.assertEqual(len(_legacy), 1)
        self.assertNotIn("conn", _legacy[0])
        # explicit child carries conn
        self.assertIn({"table": "T2", "schema": "S2", "column": "C2", "conn": "ORCL2"},
                      f[1].options_source)

    def test_detail_conn_schema_roundtrip(self):
        from fmlk_engine.compiler import FMLKFormCompiler
        d = FMLKFormCompiler(xml_text=self._XML).details()
        self.assertEqual((d[0].table, d[0].schema, d[0].conn, d[0].link_mode),
                         ("ITEMS", "IAS", "ORCL", "junction"))
        dd = d[0].to_dict()
        self.assertEqual(dd["conn"], "ORCL")
        self.assertEqual(dd["schema"], "IAS")
        self.assertEqual((d[1].conn, d[1].schema), (None, None))

    def test_fmlk_ident_and_unknown_conn(self):
        from urs.views import _fmlk_ident, _fmlk_ext_options
        self.assertTrue(_fmlk_ident("ITEMS"))
        self.assertTrue(_fmlk_ident("IAS2$T"))
        self.assertFalse(_fmlk_ident("a-b"))
        self.assertFalse(_fmlk_ident(""))
        self.assertFalse(_fmlk_ident("1ABC"))
        with self.assertRaises(ValueError):
            _fmlk_ext_options("NOPE", "", "T", "bad-col!", "", "", 10)
        with self.assertRaises(ValueError):
            _fmlk_ext_options("conn-missing-xyz", "", "T", "C", "", "", 10)

    def test_junction_detail_text_type(self):
        from urs.custom_models import pg_type
        self.assertEqual(pg_type({"type": "VARCHAR", "length": 200}), "VARCHAR(200)")
        self.assertEqual(pg_type({"type": "INTEGER"}), "INTEGER")

    def test_forms_wizard_external_markup(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "..", "fmlk_engine", "shared",
                                "forms_wizard_script.html"), encoding="utf-8") as f:
            js = f.read()
        for needle in ("fwExtConns", "fwExtConnOpts", "fwTypeRefConn", "fwTypeRefConnChanged",
                        "fwTablesFor", "fwBranchConnChanged", "fwTypeSrcsTables",
                        "قراءة فقط", "جدول ثالث", "UniExpression.openForId('fwDisp_${i}')",
                        "r.refConn = ", "conn: String(b.conn",
                        "UniDBPicker.pickTable", "UniDBPicker.pickColumn", "UniDBPicker.setSelect",
                        "fwMainTableSearch", "fwBranchTblSearch", "fwBranchSubSearch",
                        "fwTypeRefTableSearch", "fwTypeRefColSearch",
                        "fwTypeSrcTblSearch", "fwTypeSrcColSearch"):
            self.assertIn(needle, js)
        self.assertEqual(js.count("{"), js.count("}"))
        self.assertEqual(js.count("("), js.count(")"))

    def test_forms_player_external_markup(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "templates", "forms_player.html"), encoding="utf-8") as f:
            html = f.read()
        for needle in ("s.conn", "field.refConn || field.ref_conn",
                        "_dmBody", "loadOptionsSourceRef(refT, fk, dp,"):
            self.assertIn(needle, html)


class FmlkCondModalTests(SimpleTestCase):
    """datamodal المشروط: تعبير الجدول حسب مدخلات النموذج + مودال برقم مرجعي/اسم."""

    _XML = """<fml><fml_metadata name="t" displayName="t" table="T0"/>
      <fields>
        <field id="1" name="doc_type" alias="النوع" inputType="select"/>
        <field id="2" name="item_no" alias="الصنف" inputType="datamodal"
          src_expr="if(@doc_type='items',[IAS_ITM_MST],[SUP_MST])"
          key_col="NO" label_col="NAME"/>
      </fields></fml>"""

    def test_config_parse(self):
        from fmlk_engine.compiler import FMLKFormCompiler
        comp = FMLKFormCompiler(xml_text=self._XML)
        f = {x.name: x for x in comp.fields()}["item_no"]
        self.assertEqual(f.input_type, "datamodal")
        self.assertEqual(f.config.get("src_expr"),
                         "if(@doc_type='items',[IAS_ITM_MST],[SUP_MST])")
        self.assertEqual(f.config.get("key_col"), "NO")
        self.assertEqual(f.config.get("label_col"), "NAME")
        self.assertEqual(f.to_dict()["config"].get("src_expr"),
                         "if(@doc_type='items',[IAS_ITM_MST],[SUP_MST])")

    def test_registry_and_save_cover_new_keys(self):
        from fmlk_engine.field_types import CONFIG_ATTRS, get_type
        for k in ("src_expr", "key_col", "label_col"):
            self.assertIn(k, CONFIG_ATTRS["datamodal"])
        self.assertIn("src_expr", get_type("datamodal").get("attrs", []))
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "views.py"), encoding="utf-8") as fh:
            vw = fh.read()
        for k in ('"src_expr"', '"key_col"', '"label_col"'):
            self.assertIn(k, vw)

    def test_designer_markup(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "..", "fmlk_engine", "shared",
                                "forms_wizard_script.html"), encoding="utf-8") as fh:
            js = fh.read()
        for needle in ("fwTypeSrcExpr", "fwTypeKeyCol", "fwTypeLabelCol",
                       "fwSrcExprInsert", "needCondSrc", "تعبير الجدول المشروط",
                       "NO←CODE←ID", "r.config.src_expr"):
            self.assertIn(needle, js)

    def test_player_markup(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "templates", "forms_player.html"), encoding="utf-8") as fh:
            js = fh.read()
        for needle in ("fmlkDmEvalTable", "fmlkDmTableLiterals", "fmlkDmAtToBracket",
                       "fmlkDmSrcDeps", "fmlkDmKeyLabel", "fmlkDmConnId",
                       "fmlkDmColumns", "dmSyncDependents", "data-dmsrc",
                       "الرقم المرجعي", "&display="):
            self.assertIn(needle, js)

    def test_cond_expr_end_to_end(self):
        # المقيّم الحقيقي fmlkCalcEval + مساعدات التعبير المشروط (node).
        import shutil
        import subprocess
        import tempfile
        base = os.path.dirname(__file__)
        src = os.path.join(base, "templates", "forms_player.html")
        if not shutil.which("node"):
            self.skipTest("node missing")
        with open(src, encoding="utf-8") as fh:
            js = fh.read()
        c0 = js.index("window.fmlkCalcEval = window.fmlkCalcEval || function(expr, get, rownum){")
        c1 = js.index("// الرقم التسلسلي التالي", c0)
        h0 = js.index("// ── datamodal المشروط")
        h1 = js.index("// ── datamodal: اختيار القوائم", h0)
        runner = (
            "var window={};"
            + js[c0:c1] +
            "var fmlMeta={fields:[{name:'doc_type'},{name:'item_no'}]};"
            "var VALS={doc_type:'items'};"
            "function fmlkGet(k){var v=VALS[k];return v===undefined?'':v;}"
            "function notify(){};var CSS={escape:function(s){return s;}};"
            "var document={querySelector:function(){return null;},getElementById:function(){return null;}};"
            "function escHtml(s){return String(s==null?'':s);}function escAttr(s){return String(s==null?'':s);}"
            + js[h0:h1] +
            "var assert=require('assert');"
            "var f={name:'item_no',alias:'الصنف',config:{src_expr:"
            "\"if(@doc_type='items',[IAS_ITM_MST],if(@doc_type='suppliers',[SUP_MST],[CUST_MST]))\"}};"
            "assert.strictEqual(fmlkDmEvalTable(f).table,'IAS_ITM_MST');"
            "VALS={doc_type:'suppliers'};"
            "assert.strictEqual(fmlkDmEvalTable(f).table,'SUP_MST');"
            "VALS={doc_type:'other'};"
            "assert.strictEqual(fmlkDmEvalTable(f).table,'CUST_MST');"
            "assert.deepStrictEqual(fmlkDmSrcDeps(f),['doc_type']);"
            "var kl=fmlkDmKeyLabel({},[{name:'ID'},{name:'NO'},{name:'NAME'}]);"
            "assert.strictEqual(kl.key,'NO');assert.strictEqual(kl.label,'NAME');"
            "kl=fmlkDmKeyLabel({},[{name:'CODE'},{name:'DESCR'}]);"
            "assert.strictEqual(kl.key,'CODE');assert.strictEqual(kl.label,'CODE');"
            "kl=fmlkDmKeyLabel({config:{key_col:'MY_ID',label_col:'MY_NM'}},[{name:'ID'}]);"
            "assert.strictEqual(kl.key,'MY_ID');assert.strictEqual(kl.label,'MY_NM');"
            "console.log('DM-COND OK');"
        )
        with tempfile.NamedTemporaryFile("w", suffix=".js", delete=False, encoding="utf-8") as t:
            t.write(runner)
            tmp = t.name
        try:
            r = subprocess.run(["node", tmp], capture_output=True, text=True)
        finally:
            try:
                os.unlink(tmp)
            except OSError:
                pass
        self.assertEqual(r.returncode, 0, (r.stderr or "")[-2000:])
        self.assertIn("DM-COND OK", r.stdout)


class DistributedReportTests(SimpleTestCase):
    """Two databases in one report: cross-connection link, Python-side merge."""

    _XML = """<rml><rpt_metadata name="t" displayName="t"/>
      <rml_connections>
        <connection id="1" connection_id="A"/>
        <connection id="2" connection_id="B"/>
      </rml_connections>
      <fields>
        <field name="BILL_NO" table_source="T1" conn_id="1" type="TEXT"/>
        <field name="AMT" table_source="T1" conn_id="1" type="NUMERIC"/>
        <field name="CNO" table_source="T2" conn_id="2" type="TEXT"/>
        <field name="CNAME" table_source="T2" conn_id="2" type="TEXT"/>
      </fields>
      <columns>
        <column id="1" name="bno" alias="bno" expr="[T1.BILL_NO]"/>
        <column id="2" name="amt" alias="amt" expr="[T1.AMT]"/>
        <column id="3" name="cn" alias="cn" expr="[T2.CNAME]"/>
      </columns>
      <links>
        <link from_table="T1" from_col="BILL_NO" to_table="T2" to_col="CNO" rel_type="one_to_many" match="exact" pattern=""/>
      </links>
      <table_opts><table name="T1" conn="1"/></table_opts>
      </rml>"""

    class _Cur:
        def __init__(self, desc, rows):
            self.description = [(c,) for c in desc]
            self._rows = rows

        def fetchall(self):
            return list(self._rows)

        def fetchone(self):
            return self._rows[0] if self._rows else None

        def close(self):
            pass

    class _DB:
        def __init__(self, handler):
            self._handler = handler
            self.conn = True

        def connect(self):
            pass

        def _exec(self, sql, params):
            desc, rows = self._handler(sql, params)
            return DistributedReportTests._Cur(desc, rows)

    def test_two_db_link_merges_in_python(self):
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine

        def _h1(sql, params):
            return (["BILL_NO", "AMT"], [("A1", 10), ("A2", 20)])

        def _h2(sql, params):
            return (["CNO", "CNAME"], [("A1", "Alpha"), ("A3", "Gamma")])

        comp = RMLReportCompiler(xml_text=self._XML)
        eng = RMLReportEngine(comp, self._DB(_h1), databases={"A": self._DB(_h1), "B": self._DB(_h2)})
        res = eng.execute({"distributed": True, "activeTable": "T1", "page": 1, "pageSize": 50})
        self.assertNotIn("error", res)
        rows = res["rows"]
        self.assertEqual(len(rows), 2)
        by_bno = {r["bno"]: r for r in rows}
        self.assertEqual(by_bno["A1"]["cn"], "Alpha")
        self.assertIn("cn", by_bno["A2"])


class ViewsTests(SimpleTestCase):
    """Displayed tables (VIEW): RML language + multi-branch CASE + designer."""

    def test_compiler_views(self):
        from rml_python.compiler import RMLReportCompiler
        comp = RMLReportCompiler(xml_text=(
            '<rml><views><view name="view_1" display="عرض 1">'
            '<sources><source id="s1" table_name="T1" conn="2"/>'
            '<source id="s2" table_name="T2" conn="2"/></sources>'
            '<merges><merge type="join" source_1_id="s1" source_2_id="s2" '
            'join_type="LEFT" key_1="ID" key_2="TID"/>'
            '<merge type="union_all" source_1_id="s1" source_2_id="s2"/></merges>'
            '<columns><column name="total" alias="total" expr="SUM([s1.AMT])" type="NUMERIC"/>'
            '</columns>'
            '<groups><group source_column="[s1.CAT]" expression="SUM([s1.AMT])" '
            'type="normal" agg="SUM" agg_col="[s1.AMT]"/>'
            '<group source_column="[s1.CAT]" type="conditional">'
            '<conditions><g_if if="[s1.CAT]" value="1" op="=" then="[s1.AMT]"/>'
            '<g_if if="[s1.CAT]" value="2" then="[s1.DIS]"/></conditions>'
            '</group></groups>'
            '</view></views></rml>'))
        views = comp.views()
        self.assertEqual(len(views), 1)
        self.assertEqual(views[0]["name"], "view_1")
        self.assertEqual(len(views[0]["sources"]), 2)
        self.assertEqual(views[0]["merges"][0]["key_1"], "ID")
        self.assertEqual(views[0]["merges"][1]["type"], "union_all")
        self.assertEqual(views[0]["columns"][0]["expr"], "SUM([s1.AMT])")
        self.assertEqual(views[0]["groups"][0]["agg"], "SUM")
        self.assertEqual(len(views[0]["groups"][1]["conditions"]), 2)

    def test_write_views_roundtrip(self):
        import xml.etree.ElementTree as ET
        rml = ET.Element("rml")
        _v._write_views_el(rml, ET, [{
            "name": "view_1", "display": "عرض 1",
            "sources": [{"id": "s1", "table_name": "T1", "conn": "2"},
                        {"id": "s2", "table_name": "", "conn": "2"}],
            "merges": [{"type": "join", "s1": "s1", "s2": "s2",
                        "join_type": "LEFT", "key_1": "ID", "key_2": "TID"}],
            "columns": [{"name": "total", "alias": "total",
                         "expr": "SUM([s1.AMT])", "type": "NUMERIC"},
                        {"name": "  "}],
            "groups": [{"name": "total_qty", "group_by": "total",
                        "expr": "sum([s1.P_qty])"}],
        }])
        _v._write_views_el(rml, ET, [{"display": "noname"}])
        v = rml.find("views/view")
        self.assertIsNotNone(v)
        self.assertEqual(v.get("name"), "view_1")
        self.assertEqual(len(v.findall("sources/source")), 1)
        self.assertEqual(len(v.findall("columns/column")), 1)
        self.assertEqual(v.find("merges/merge").get("key_1"), "ID")
        g = v.find("groups/group")
        self.assertEqual(g.get("name"), "total_qty")
        self.assertEqual(g.get("group_by"), "total")
        self.assertEqual(g.get("expr"), "sum([s1.P_qty])")
        from rml_python.compiler import RMLReportCompiler
        comp = RMLReportCompiler(
            xml_text="<rml>" + ET.tostring(rml, encoding="unicode") + "</rml>")
        back = comp.views()
        self.assertEqual(len(back), 1)
        self.assertEqual(back[0]["groups"][0]["name"], "total_qty")
        self.assertEqual(back[0]["groups"][0]["group_by"], "total")
        self.assertEqual(back[0]["groups"][0]["expr"], "sum([s1.P_qty])")

    def test_compiler_views_legacy_groups_preserved(self):
        # Saved RML with the old group shape still parses with legacy keys
        # (designer migrates on open; engine keeps the legacy GROUP BY path).
        from rml_python.compiler import RMLReportCompiler
        comp = RMLReportCompiler(xml_text=(
            '<rml><views><view name="view_1" display="عرض 1">'
            '<sources><source id="s1" table_name="T1" conn="2"/></sources>'
            '<merges/>' '<columns/>'
            '<groups><group source_column="[s1.CAT]" expression="SUM([s1.AMT])" '
            'type="normal" agg="SUM" agg_col="[s1.AMT]"/>'
            '</groups>'
            '</view></views></rml>'))
        views = comp.views()
        self.assertEqual(len(views), 1)
        g = views[0]["groups"][0]
        self.assertEqual(g["agg"], "SUM")
        self.assertEqual(g["source_column"], "[s1.CAT]")
        self.assertNotIn("name", g)

    def test_multibranch_case_clause(self):
        from rml_python.engine import _build_group_clause
        cond = [{"field": "bucket", "cond": {
            "when_field": "cat", "op": "=", "value": "", "then_field": "sales",
            "else_field": "",
            "branches": [
                {"when_field": "cat", "op": "=", "value": "1", "then_field": "sales"},
                {"when_field": "cat", "op": "=", "value": "2", "then_field": "disc"},
            ]}}]
        clause = _build_group_clause(cond, None, [], {}, {}, [], None)
        self.assertIn("WHEN", clause)
        self.assertEqual(clause.count("WHEN"), 2)
        self.assertIn("ELSE NULL END", clause)
        legacy = [{"field": "b", "cond": {"when_field": "cat", "op": "=",
                                          "value": "1", "then_field": "sales"}}]
        self.assertIn("ELSE NULL END",
                      _build_group_clause(legacy, None, [], {}, {}, [], None))

    def test_designer_views_markup(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "..", "rml_python", "shared",
                                "rml_wizard_script.html"), encoding="utf-8") as f:
            js = f.read()
        for needle in ("wizardViews", "wbViewsHtml", "wbOpenViewModal",
                       "applyViewModal", "wbViewBuildSQL", "wbImportView",
                       "wbDelView", "viewAddMerge", "viewAddGroup",
                       "union_partial",
                       "اسم التجميع", "تعبير التجميع", "تجميع حسب العمود",
                       "wbViewGrpChipsHTML", "viewgexpr", "wbViewMigrateGroups",
                       "wbViewValidateGroups", "wbViewExpandGrp", "wbViewGroupByOptions",
                       "اضافة جدول معروض"):
            self.assertIn(needle, js)
        for gone in ("viewAddCondRow", "viewDelCondRow", "VW_AGG",
                     "viewgsrc", "viewgagg"):
            self.assertNotIn(gone, js)

    _VIEW_XML = """<rml>
      <rpt_metadata name="vt" displayName="vt" schema=""/>
      <fields>
        <field name="BILL_NO" table_source="BILL_MST" type="TEXT"/>
        <field name="total" table_source="view_1" type="NUMERIC"/>
        <field name="BILL_NO" table_source="view_1" type="TEXT"/>
      </fields>
      <columns>
        <column id="1" name="total" alias="الاجمالي" expr="[view_1.total]"/>
      </columns>
      <links>
        <link from_table="BILL_MST" from_col="BILL_NO" to_table="view_1" to_col="BILL_NO" rel_type="one_to_many" match="exact" pattern=""/>
      </links>
      <views><view name="view_1" display="v1">
        <sources><source id="s1" table_name="BILL_MST" conn=""/><source id="s2" table_name="BILL_DTL" conn=""/></sources>
        <merges><merge type="join" source_1_id="s1" source_2_id="s2" join_type="LEFT" key_1="BILL_NO" key_2="BILL_NO"/></merges>
        <columns><column name="total" alias="total" expr="SUM([s1.AMT])" type="NUMERIC"/></columns>
        <groups/>
      </view></views>
    </rml>"""

    def _view_engine(self, xml=None, db=None, databases=None):
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        comp = RMLReportCompiler(xml_text=xml or self._VIEW_XML)
        eng = RMLReportEngine(comp, db if db is not None else mock.Mock())
        if databases:
            eng.databases.update(databases)
        return eng

    def _patch_physical_cols(self, eng):
        from rml_python.engine import RMLReportEngine
        real_tc = RMLReportEngine._table_columns

        def fake_tc(self, table, schema, db=None):
            if self._norm_table(table) == "VIEW_1":
                return real_tc(self, table, schema, db)
            return {"BILL_NO": "TEXT", "AMT": "NUMERIC"}

        eng._table_columns = fake_tc.__get__(eng, RMLReportEngine)
        return eng

    def test_view_outputs_rule(self):
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        comp = RMLReportCompiler(xml_text=self._VIEW_XML)
        eng = RMLReportEngine(comp, mock.Mock())
        v = eng._view_map()["VIEW_1"]
        # explicit alias + implicit join key (deduped across key_1/key_2)
        self.assertEqual(eng._view_outputs(v), ["total", "BILL_NO"])

    def test_compile_view_sql_join(self):
        eng = self._view_engine()
        sql = eng._compile_view_sql(eng._view_map()["VIEW_1"])
        self.assertIn('LEFT JOIN "BILL_DTL" "S2"', sql)
        self.assertIn('SUM("S1"."AMT") AS "total"', sql)
        self.assertIn('"S1"."BILL_NO" AS "BILL_NO"', sql)

    def test_compile_view_sql_union_nullfill(self):
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        comp = RMLReportCompiler(xml_text=(
            '<rml><rpt_metadata name="t" displayName="t"/><views><view name="vu">'
            '<sources><source id="s1" table_name="T1" conn=""/>'
            '<source id="s2" table_name="T2" conn=""/></sources>'
            '<merges><merge type="union_all" source_1_id="s1" source_2_id="s2"/></merges>'
            '<columns><column name="amt" alias="amt" expr="[s1.AMT]"/></columns>'
            '<groups/>' '</view></views>'
            '<columns><column id="1" name="d" alias="d" expr="1"/></columns></rml>'))
        eng = RMLReportEngine(comp, mock.Mock())
        sql = eng._compile_view_sql(eng._view_map()["VU"])
        self.assertIn("UNION ALL", sql)
        self.assertIn('NULL AS "amt"', sql)

    def test_compile_view_sql_cross_join_unlinked(self):
        # ORA-00904 regression: a source no merge touches must still appear
        # in FROM (CROSS JOIN) instead of dropping silently.
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        comp = RMLReportCompiler(xml_text=(
            '<rml><rpt_metadata name="t" displayName="t"/><views><view name="Sales">'
            '<sources><source id="s1" table_name="IAS_ITM_MST" conn=""/>'
            '<source id="s3" table_name="IAS_BILL_MST" conn=""/>'
            '<source id="s4" table_name="IAS_BILL_DTL" conn=""/></sources>'
            '<merges><merge type="join" source_1_id="s1" source_2_id="s4" '
            'join_type="INNER" key_1="I_CODE" key_2="I_CODE"/></merges>'
            '<columns><column name="item_code" alias="item_code" expr="[s1.I_CODE]"/>'
            '<column name="bill_ser" alias="BILL_SER" expr="[s3.BILL_SER]"/>'
            '<column name="sum_p_qty" alias="sum_p_qty" expr="SUM([s4.P_QTY])"/>'
            '</columns><groups/>' '</view></views>'
            '<columns><column id="1" name="d" alias="d" expr="1"/></columns></rml>'))
        eng = RMLReportEngine(comp, mock.Mock())
        sql = eng._compile_view_sql(eng._view_map()["SALES"])
        self.assertIn('INNER JOIN "IAS_BILL_DTL" "S4"', sql)
        self.assertIn('CROSS JOIN "IAS_BILL_MST" "S3"', sql)
        self.assertIn('"S3"."BILL_SER" AS "bill_ser"', sql)

    def test_compile_view_sql_cross_join_no_merges(self):
        # Multi-source view with no merges at all: CROSS JOIN, not an error.
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        comp = RMLReportCompiler(xml_text=(
            '<rml><rpt_metadata name="t" displayName="t"/><views><view name="v2">'
            '<sources><source id="s1" table_name="T1" conn=""/>'
            '<source id="s2" table_name="T2" conn=""/></sources>'
            '<merges/>' '<columns>'
            '<column name="a" alias="a" expr="[s1.A]"/>'
            '<column name="b" alias="b" expr="[s2.B]"/>'
            '</columns><groups/>' '</view></views>'
            '<columns><column id="1" name="d" alias="d" expr="1"/></columns></rml>'))
        eng = RMLReportEngine(comp, mock.Mock())
        sql = eng._compile_view_sql(eng._view_map()["V2"])
        self.assertIn('FROM "T1" "S1" CROSS JOIN "T2" "S2"', sql)

    def test_compile_view_sql_join_chain_order(self):
        # ORA-00904 regression (real Sales def): first merge (s3-s4) must not
        # emit ON naming S3 before S3 enters FROM — the chain seeds FROM with
        # the first merge's s1 and JOINs each new side only; the both-joined
        # merge becomes a trailing AND (no duplicate table).
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        comp = RMLReportCompiler(xml_text=(
            '<rml><rpt_metadata name="t" displayName="t"/><views><view name="Sales">'
            '<sources><source id="s1" table_name="IAS_ITM_MST" conn=""/>'
            '<source id="s2" table_name="IAS_ITM_DTL" conn=""/>'
            '<source id="s3" table_name="IAS_BILL_MST" conn=""/>'
            '<source id="s4" table_name="IAS_BILL_DTL" conn=""/></sources>'
            '<merges>'
            '<merge type="join" source_1_id="s3" source_2_id="s4" join_type="INNER" key_1="BILL_SER" key_2="BILL_SER"/>'
            '<merge type="join" source_1_id="s4" source_2_id="s1" join_type="INNER" key_1="I_CODE" key_2="I_CODE"/>'
            '<merge type="join" source_1_id="s4" source_2_id="s2" join_type="INNER" key_1="ITM_UNT" key_2="ITM_UNT"/>'
            '<merge type="join" source_1_id="s1" source_2_id="s2" join_type="LEFT" key_1="I_CODE" key_2="I_CODE"/>'
            '</merges>'
            '<columns><column name="item_code" alias="item_code" expr="[[s1.I_CODE]]"/>'
            '<column name="item_unit" alias="item_unit" expr="[[s2.ITM_UNT]]"/>'
            '<column name="sum_p_qty" alias="sum_p_qty" expr="%%p_qty%%"/>'
            '</columns>'
            '<groups><group name="p_qty" group_by="item_code" expr="SUM([[s4.P_QTY]])"/>'
            '</groups>' '</view></views>'
            '<columns><column id="1" name="d" alias="d" expr="1"/></columns></rml>'))
        eng = RMLReportEngine(comp, mock.Mock())
        sql = eng._compile_view_sql(eng._view_map()["SALES"])
        # FROM opens with the first merge's s1; every ON-table precedes its use
        _from = sql.split("FROM", 1)[1]
        self.assertTrue(_from.lstrip().startswith('"IAS_BILL_MST" "S3"'),
                        "chain must seed FROM with s3 (first merge's s1): " + _from[:80])
        self.assertIn('INNER JOIN "IAS_BILL_DTL" "S4" ON "S3"."BILL_SER" = "S4"."BILL_SER"', sql)
        self.assertIn('INNER JOIN "IAS_ITM_MST" "S1" ON "S4"."I_CODE" = "S1"."I_CODE"', sql)
        self.assertIn('INNER JOIN "IAS_ITM_DTL" "S2" ON "S4"."ITM_UNT" = "S2"."ITM_UNT"', sql)
        # both-joined (s1-s2) merge → trailing AND, S2's table joined exactly once
        self.assertIn('AND "S1"."I_CODE" = "S2"."I_CODE"', sql)
        self.assertEqual(_from.count('"IAS_ITM_DTL" "S2"'), 1)
        self.assertIn('SUM("S4"."P_QTY") AS "p_qty"', sql)
        # ORA-00979: bare select items extend GROUP BY (user key stays first)
        self.assertIn('GROUP BY "S1"."I_CODE", "S2"."ITM_UNT", "S3"."BILL_SER", '
                      '"S4"."I_CODE", "S4"."ITM_UNT"', sql)
        _gb = sql.split("GROUP BY", 1)[1]
        self.assertNotIn("SUM(", _gb)

    def test_compile_view_sql_self_merge_rejected(self):
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        comp = RMLReportCompiler(xml_text=(
            '<rml><rpt_metadata name="t" displayName="t"/><views><view name="v9">'
            '<sources><source id="s1" table_name="T1" conn=""/></sources>'
            '<merges><merge type="join" source_1_id="s1" source_2_id="s1" '
            'join_type="INNER" key_1="A" key_2="A"/></merges>'
            '<columns/>' '<groups/>' '</view></views>'
            '<columns><column id="1" name="d" alias="d" expr="1"/></columns></rml>'))
        eng = RMLReportEngine(comp, mock.Mock())
        with self.assertRaises(ValueError) as ctx:
            eng._compile_view_sql(eng._view_map()["V9"])
        self.assertIn("بنفسه", str(ctx.exception))

    def test_compile_view_sql_errors(self):
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        multi = RMLReportCompiler(xml_text=(
            '<rml><rpt_metadata name="t" displayName="t"/><views><view name="vx">'
            '<sources><source id="s1" table_name="A" conn="1"/>'
            '<source id="s2" table_name="B" conn="2"/></sources>'
            '<merges/>' '<columns/>' '<groups/>'
            '</view></views>'
            '<columns><column id="1" name="d" alias="d" expr="1"/></columns></rml>'))
        eng = RMLReportEngine(multi, mock.Mock())
        with self.assertRaises(ValueError) as ctx:
            eng._compile_view_sql(eng._view_map()["VX"])
        self.assertIn("اتصالات مختلفة", str(ctx.exception))
        badref = RMLReportCompiler(xml_text=(
            '<rml><rpt_metadata name="t" displayName="t"/><views><view name="vy">'
            '<sources><source id="s1" table_name="A" conn=""/></sources>'
            '<merges/>' '<columns><column name="x" alias="x" expr="[s9.NOPE]"/></columns>'
            '<groups/>'
            '</view></views>'
            '<columns><column id="1" name="d" alias="d" expr="1"/></columns></rml>'))
        eng2 = RMLReportEngine(badref, mock.Mock())
        with self.assertRaises(ValueError) as ctx2:
            eng2._compile_view_sql(eng2._view_map()["VY"])
        self.assertIn("[s9.NOPE]", str(ctx2.exception))

    def test_named_view_groups_join_scope(self):
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        comp = RMLReportCompiler(xml_text=(
            '<rml><rpt_metadata name="t" displayName="t"/><views><view name="v1">'
            '<sources><source id="s1" table_name="T1" conn=""/></sources>'
            '<merges/>' '<columns>'
            '<column name="bno" alias="bno" expr="[s1.BILL_NO]"/>'
            '<column name="net" alias="net" expr="%%total_qty%% * 2"/>'
            '</columns>' '<groups>'
            '<group name="total_qty" group_by="bno" expr="sum([s1.P_qty])"/>'
            '</groups>' '</view></views>'
            '<columns><column id="1" name="d" alias="d" expr="1"/></columns></rml>'))
        eng = RMLReportEngine(comp, mock.Mock())
        v = eng._view_map()["V1"]
        self.assertEqual(eng._view_outputs(v), ["bno", "net", "total_qty"])
        sql = eng._compile_view_sql(v)
        self.assertIn('sum("S1"."P_qty") AS "total_qty"', sql)
        self.assertIn('sum("S1"."P_qty") * 2 AS "net"', sql)
        self.assertIn('GROUP BY "S1"."BILL_NO"', sql)

    def test_named_view_groups_union_scope(self):
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        comp = RMLReportCompiler(xml_text=(
            '<rml><rpt_metadata name="t" displayName="t"/><views><view name="vu">'
            '<sources><source id="s1" table_name="T1" conn=""/>'
            '<source id="s2" table_name="T2" conn=""/></sources>'
            '<merges><merge type="union_all" source_1_id="s1" source_2_id="s2"/></merges>'
            '<columns><column name="amt" alias="amt" expr="[s1.AMT]"/>'
            '<column name="cat" alias="cat" expr="[s1.CAT]"/></columns>'
            '<groups><group name="total_amt" group_by="cat" expr="sum([s1.AMT])"/>'
            '</groups>' '</view></views>'
            '<columns><column id="1" name="d" alias="d" expr="1"/></columns></rml>'))
        eng = RMLReportEngine(comp, mock.Mock())
        sql = eng._compile_view_sql(eng._view_map()["VU"])
        self.assertIn("UNION ALL", sql)
        self.assertIn('sum("U"."amt") AS "total_amt"', sql)
        self.assertIn('GROUP BY "U"."cat", "U"."amt"', sql)
        # branches stay core: no group aggregate inside the UNION legs
        _core = sql.split('FROM (', 1)[1].rsplit(') "U"', 1)[0]
        self.assertNotIn("total_amt", _core)

    def test_named_view_groups_errors(self):
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine

        def _sql(groups_xml):
            comp = RMLReportCompiler(xml_text=(
                '<rml><rpt_metadata name="t" displayName="t"/><views><view name="v1">'
                '<sources><source id="s1" table_name="T1" conn=""/></sources>'
                '<merges/>' '<columns>'
                '<column name="bno" alias="bno" expr="[s1.BILL_NO]"/>'
                '<column name="tot" alias="tot" expr="sum([s1.AMT])"/>'
                '</columns>' '<groups>' + groups_xml + '</groups>' '</view></views>'
                '<columns><column id="1" name="d" alias="d" expr="1"/></columns></rml>'))
            eng = RMLReportEngine(comp, mock.Mock())
            return eng._compile_view_sql(eng._view_map()["V1"])

        with self.assertRaises(ValueError) as ctx:
            _sql('<group name="a" group_by="bno" expr="%%nope%%"/>')
        self.assertIn("%%nope%%", str(ctx.exception))
        with self.assertRaises(ValueError) as ctx:
            _sql('<group name="a" group_by="bno" expr="%%a%%"/>')
        self.assertIn("دائري", str(ctx.exception))
        with self.assertRaises(ValueError) as ctx:
            _sql('<group name="bno" group_by="bno" expr="sum([s1.P_qty])"/>')
        self.assertIn("يطابق عموداً", str(ctx.exception))
        with self.assertRaises(ValueError) as ctx:
            _sql('<group name="bad name" group_by="bno" expr="sum([s1.P_qty])"/>')
        self.assertIn("غير صالح", str(ctx.exception))
        with self.assertRaises(ValueError) as ctx:
            _sql('<group name="a" group_by="tot" expr="sum([s1.P_qty])"/>')
        self.assertIn("عمود تجميعي", str(ctx.exception))
        with self.assertRaises(ValueError) as ctx:
            _sql('<group name="a" group_by="" expr="sum([s1.P_qty])"/>')
        self.assertIn("بلا عمود تجميع", str(ctx.exception))

    def test_build_plan_view_secondary(self):
        import dataclasses
        eng = self._patch_physical_cols(self._view_engine())
        plan = eng._build_plan("BILL_MST", [], None, None, dataclasses)
        self.assertIn("LEFT JOIN (SELECT", plan["from_q"])
        self.assertIn('"view_1"."BILL_NO" = "T0"."BILL_NO"', plan["from_q"])
        self.assertEqual(plan["table_map"].get("total"), "view_1")
        sql, _params = eng._compile_sql2(plan, [], None, 1, 50, None)
        self.assertIn('SELECT "view_1"."total"', sql)

    def test_build_plan_view_base(self):
        import dataclasses
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        comp = RMLReportCompiler(xml_text=self._VIEW_XML)
        eng = RMLReportEngine(comp, mock.Mock())
        plan = eng._build_plan("view_1", [], None, None, dataclasses)
        self.assertIn('(SELECT SUM("S1"."AMT")', plan["from_q"])
        self.assertTrue(plan["from_q"].strip().endswith('"view_1"'))
        self.assertEqual(plan["table_map"].get("total"), "view_1")
        self.assertEqual(plan.get("base_alias"), "view_1")

    def test_build_plan_view_cross_db_rejected(self):
        import dataclasses
        xml = self._VIEW_XML.replace(
            '<field name="total" table_source="view_1" type="NUMERIC"/>',
            '<field name="total" table_source="view_1" type="NUMERIC" conn_id="9"/>'
        ).replace(
            '<source id="s1" table_name="BILL_MST" conn=""/>',
            '<source id="s1" table_name="BILL_MST" conn="9"/>'
        ).replace(
            '<source id="s2" table_name="BILL_DTL" conn=""/>',
            '<source id="s2" table_name="BILL_DTL" conn="9"/>'
        )
        eng = self._patch_physical_cols(self._view_engine(
            xml=xml, db=mock.Mock(dsn="BASE"),
            databases={"9": mock.Mock(dsn="OTHER")}))
        with self.assertRaises(ValueError) as ctx:
            eng._build_plan("BILL_MST", [], None, None, dataclasses)
        self.assertIn("قاعدة مختلفة", str(ctx.exception))

    def test_view_nested_tags_dont_shadow_report(self):
        # Views nest their own <columns>/<groups> BEFORE report-level ones —
        # the report parser must prefer direct children.
        from rml_python.compiler import RMLReportCompiler
        comp = RMLReportCompiler(xml_text=(
            '<rml><rpt_metadata name="t" displayName="t"/>'
            '<views><view name="vv">'
            '<sources><source id="s1" table_name="T1" conn=""/></sources>'
            '<merges/>'
            '<columns><column name="vcol" alias="vcol" expr="[s1.X]"/></columns>'
            '<groups><group source_column="[s1.X]" type="normal" agg="SUM" agg_col="[s1.X]"/></groups>'
            '</view></views>'
            '<columns><column id="1" name="d" alias="d" expr="1"/></columns>'
            '<groups><group id="1" name="H1" order="1"><column alias="d"/></group></groups>'
            '</rml>'))
        cols = comp.columns()
        self.assertEqual(len(cols), 1)
        self.assertEqual(cols[0].alias, "d")
        grps = comp.groups()
        self.assertEqual(len(grps), 1)
        self.assertEqual(grps[0].name, "H1")
        self.assertEqual(len(comp.views()), 1)

    def test_designer_view_table_markup(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "..", "rml_python", "shared",
                                "rml_wizard_script.html"), encoding="utf-8") as f:
            js = f.read()
        for needle in ("wbSyncViewFields", "wbViewOutputs", "is_view",
                       "wbViewTokCol"):
            self.assertIn(needle, js)
        with open(os.path.join(base, "..", "urs", "templates",
                                "data_diagram.html"), encoding="utf-8") as f:
            dg = f.read()
        for needle in ("DG.views", "dgViewOutputs", "dgViewTokCol"):
            self.assertIn(needle, dg)


        for needle in ("DG.views", "dgViewOutputs", "dgViewTokCol"):
            self.assertIn(needle, dg)


class ViewWithImportTests(SimpleTestCase):
    """WITH … UNION ALL organized import → one view per CTE + main."""

    _WITH_SQL = """WITH trn AS (
        SELECT TRUNC(m.BILL_DATE) AS BILL_DATE, d.I_CODE,
            NVL(d.P_QTY, 0) AS SALES_QTY, 0 AS ACTUAL_QTY
        FROM IAS_BILL_MST m JOIN IAS_BILL_DTL d ON m.BILL_SER = d.BILL_SER
        WHERE d.CC_CODE IN ('11', '12')
        UNION ALL
        SELECT TRUNC(m.BILL_DATE) AS BILL_DATE, d.I_CODE,
            0 AS SALES_QTY, NVL(d.FREE_QTY, 0) AS ACTUAL_QTY
        FROM IAS_BILL_MST m JOIN IAS_BILL_DTL d ON m.BILL_SER = d.BILL_SER
        WHERE d.CC_CODE = '18'
    ),
    agg AS (
        SELECT t.I_CODE, t.BILL_DATE, SUM(t.SALES_QTY) AS SALES_QTY
        FROM trn t GROUP BY t.I_CODE, t.BILL_DATE
    )
    SELECT m.I_CODE, NVL(a.SALES_QTY, 0) AS SALES_QTY
    FROM IAS_ITM_MST m JOIN agg a ON m.I_CODE = a.I_CODE
    WHERE m.G_CODE NOT IN ('9000', '9500')"""

    def test_parse_with_views(self):
        from rml_python.sql_import import parse_with
        res = parse_with(self._WITH_SQL)
        self.assertEqual(res["mode"], "with")
        self.assertEqual([v["name"] for v in res["views"]], ["trn", "agg", "main"])
        trn = res["views"][0]
        # legs are independent FROMs: same tables twice, distinct sids
        self.assertEqual(len(trn["sources"]), 4)
        self.assertEqual([s["id"] for s in trn["sources"]], ["s1", "s2", "s3", "s4"])
        self.assertEqual(trn["sources"][0]["filter"], "[s2.CC_CODE] IN ('11', '12')")
        self.assertEqual(trn["sources"][2]["filter"], "[s4.CC_CODE] = '18'")
        self.assertIn(("join", "s1", "s2"), [(m["type"], m["s1"], m["s2"]) for m in trn["merges"]])
        self.assertIn(("union_all", "s1", "s3"), [(m["type"], m["s1"], m["s2"]) for m in trn["merges"]])
        self.assertEqual([c["name"] for c in trn["columns"]],
                         ["BILL_DATE", "I_CODE", "SALES_QTY", "ACTUAL_QTY"])
        agg = res["views"][1]
        self.assertEqual([(s["id"], s["table_name"]) for s in agg["sources"]], [("s1", "trn")])
        self.assertIn("SUM([s1.SALES_QTY])", [c["expr"] for c in agg["columns"]])
        main = res["views"][2]
        self.assertEqual(main["name"], "main")
        self.assertIn("agg", [s["table_name"] for s in main["sources"]])
        self.assertTrue(any("G_CODE" in (s.get("filter") or "") for s in main["sources"]))
        self.assertTrue(any("GROUP BY" in w for w in res["warnings"]))

    def test_parse_with_errors(self):
        from rml_python.sql_import import parse_with
        with self.assertRaises(ValueError):
            parse_with("SELECT 1")
        with self.assertRaises(ValueError):
            parse_with("WITH trn AS (SELECT 1), agg AS (SELECT 1) SELECT 1")
        with self.assertRaises(ValueError):
            parse_with("WITH a AS (SELECT 1 FROM T1), b AS (SELECT 1 FROM T2) "
                       "SELECT * FROM a UNION ALL SELECT * FROM b")

    def test_parse_with_cte_shadows_physical_like_sql(self):
        # دلالات SQL: الـ CTE يظلل الجدول الفيزيائي بنفس الاسم — المراجع تربط بالـ CTE
        from rml_python.sql_import import parse_with
        res = parse_with("WITH t AS (SELECT a.x FROM T1 a) "
                         "SELECT t.x, T1.y FROM t JOIN T1 ON t.x = T1.y")
        by_name = {v["name"]: v for v in res["views"]}
        self.assertIn("t", by_name)
        self.assertIn("main", by_name)
        t_srcs = [(s["id"], s["table_name"]) for s in by_name["main"]["sources"]]
        self.assertIn(("s1", "t"), t_srcs)
        self.assertIn(("s2", "T1"), t_srcs)

    def test_parse_with_endpoint(self):
        req = RequestFactory().post("/api/rml/parse-sql/",
                                    data=json.dumps({"sql": self._WITH_SQL}),
                                    content_type="application/json")
        res = _v.api_rml_parse_sql(req)
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        self.assertEqual(body["mode"], "with")
        self.assertEqual([v["name"] for v in body["views"]], ["trn", "agg", "main"])

    def test_with_views_compile_end_to_end(self):
        # imported model compiles: per-leg WHERE + nested agg GROUP BY + main join
        from rml_python.sql_import import parse_with
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        import xml.etree.ElementTree as ET
        res = parse_with(self._WITH_SQL)
        rml = ET.Element("rml")
        _v._write_views_el(rml, ET, res["views"])
        md = ET.SubElement(rml, "rpt_metadata")
        md.set("name", "t")
        md.set("displayName", "t")
        cols_el = ET.SubElement(rml, "columns")
        dc = ET.SubElement(cols_el, "column")
        dc.set("id", "1")
        dc.set("name", "d")
        dc.set("alias", "d")
        dc.set("expr", "1")
        comp = RMLReportCompiler(xml_text=ET.tostring(rml, encoding="unicode"))
        eng = RMLReportEngine(comp, mock.Mock())
        trn_sql = eng._compile_view_sql(eng._view_map()["TRN"])
        self.assertEqual(trn_sql.count("UNION ALL"), 1)
        self.assertIn('"S2"."CC_CODE" IN (\'11\', \'12\')', trn_sql)
        self.assertIn('"S4"."CC_CODE" = \'18\'', trn_sql)
        agg_sql = eng._compile_view_sql(eng._view_map()["AGG"])
        self.assertIn("SUM(", agg_sql)
        self.assertIn("GROUP BY", agg_sql)
        main_sql = eng._compile_view_sql(eng._view_map()["MAIN"])
        self.assertIn("NOT IN ('9000', '9500')", main_sql)

    def test_designer_with_import_markup(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "..", "rml_python", "shared",
                                "rml_wizard_script.html"), encoding="utf-8") as f:
            js = f.read()
        for needle in ("wbViewSrcRows", "wbViewSrcSetFilter", "viewUnionAll",
                        "wbViewComponents", "wbViewLegWhere", "wbUnionAdaptAsView",
                        "wbViewBuildCore", "wbViewFindNested", "sqlImportWithHTML",
                        "تفعيل UNION ALL للكل", "مرشح الفرع"):
            self.assertIn(needle, js)


class ViewLibraryTests(SimpleTestCase):
    """Shared view library: namespace.view_name calls from any report."""

    def test_designer_view_table_markup(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "..", "rml_python", "shared",
                                "rml_wizard_script.html"), encoding="utf-8") as f:
            js = f.read()
        for needle in ("wbSyncViewFields", "wbViewOutputs", "is_view",
                       "wbViewTokCol"):
            self.assertIn(needle, js)
        with open(os.path.join(base, "..", "urs", "templates",
                                "data_diagram.html"), encoding="utf-8") as f:
            dg = f.read()
        for needle in ("DG.views", "dgViewOutputs", "dgViewTokCol"):
            self.assertIn(needle, dg)

    def test_view_refs_roundtrip(self):
        import xml.etree.ElementTree as ET
        rml = ET.Element("rml")
        _v._write_view_refs_el(rml, ET, [
            {"namespace": "store", "name": "sales_net"},
            {"namespace": "store", "name": "sales_net"},
            {"namespace": "", "name": "  "},
            {"name": "local_v"},
        ])
        refs = rml.findall("view_refs/ref")
        self.assertEqual(len(refs), 2)
        self.assertEqual(refs[0].get("namespace"), "store")
        from rml_python.compiler import RMLReportCompiler
        comp = RMLReportCompiler(
            xml_text="<rml>" + ET.tostring(rml, encoding="unicode") + "</rml>")
        back = comp.view_refs()
        self.assertEqual(len(back), 2)
        self.assertEqual(back[1], {"namespace": "", "name": "local_v"})

    def test_library_merge_and_shadow(self):
        import pathlib
        import tempfile
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        with tempfile.TemporaryDirectory() as td:
            td = pathlib.Path(td)
            (td / "_view_library.rml").write_text(
                '<rml><rpt_metadata name="__view_library__" displayName="lib" namespace="store"/>'
                '<views><view name="sales_net" display="صافي">'
                '<sources><source id="s1" table_name="T1" conn=""/></sources>'
                '<merges/><columns><column name="amt" alias="amt" expr="[s1.AMT]"/></columns>'
                '<groups/></view>'
                '<view name="view_1" display="lib one">'
                '<sources><source id="s1" table_name="TLIB" conn=""/></sources>'
                '<merges/><columns><column name="x" alias="x" expr="[s1.X]"/></columns>'
                '<groups/></view></views></rml>', encoding="utf-8")
            (td / "rep.rml").write_text(
                '<rml><rpt_metadata name="r" displayName="r"/>'
                '<columns><column id="1" name="d" alias="d" expr="1"/></columns>'
                '<views><view name="view_1" display="local one">'
                '<sources><source id="s1" table_name="TLOC" conn=""/></sources>'
                '<merges/><columns><column name="y" alias="y" expr="[s1.Y]"/></columns>'
                '<groups/></view></views></rml>', encoding="utf-8")
            comp = RMLReportCompiler(path=td / "rep.rml")
            eng = RMLReportEngine(comp, mock.Mock())
            lib = eng._library_views()
            self.assertEqual(len(lib), 2)
            self.assertEqual(lib[0].get("_ns"), "store")
            # namespaced full map
            full = eng._view_full_map()
            self.assertIn("store.sales_net", full)
            # local shadows library on stripped clash
            m = eng._view_map()
            self.assertEqual(m["VIEW_1"]["sources"][0]["table_name"], "TLOC")
            self.assertEqual(m["SALES_NET"]["sources"][0]["table_name"], "T1")
            self.assertTrue(eng._is_view_call("store", "sales_net"))
            self.assertFalse(eng._is_view_call("store", "nope"))

    def test_ns_view_call_plan(self):
        import dataclasses
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        comp = RMLReportCompiler(xml_text=(
            '<rml><rpt_metadata name="vt" displayName="vt" schema=""/>'
            '<fields>'
            '<field name="BILL_NO" table_source="BILL_MST" type="TEXT"/>'
            '<field name="total" table_source="store.view_1" type="NUMERIC"/>'
            '<field name="BILL_NO" table_source="store.view_1" type="TEXT"/>'
            '</fields>'
            '<columns><column id="1" name="total" alias="t" expr="[store.view_1.total]"/></columns>'
            '<links><link from_table="BILL_MST" from_col="BILL_NO" to_table="store.view_1" to_col="BILL_NO" rel_type="one_to_many" match="exact" pattern=""/></links>'
            '<view_refs><ref namespace="store" name="view_1"/></view_refs>'
            '<views><view name="view_1">'
            '<sources><source id="s1" table_name="BILL_MST" conn=""/><source id="s2" table_name="BILL_DTL" conn=""/></sources>'
            '<merges><merge type="join" source_1_id="s1" source_2_id="s2" join_type="LEFT" key_1="BILL_NO" key_2="BILL_NO"/></merges>'
            '<columns><column name="total" alias="total" expr="SUM([s1.AMT])" type="NUMERIC"/></columns>'
            '<groups/></view></views></rml>'))
        eng = RMLReportEngine(comp, mock.Mock())
        self.assertTrue(eng._is_view_call("store", "view_1"))
        self.assertEqual(eng._match_qualified("store.view_1.total", eng._field_table_map()), "total")
        real_tc = RMLReportEngine._table_columns

        def fake_tc(self, table, schema, db=None):
            if self._norm_table(table) == "VIEW_1":
                return real_tc(self, table, schema, db)
            return {"BILL_NO": "TEXT", "AMT": "NUMERIC"}

        eng._table_columns = fake_tc.__get__(eng, RMLReportEngine)
        plan = eng._build_plan("BILL_MST", [], None, None, dataclasses)
        self.assertIn("LEFT JOIN (SELECT", plan["from_q"])
        sql, _ = eng._compile_sql2(plan, [], None, 1, 50, None)
        self.assertIn('SELECT "view_1"."total"', sql)

    def test_library_endpoints(self):
        import pathlib
        import tempfile
        from django.test import RequestFactory
        rf = RequestFactory()
        with tempfile.TemporaryDirectory() as td:
            td = pathlib.Path(td)
            target = td / "_view_library.rml"
            with mock.patch.object(_v, "_resolve_rml_target", return_value=(td, target, None)):
                req = rf.get("/api/apps/demo/views/library/")
                res = _v.api_view_library(req, "demo")
                self.assertEqual(res.status_code, 200)
                self.assertEqual(json.loads(res.content), {"namespace": "demo", "views": []})
                with mock.patch.object(_v, "_ws_require_design", return_value=None):
                    view = {"name": "sales_net", "display": "صافي",
                            "sources": [{"id": "s1", "table_name": "T1", "conn": ""}],
                            "merges": [], "columns": [], "groups": []}
                    req2 = rf.post("/api/apps/demo/views/library/", data=json.dumps(
                        {"action": "save", "view": view}), content_type="application/json")
                    res2 = _v.api_view_library(req2, "demo")
                    self.assertEqual(res2.status_code, 200)
                    body = json.loads(res2.content)
                    self.assertTrue(body["ok"])
                    self.assertEqual(len(body["views"]), 1)
                    self.assertTrue(target.exists())
                    req3 = rf.get("/api/apps/demo/views/library/")
                    res3 = _v.api_view_library(req3, "demo")
                    self.assertEqual(len(json.loads(res3.content)["views"]), 1)
                    req4 = rf.post("/api/apps/demo/views/library/", data=json.dumps(
                        {"action": "delete", "name": "sales_net"}), content_type="application/json")
                    res4 = _v.api_view_library(req4, "demo")
                    self.assertEqual(json.loads(res4.content)["views"], [])

    def test_library_route(self):
        from django.urls import resolve
        m = resolve("/api/apps/demo/views/library/")
        self.assertEqual(m.func, _v.api_view_library)

    def test_view_variables_roundtrip(self):
        import xml.etree.ElementTree as ET
        from rml_python.compiler import RMLReportCompiler
        rml = ET.Element("rml")
        _v._write_views_el(rml, ET, [{"name": "v1", "display": "V",
            "sources": [{"id": "s1", "table_name": "T1", "conn": ""}],
            "merges": [], "columns": [{"name": "a", "alias": "a", "expr": "[s1.A]"}],
            "groups": [], "variables": [{"name": "W", "type": "number", "default": "1001"},
                                        {"name": "", "type": "text", "default": "x"},
                                        {"name": "D", "type": "bogus", "default": ""}]}])
        comp = RMLReportCompiler(
            xml_text="<rml>" + ET.tostring(rml, encoding="unicode") + "</rml>")
        vs = comp.views()
        self.assertEqual(len(vs), 1)
        self.assertEqual(vs[0]["variables"], [{"name": "W", "type": "number", "default": "1001"},
                                              {"name": "D", "type": "text", "default": ""}])

    def test_doc_param_var_roundtrip(self):
        import xml.etree.ElementTree as ET
        from rml_python.compiler import RMLReportCompiler
        rml = ET.Element("rml")
        _v._write_doc_params_el(rml, ET, [{"column": "w_code", "label": "الكود",
                                           "type": "number", "var": "W"}])
        comp = RMLReportCompiler(
            xml_text="<rml>" + ET.tostring(rml, encoding="unicode") + "</rml>")
        dps = comp.doc_params()
        self.assertEqual(len(dps), 1)
        self.assertEqual(dps[0].var, "W")
        self.assertEqual(dps[0].to_dict()["var"], "W")

    def _var_eng(self, views=(), params=(), filters=None):
        from rml_python.engine import RMLReportEngine
        eng = RMLReportEngine.__new__(RMLReportEngine)
        eng.doc_params = list(params)
        eng.report_views = list(views)
        eng._active_filters = list(filters or [])
        eng.columns = []
        eng.metadata = {}
        return eng

    def test_expand_view_vars_number_default(self):
        eng = self._var_eng(
            views=[{"name": "v1", "variables": [{"name": "W", "type": "number", "default": "1001"}]}])
        self.assertEqual(eng._expand_view_vars("[T.W_CODE] = &W"), "[T.W_CODE] = 1001")
        # case-insensitive + string literal untouched
        self.assertEqual(eng._expand_view_vars("[T.C] = &w AND [T.D] = '&W'"),
                         "[T.C] = 1001 AND [T.D] = '&W'")

    def test_expand_view_vars_text_quoting(self):
        eng = self._var_eng(
            views=[{"name": "v1", "variables": [{"name": "N", "type": "text", "default": "O'Brien"}]}])
        self.assertEqual(eng._expand_view_vars("[T.NM] = &N"), "[T.NM] = 'O''Brien'")

    def test_expand_view_vars_filter_wins_over_default(self):
        dp = SimpleNamespace(var="W", param_refname="", refname="", column="w_code",
                             type="number", default_value="")
        eng = self._var_eng(
            views=[{"name": "v1", "variables": [{"name": "W", "type": "number", "default": "1001"}]}],
            params=[dp], filters=[{"field": "w_code", "op": "equals", "value": "7"}])
        self.assertEqual(eng._expand_view_vars("[T.W_CODE] = &W"), "[T.W_CODE] = 7")

    def test_expand_view_vars_unknown_raises(self):
        eng = self._var_eng(views=[{"name": "v1", "variables": []}])
        with self.assertRaises(ValueError):
            eng._expand_view_vars("[T.C] = &NOPE")

    def test_expand_view_vars_missing_becomes_null(self):
        eng = self._var_eng(
            views=[{"name": "v1",
                    "columns": [{"name": "c", "alias": "c", "expr": "[s1.C] = &W"}],
                    "variables": [{"name": "W", "type": "number", "default": ""}]}])
        self.assertEqual(eng._expand_view_vars("[T.W_CODE] = &W"), "[T.W_CODE] = NULL")
        self.assertEqual(eng._view_vars_missing, {"W"})
        self.assertEqual(eng._view_vars_missing_for([]), {"W"})

    def test_execute_empty_until_var_input(self):
        from rml_python.engine import RMLReportEngine
        eng = RMLReportEngine.__new__(RMLReportEngine)
        eng.metadata = {}
        eng.columns = []
        eng.report_views = [{"name": "v1",
                             "variables": [{"name": "W", "type": "number", "default": ""}]}]
        eng.doc_params = []
        eng.compiler = SimpleNamespace(general_where=lambda: "[T.W_CODE] = &W",
                                       doc_params=lambda: [])
        res = eng.execute({})
        self.assertEqual(res["rows"], [])
        self.assertEqual(res["total"], 0)
        self.assertTrue(res["required_inputs_missing"])
        self.assertEqual(res["awaiting_vars"], ["W"])
        # with a filter value... no linked doc_param and no default → still empty
        res2 = eng.execute({"filters": [{"field": "other", "op": "equals", "value": "1"}]})
        self.assertEqual(res2["rows"], [])

    def test_designer_view_vars_markup(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "..", "rml_python", "shared",
                                "rml_wizard_script.html"), encoding="utf-8") as f:
            js = f.read()
        for needle in ("wbViewVarsHTML", "viewAddVar", "viewDelVar",
                        "wbViewValidateVars", "wbVarPool", "wbVarLinkOptions",
                        "kind:'var'", "خيارات الإدخال"):
            self.assertIn(needle, js)

    def test_preview_modal_on_top(self):
        # الاستعلام التجريبي فوق كل المودالات (استيراد/عرض/مكتبة): عنصر بمستوى
        # الجذر مع z-index صريح (inline) فلا يعتمد على توليد كلاس Tailwind —
        # وإلا سقط خلفها بترتيب DOM أو سياق التكديس
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "..", "rml_python", "shared",
                                "rml_wizard_modal.html"), encoding="utf-8") as f:
            modal = f.read()
        self.assertIn('id="tblPreviewModal"', modal)
        seg = modal.split('id="tblPreviewModal"')[1][:400]
        self.assertIn('z-index:9999', seg)

    def test_sidebar_skips_internal_rml(self):
        import pathlib
        import tempfile
        rep = ('<rml><rpt_metadata name="r" displayName="r"/>'
               '<columns><column id="1" name="d" alias="d" expr="1"/></columns></rml>')
        with tempfile.TemporaryDirectory() as td:
            td = pathlib.Path(td)
            (td / "rep.rml").write_text(rep, encoding="utf-8")
            (td / "_view_library.rml").write_text(
                '<rml><rpt_metadata name="__view_library__" displayName="lib"/>'
                '<views/></rml>', encoding="utf-8")
            (td / "__draft_rep.rml").write_text(rep, encoding="utf-8")
            (td / "__preview_abc.rml").write_text(rep, encoding="utf-8")
            with mock.patch.object(_v, "ws_app_roots", return_value=[td]):
                _meta, _fml, rml_files = _v._load_app_context("demo")
                self.assertEqual([f["file"] for f in rml_files], ["rep.rml"])
                from django.test import RequestFactory
                res = _v.api_app_files(RequestFactory().get("/"), "demo")
                body = json.loads(res.content)
                self.assertEqual([f["file"] for f in body["rml"]], ["rep.rml"])

    def test_designer_library_sql_markup(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "..", "rml_python", "shared",
                                "rml_wizard_script.html"), encoding="utf-8") as f:
            js = f.read()
        for needle in ("wizardViewRefs", "view_refs", "wizardLibViews",
                       "openViewLibrary", "publishViewToLib", "callLibView",
                       "unrefLibView", "refreshLibViews", "deleteLibView",
                       "wbViewSQLLib", "loadLibraryViews", "wbValidViewName",
                       "sqlImportTranslationHTML", "sqlImportCreateView",
                       "sqlImportViewBtn", "استيراد SQL"):
            self.assertIn(needle, js)

    _NESTED_XML = """<rml>
      <rpt_metadata name="vt" displayName="vt" schema=""/>
      <fields>
        <field name="BILL_NO" table_source="BILL_MST" type="TEXT"/>
        <field name="t2" table_source="view_2" type="NUMERIC"/>
        <field name="BILL_NO" table_source="view_2" type="TEXT"/>
      </fields>
      <columns>
        <column id="1" name="t2" alias="t2" expr="[view_2.t2]"/>
      </columns>
      <links>
        <link from_table="BILL_MST" from_col="BILL_NO" to_table="view_2" to_col="BILL_NO" rel_type="one_to_many" match="exact" pattern=""/>
      </links>
      <views>
      <view name="view_1">
        <sources><source id="s1" table_name="BILL_MST" conn=""/></sources>
        <merges/>
        <columns><column name="total" alias="total" expr="SUM([s1.AMT])" type="NUMERIC"/></columns>
        <groups/>
      </view>
      <view name="view_2">
        <sources><source id="a" table_name="view_1" conn=""/></sources>
        <merges/>
        <columns>
          <column name="t2" alias="t2" expr="[a.total]" type="NUMERIC"/>
          <column name="big" alias="big" expr="IF([a.total] &gt; 100, 1, 0)" type="NUMERIC"/>
        </columns>
        <groups/>
      </view>
      </views>
    </rml>"""

    def test_compile_view_sql_nested_flatten(self):
        # Derived source (table = another view): no physical table exists —
        # the inner view SQL is flattened inline as a CTE-like derived table.
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        comp = RMLReportCompiler(xml_text=self._NESTED_XML.replace("&gt;", ">"))
        eng = RMLReportEngine(comp, mock.Mock())
        sql = eng._compile_view_sql(eng._view_map()["VIEW_2"])
        self.assertIn('FROM (SELECT SUM("S1"."AMT") AS "total"', sql)
        self.assertIn(') "A"', sql)
        self.assertIn('"A"."total" AS "t2"', sql)
        # XSQL language still applies through the flattened scope
        self.assertIn("CASE WHEN", sql)
        self.assertNotIn('FROM "view_1"', sql)
        self.assertNotIn("FROM \"VIEW_1\"", sql)

    def test_compile_view_sql_cycle_rejected(self):
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        xml = self._NESTED_XML.replace(
            '<source id="s1" table_name="BILL_MST" conn=""/>',
            '<source id="s1" table_name="view_2" conn=""/>').replace("&gt;", ">")
        eng = RMLReportEngine(RMLReportCompiler(xml_text=xml), mock.Mock())
        with self.assertRaises(ValueError) as ctx:
            eng._compile_view_sql(eng._view_map()["VIEW_2"])
        self.assertIn("دائري", str(ctx.exception))
        # gid walk also terminates on cycles (no RecursionError)
        eng._view_source_gids(eng._view_map()["VIEW_2"])

    def test_build_plan_nested_view_base(self):
        import dataclasses
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        comp = RMLReportCompiler(xml_text=self._NESTED_XML.replace("&gt;", ">"))
        eng = RMLReportEngine(comp, mock.Mock())
        plan = eng._build_plan("view_2", [], None, None, dataclasses)
        self.assertIn("FROM (SELECT", plan["from_q"])
        self.assertTrue(plan["from_q"].strip().endswith('"view_2"'))
        self.assertEqual(plan.get("base_alias"), "view_2")

    _LEG_XML = """<rml>
      <rpt_metadata name="t" displayName="t"/>
      <views><view name="trn">
      <sources>
        <source id="s1" table_name="BILL_MST" conn=""/>
        <source id="s2" table_name="BILL_DTL" conn="" filter="[s2.CC_CODE] IN ('11', '12')"/>
        <source id="s3" table_name="RT_DTL" conn="" filter="[s3.CC_CODE] = '18'"/>
      </sources>
      <merges>
        <merge type="join" source_1_id="s1" source_2_id="s2" join_type="INNER" key_1="BILL_SER" key_2="BILL_SER"/>
        <merge type="union_all" source_1_id="s1" source_2_id="s3"/>
      </merges>
      <columns>
        <column name="CODE" alias="CODE" expr="[s1.CODE]"/>
        <column name="QTY" alias="QTY" expr="NVL([s2.P_QTY], 0)"/>
      </columns><groups/>
      </view></views>
      <columns><column id="1" name="d" alias="d" expr="1"/></columns></rml>"""

    def test_view_source_filter_roundtrip(self):
        import xml.etree.ElementTree as ET
        from rml_python.compiler import RMLReportCompiler
        rml = ET.Element("rml")
        _v._write_views_el(rml, ET, [{"name": "v1",
            "sources": [{"id": "s1", "table_name": "T1", "conn": "",
                         "filter": "[s1.C] IN ('11', '12')"}],
            "merges": [], "columns": [{"name": "a", "alias": "a", "expr": "[s1.A]"}],
            "groups": []}])
        src = rml.find("views/view/sources/source")
        self.assertIsNotNone(src)
        self.assertEqual(src.get("filter"), "[s1.C] IN ('11', '12')")
        comp = RMLReportCompiler(
            xml_text="<rml>" + ET.tostring(rml, encoding="unicode") + "</rml>")
        self.assertEqual(comp.views()[0]["sources"][0]["filter"], "[s1.C] IN ('11', '12')")

    def test_compile_view_sql_leg_filters(self):
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        eng = RMLReportEngine(RMLReportCompiler(xml_text=self._LEG_XML), mock.Mock())
        sql = eng._compile_view_sql(eng._view_map()["TRN"])
        parts = sql.split("UNION ALL")
        self.assertEqual(len(parts), 2)
        self.assertIn('"S2"."CC_CODE" IN (\'11\', \'12\')', parts[0])
        self.assertIn('"S3"."CC_CODE" = \'18\'', parts[1])
        # scoped legs: no foreign tables leak in as CROSS JOINs
        self.assertNotIn("CROSS JOIN", parts[1])
        self.assertNotIn("BILL_MST", parts[1])

    def test_compile_view_sql_leg_scope_violation(self):
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        xml = self._LEG_XML.replace("[s2.CC_CODE] IN ('11', '12')", "[s3.CC_CODE] = '18'")
        eng = RMLReportEngine(RMLReportCompiler(xml_text=xml), mock.Mock())
        with self.assertRaises(ValueError) as ctx:
            eng._compile_view_sql(eng._view_map()["TRN"])
        self.assertIn("خارج فرعه", str(ctx.exception))

    def test_compile_view_sql_leg_bare_ref(self):
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        xml = self._LEG_XML.replace("[s2.CC_CODE] IN ('11', '12')", "[CC_CODE] = '18'")
        eng = RMLReportEngine(RMLReportCompiler(xml_text=xml), mock.Mock())
        with self.assertRaises(ValueError) as ctx:
            eng._compile_view_sql(eng._view_map()["TRN"])
        self.assertIn("[sid.col]", str(ctx.exception))

    def test_compile_view_sql_grouped_union_select(self):
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        xml = """<rml><rpt_metadata name="t" displayName="t"/>
          <views><view name="u">
          <sources><source id="s1" table_name="T1" conn=""/><source id="s2" table_name="T2" conn=""/></sources>
          <merges><merge type="union_all" source_1_id="s1" source_2_id="s2"/></merges>
          <columns><column name="BILL_NO" alias="BILL_NO" expr="[s1.BILL_NO]"/>
          <column name="total" alias="total" expr="SUM([s1.AMT])"/></columns>
          <groups><group name="t" group_by="BILL_NO" expr="SUM([s1.AMT])"/></groups>
          </view></views>
          <columns><column id="1" name="d" alias="d" expr="1"/></columns></rml>"""
        eng = RMLReportEngine(RMLReportCompiler(xml_text=xml), mock.Mock())
        sql = eng._compile_view_sql(eng._view_map()["U"])
        self.assertNotIn("SELECT  FROM", sql)
        self.assertIn('AS "t"', sql)
        self.assertIn("GROUP BY", sql)

    _UNION_XML = """<rml>
      <rpt_metadata name="vt" displayName="vt" schema=""/>
      <fields>
        <field name="BILL_NO" table_source="BILL_MST" type="TEXT"/>
        <field name="total" table_source="union_1" type="NUMERIC"/>
        <field name="BILL_NO" table_source="union_1" type="TEXT"/>
      </fields>
      <columns>
        <column id="1" name="total" alias="total" expr="[union_1.total]"/>
      </columns>
      <links>
        <link from_table="BILL_MST" from_col="BILL_NO" to_table="union_1" to_col="BILL_NO" rel_type="one_to_many" match="exact" pattern=""/>
      </links>
      <views>
      <view name="view_1">
        <sources><source id="s1" table_name="BILL_MST" conn=""/></sources>
        <merges/>
        <columns><column name="total" alias="total" expr="SUM([s1.AMT])" type="NUMERIC"/><column name="BILL_NO" alias="BILL_NO" expr="[s1.BILL_NO]" type="TEXT"/></columns>
        <groups/>
      </view>
      <view name="view_2">
        <sources><source id="s1" table_name="BILL_DTL" conn=""/></sources>
        <merges/>
        <columns><column name="qty" alias="qty" expr="SUM([s1.QTY])" type="NUMERIC"/><column name="BILL_NO" alias="BILL_NO" expr="[s1.BILL_NO]" type="TEXT"/></columns>
        <groups/>
      </view>
      </views>
      <unions><union name="union_1" display="u1">
        <view name="view_1"/><view name="view_2"/>
      </union></unions>
    </rml>"""

    def test_compiler_unions(self):
        from rml_python.compiler import RMLReportCompiler
        comp = RMLReportCompiler(xml_text=self._UNION_XML)
        unions = comp.unions()
        self.assertEqual(len(unions), 1)
        self.assertEqual(unions[0]["name"], "union_1")
        self.assertEqual(unions[0]["views"], ["view_1", "view_2"])
        # union members are REFERENCES — never parsed as (empty) views
        self.assertEqual({v["name"] for v in comp.views()}, {"view_1", "view_2"})

    def test_write_unions_roundtrip(self):
        import xml.etree.ElementTree as ET
        rml = ET.Element("rml")
        _v._write_unions_el(rml, ET, [
            {"name": "union_1", "display": "u1", "views": ["view_1", "view_2"], "imported": True},
            {"name": "thin", "views": ["only_one"]},
            {"name": "  "},
        ])
        els = rml.findall("unions/union")
        self.assertEqual(len(els), 1)
        self.assertEqual(els[0].get("name"), "union_1")
        self.assertEqual(els[0].get("imported"), "1")
        self.assertEqual([x.get("name") for x in els[0].findall("view")], ["view_1", "view_2"])
        from rml_python.compiler import RMLReportCompiler
        back = RMLReportCompiler(
            xml_text="<rml>" + ET.tostring(rml, encoding="unicode") + "</rml>").unions()
        self.assertEqual(len(back), 1)
        self.assertEqual(back[0]["views"], ["view_1", "view_2"])
        self.assertTrue(back[0].get("imported"))

    def test_union_as_view_adapt(self):
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        eng = RMLReportEngine(RMLReportCompiler(xml_text=self._UNION_XML), mock.Mock())
        adapted = eng._view_map()["UNION_1"]
        self.assertEqual([c["alias"] for c in adapted["columns"]], ["total", "BILL_NO", "qty"])
        self.assertEqual(adapted["merges"][0]["type"], "union_all")
        self.assertEqual([s["table_name"] for s in adapted["sources"]], ["view_1", "view_2"])

    def test_compile_union_sql(self):
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        eng = RMLReportEngine(RMLReportCompiler(xml_text=self._UNION_XML), mock.Mock())
        sql = eng._compile_view_sql(eng._view_map()["UNION_1"])
        self.assertIn("UNION ALL", sql)
        self.assertIn('NULL AS "qty"', sql)
        self.assertIn('NULL AS "total"', sql)
        self.assertNotIn('FROM "view_1"', sql)
        self.assertNotIn('FROM "VIEW_1"', sql)

    def test_union_as_view_errors(self):
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        eng = RMLReportEngine(RMLReportCompiler(xml_text=self._UNION_XML), mock.Mock())
        raw = dict(eng._view_map())
        with self.assertRaises(ValueError):
            eng._union_as_view({"name": "u", "views": ["view_1"]}, raw)
        with self.assertRaises(ValueError):
            eng._union_as_view({"name": "u", "views": ["view_1", "nope"]}, raw)
        with self.assertRaises(ValueError):
            eng._union_as_view({"name": "u", "views": ["view_1", "union_1"]}, raw)

    def test_build_plan_union_secondary(self):
        import dataclasses
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        eng = RMLReportEngine(RMLReportCompiler(xml_text=self._UNION_XML), mock.Mock())
        real_tc = RMLReportEngine._table_columns

        def fake_tc(self, table, schema, db=None):
            if self._norm_table(table) in ("VIEW_1", "VIEW_2", "UNION_1"):
                return real_tc(self, table, schema, db)
            return {"BILL_NO": "TEXT", "AMT": "NUMERIC", "QTY": "NUMERIC"}

        eng._table_columns = fake_tc.__get__(eng, RMLReportEngine)
        plan = eng._build_plan("BILL_MST", [], None, None, dataclasses)
        self.assertIn("LEFT JOIN (SELECT", plan["from_q"])
        self.assertIn("UNION ALL", plan["from_q"])
        self.assertIn('"union_1"."BILL_NO" = "T0"."BILL_NO"', plan["from_q"])
        self.assertEqual(plan["table_map"].get("total"), "union_1")
        sql, _ = eng._compile_sql2(plan, [], None, 1, 50, None)
        self.assertIn('SELECT "union_1"."total"', sql)

    def test_designer_union_markup(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "..", "rml_python", "shared",
                                "rml_wizard_script.html"), encoding="utf-8") as f:
            js = f.read()
        for needle in ("wizardUnions", "openViewsBrowser", "closeViewsBrowser",
                       "wbBrowseRender", "wbBrowseDoMerge", "wbBrowseMergeForm",
                       "wbValidUnionName", "wbUnionOutputs", "wbSyncOneUnion",
                       "wbImportUnionSync", "wbImportViewSync", "wbDelUnion",
                       "wbCascadeViewDel", "wbMigrateImported",
                       "openUnionEditModal", "closeUnionEditModal",
                       "wbUnionEditRender", "applyUnionEditModal",
                       "wbUnionRenamePropagate", "تعديل بيانات الدمج",
                       "unions:", "مدموج", "دمج كلي",
                       "استعراض الجداول المعروضة"):
            self.assertIn(needle, js)

    _DOC_EVENT_XML = """<rml>
      <rpt_metadata name="t" displayName="t" schema="" type="doc"/>
      <fields>
        <field name="BILL_NO" table_source="BILL_MST" type="TEXT"/>
        <field name="AMT" table_source="BILL_MST" type="NUMERIC"/>
        <field name="PRICE" table_source="BILL_MST" type="NUMERIC"/>
        <field name="COST" table_source="BILL_MST" type="NUMERIC"/>
      </fields>
      <columns>
        <column id="1" name="BILL_NO" alias="bno" expr="[BILL_MST.BILL_NO]"/>
        <column id="2" name="PX" alias="px" expr="[BILL_MST.PRICE]"/>
      </columns>
      <doc_params>
        <doc_param id="dp1" column="BILL_NO" label="g" type="text" op="group_by" event="group" group_expr="SUM([AMT])">
          <group_var n="1" label="bill" table="BILL_MST" column="BILL_NO"/>
        </doc_param>
        <doc_param id="dp2" column="PX" label="v" type="text" op="var_select" event="variable" target_column="px">
          <var_source n="1" label="cost" expr="[BILL_MST.COST]"/>
          <var_source n="2" label="sale" expr="[BILL_MST.PRICE]"/>
        </doc_param>
        <doc_param id="dp3" column="BILL_NO" label="f" type="text" op="contains" event="filter" src="input"/>
      </doc_params>
    </rml>"""

    def test_compiler_doc_param_events(self):
        from rml_python.compiler import RMLReportCompiler
        comp = RMLReportCompiler(xml_text=self._DOC_EVENT_XML)
        dps = {d.id: d for d in comp.doc_params()}
        self.assertEqual(dps["dp1"].event, "group")
        self.assertEqual(dps["dp1"].group_expr, "SUM([AMT])")
        self.assertEqual(dps["dp1"].group_vars,
                         [{"n": "1", "label": "bill", "table": "BILL_MST", "column": "BILL_NO"}])
        self.assertEqual(dps["dp2"].event, "variable")
        self.assertEqual(dps["dp2"].target_column, "px")
        self.assertEqual(len(dps["dp2"].var_sources), 2)
        self.assertEqual(dps["dp2"].var_sources[1]["label"], "sale")
        self.assertEqual(dps["dp1"].to_dict()["event"], "group")
        # legacy default stays filter
        comp2 = RMLReportCompiler(xml_text=self._DOC_EVENT_XML.replace(' event="filter" src="input"', ""))
        dps2 = {d.id: d for d in comp2.doc_params()}
        self.assertEqual(dps2["dp1"].event, "group")
        self.assertEqual(dps2["dp3"].event, "filter")

    def test_write_doc_params_events_roundtrip(self):
        import xml.etree.ElementTree as ET
        rml = ET.Element("rml")
        self.assertTrue(_v._write_doc_params_el(rml, ET, [{
            "column": "BILL_NO", "label": "g", "event": "group",
            "group_expr": "SUM([AMT])",
            "group_vars": [{"n": "1", "label": "bill", "table": "BILL_MST", "column": "BILL_NO"}],
        }, {
            "column": "PX", "label": "v", "event": "variable", "target_column": "px",
            "var_sources": [{"n": "1", "label": "cost", "expr": "[COST]"}],
        }]))
        from rml_python.compiler import RMLReportCompiler
        back = {d.column: d for d in RMLReportCompiler(
            xml_text="<rml>" + ET.tostring(rml, encoding="unicode") + "</rml>").doc_params()}
        self.assertEqual(back["BILL_NO"].event, "group")
        self.assertEqual(back["BILL_NO"].group_vars[0]["column"], "BILL_NO")
        self.assertEqual(back["PX"].var_sources[0]["expr"], "[COST]")

    def test_where_new_ops(self):
        from rml_python.engine import _build_where
        sql, params = _build_where([{"field": "AMT", "op": "not_between", "valFrom": "1", "valTo": "5"}])
        self.assertIn("NOT BETWEEN", sql)
        sql, _ = _build_where([{"field": "DT", "op": "between_dates", "valFrom": "2026-01-01", "valTo": "2026-01-31"}])
        self.assertIn("BETWEEN", sql)
        sql, _ = _build_where([{"field": "DT", "op": "on_date", "value": "2026-01-15"}])
        self.assertIn("CAST(", sql)
        sql, params = _build_where([{"field": "DT", "op": "in_hour", "value": "9"}])
        self.assertIn("EXTRACT(HOUR FROM", sql)
        self.assertEqual(params.get("p1"), 9)
        sql, params = _build_where([{"field": "DT", "op": "in_minute", "value": "30"}])
        self.assertIn("EXTRACT(MINUTE FROM", sql)
        sql, _ = _build_where([{"field": "PX", "op": "var_select", "value": "sale"}])
        self.assertEqual(sql.strip(), "")
        sql, _ = _build_where([{"field": "BILL_NO", "op": "group_by", "value": "bill"}])
        self.assertEqual(sql.strip(), "")
        with self.assertRaises(ValueError):
            _build_where([{"field": "DT", "op": "in_hour", "value": "99"}])

    def test_doc_events_no_filter_keys(self):
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        eng = RMLReportEngine(RMLReportCompiler(xml_text=self._DOC_EVENT_XML), mock.Mock())
        keys = eng._doc_no_filter_keys()
        self.assertIn("bill_no", keys)
        self.assertIn("px", keys)

    def test_doc_events_compile_sql(self):
        import dataclasses
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        eng = RMLReportEngine(RMLReportCompiler(xml_text=self._DOC_EVENT_XML), mock.Mock())
        real_tc = RMLReportEngine._table_columns
        eng._table_columns = (lambda self, table, schema, db=None: {
            "BILL_NO": "TEXT", "AMT": "NUMERIC", "PRICE": "NUMERIC", "COST": "NUMERIC"}
            ).__get__(eng, RMLReportEngine)
        plan = eng._build_plan("BILL_MST", [], None, None, dataclasses)
        filts = [{"field": "BILL_NO", "refname": "g", "op": "group_by", "value": "bill"},
                 {"field": "PX", "refname": "v", "op": "var_select", "value": "sale"}]
        sql, params = eng._compile_sql2(plan, filts, None, 1, 50, None)
        self.assertIn("OVER (PARTITION BY", sql)
        self.assertIn("GROUP BY", sql)
        self.assertIn("DECODE(", sql)  # oracle dialect for the Mock connection
        self.assertNotIn("WHERE", sql)
        # unknown choice fails fast with guidance
        with self.assertRaises(ValueError):
            eng._compile_sql2(plan, [{"field": "PX", "refname": "v", "op": "var_select", "value": "nope"}],
                              None, 1, 50, None)

    def test_doc_variable_case_dialect(self):
        from types import SimpleNamespace as _NS
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine
        eng = RMLReportEngine(RMLReportCompiler(xml_text=self._DOC_EVENT_XML), mock.Mock())
        cols, groups, extra = eng._apply_doc_param_events(
            {"base_norm": "BILL_MST", "local_sec": []},
            [_NS(alias="px", name="PX", expr="[BILL_MST.PRICE]")],
            [{"field": "PX", "refname": "v", "op": "var_select", "value": "cost"}],
            {}, "pg")
        # group event falls back to its first variable when no value is sent
        self.assertEqual(groups, ['"BILL_NO"'])
        self.assertEqual(len(extra), 1)
        self.assertIn("PARTITION BY", extra[0][0])
        self.assertEqual(len(cols), 1)
        self.assertIn("CASE WHEN", cols[0].expr)
        self.assertNotIn("DECODE", cols[0].expr)

    def test_designer_doc_events_markup(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "..", "rml_python", "shared",
                                "rml_wizard_script.html"), encoding="utf-8") as f:
            js = f.read()
        for needle in ("dpEventOf", "dpSrcOf", "DP_OPS", "dpSetEvent", "dpSetSrc",
                       "dpEventBodyHTML", "dpAddGroupVar", "dpDelGroupVar",
                       "dpAddVarSource", "dpDelVarSource", "dpTestSrcValues",
                       "setRptPrevFilterRange", "dpPrevDynLoad", "dpPrevDefault",
                       "حدث المدخل", "متغيرات التجميع", "مصادر العمود",
                       "العمود المتأثر", "not_between", "between_dates",
                       "in_hour", "in_minute", "on_date",
                       "فرع المدخلات المطلوبة (كل أنواع التقارير)"):
            self.assertIn(needle, js)

    def test_uniexpression_module(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "..", "rml_python", "shared",
                                "UniExpression.js"), encoding="utf-8") as f:
            js = f.read()
        # main UI + tree modal + catalog + parser
        for needle in ("UniExpression", "fieldHTML", "openForId", "openTree",
                       "closeTree", "applyTree", "refreshCatalog",
                        "محرر التعبير الشجري",
                       "parseExpr", "splitTop", "UE_SQL", "+f",
                       "dbSearchOpen", "ueTreeModal",
                       "exprAuto", "_ueFieldsProvider", "ueFallbackComplete"):
            self.assertIn(needle, js)
        self.assertEqual(js.count("{"), js.count("}"))
        self.assertEqual(js.count("("), js.count(")"))
        # wizard + forms call it (no local duplicates of the component)
        with open(os.path.join(base, "..", "rml_python", "shared",
                                "rml_wizard_script.html"), encoding="utf-8") as f:
            wz = f.read()
        self.assertNotIn("function openTree(", wz)
        for caller in ("UniExpression.fieldHTML", "UniExpression.openForId",
                       "uexpBtn('wxe_'", "uexpBtn('wxw_'", "uexpBtn('wxed_'",
                       "openForId(\\'vwColExpr_'",
                       "UniExpression.openForId('dpGExpr_${i}')"):
            self.assertIn(caller, wz)
        with open(os.path.join(base, "..", "fmlk_engine", "shared",
                                "forms_wizard_script.html"), encoding="utf-8") as f:
            fw = f.read()
        self.assertIn("UniExpression.openForId('fwExpr_${i}')", fw)
        # included before the wizards that call it
        for tpl, wiz in (("report_designer.html", "rml_wizard_script.html"),
                         ("forms_designer.html", "forms_wizard_script.html")):
            with open(os.path.join(base, "templates", tpl), encoding="utf-8") as f:
                html = f.read()
            self.assertIn('{% include "UniExpression.js" %}', html)
            wiz_inc = '{% include "' + wiz + '" %}'
            self.assertLess(html.index('{% include "UniExpression.js" %}'),
                            html.index(wiz_inc))

    def test_unidbpicker_module(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "..", "rml_python", "shared",
                                "UniDBPicker.js"), encoding="utf-8") as f:
            js = f.read()
        # unified pattern: select menu + search button + search modal
        for needle in ("UniDBPicker", "dbSearchOpen", "dbSearchClose",
                       "dbSearchFilter", "dbSearchPick", "dbPickSetSelect",
                       "dbPickBtnHTML", "dbPickWrapHTML",
                       "pickTable", "pickColumn",
                       "fillTableSelect", "fillColumnSelect", "selOptions"):
            self.assertIn(needle, js)
        # balanced braces sanity
        self.assertEqual(js.count("{"), js.count("}"))
        self.assertEqual(js.count("("), js.count(")"))
        # wizard calls it (no local duplicates of the moved core)
        with open(os.path.join(base, "..", "rml_python", "shared",
                                "rml_wizard_script.html"), encoding="utf-8") as f:
            wz = f.read()
        self.assertNotIn("function dbSearchOpen(", wz)
        self.assertNotIn("function dbPickSetSelect(", wz)
        for caller in ("UniDBPicker.pickTable", "UniDBPicker.pickColumn",
                       "UniDBPicker.fillTableSelect", "UniDBPicker.fillColumnSelect",
                       "UniDBPicker.selOptions", "vwTableSearch", "importTableSearch",
                       "importColSearch", "fldTableSearch", "fldColSearch",
                       "unionTableSearch", "wbLinkSearchTbl", "wbLinkSearchFld",
                       "dpGroupTblSearch", "dpGroupColSearch"):
            self.assertIn(caller, wz)
        # included before the wizard in every host template
        for tpl in ("report_designer.html", "report_player.html", "app_detail.html"):
            with open(os.path.join(base, "templates", tpl), encoding="utf-8") as f:
                html = f.read()
            self.assertIn('{% include "UniDBPicker.js" %}', html)
            self.assertLess(html.index('{% include "UniDBPicker.js" %}'),
                            html.index('{% include "rml_wizard_script.html" %}'))


class ImportMultiConnTests(SimpleTestCase):
    """Import modal keeps selections across connections (multi-DB report)."""

    def _wiz(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "..", "rml_python", "shared",
                                "rml_wizard_script.html"), encoding="utf-8") as f:
            return f.read()

    def test_import_keeps_tables_across_connections(self):
        js = self._wiz()
        # composite entry key conn||table
        for needle in ("function importEntryKey(conn, sh)",
                       "renderImportConnSel", "importConnSelCounts",
                       "importToggleTable(full, on, key)",
                       "previewImportTable(onlyKey)",
                       "importUnionSel[key]",
                       "importUnionConnsHTML(key, e)"):
            self.assertIn(needle, js)
        # switching the connection must NOT wipe previously added tables
        start = js.index("async function loadImportTables(){")
        end = js.index("/* ── Per-table General Condition Builder", start)
        body = js[start:end]
        self.assertNotIn("importMulti = {}", body)
        self.assertNotIn("importActiveTbl = ''", body)
        self.assertIn("تبقى كما هي", body)
        # modal open still starts fresh (incl. stale UNION picks)
        self.assertIn("importMulti = {}; importActiveTbl = ''; importMergeMode = 'links'; importUnionSel = {};", js)

class HideEmptyRowsTests(SimpleTestCase):
    """Column hide_blank/hide_zero hide the RECORD (row filter), not just the cell."""

    _XML = """<rml><rpt_metadata name="t" displayName="t"/>
      <rml_connections><connection id="1" connection_id="A"/></rml_connections>
      <fields>
        <field name="CNAME" table_source="T1" conn_id="1" type="TEXT"/>
        <field name="AMT" table_source="T1" conn_id="1" type="NUMERIC"/>
      </fields>
      <columns>
        <column id="1" name="cn" alias="cn" expr="[T1.CNAME]" hide_blank="1"/>
        <column id="2" name="amt" alias="amt" expr="[T1.AMT]" hide_zero="1"/>
        <column id="3" name="tot" alias="tot" expr="SUM([T1.AMT])" hide_zero="1"/>
        <column id="4" name="amt_nb" alias="amt_nb" expr="[T1.AMT]" hide_blank="1"/>
      </columns>
      <table_opts><table name="T1" conn="1"/></table_opts>
      </rml>"""

    class _Cur:
        def __init__(self, desc, rows):
            self.description = [(c,) for c in desc]
            self._rows = rows

        def fetchall(self):
            return list(self._rows)

        def fetchone(self):
            return self._rows[0] if self._rows else None

        def close(self):
            pass

    class _DB:
        def __init__(self, handler):
            self._handler = handler
            self.conn = True

        def connect(self):
            pass

        def _exec(self, sql, params):
            desc, rows = self._handler(sql, params)
            return HideEmptyRowsTests._Cur(desc, rows)

    def _eng(self):
        from rml_python.compiler import RMLReportCompiler
        from rml_python.engine import RMLReportEngine

        def _h(sql, params):
            return (["CNAME", "AMT"], [("x", 5)])

        comp = RMLReportCompiler(xml_text=self._XML)
        return RMLReportEngine(comp, self._DB(_h), databases={"A": self._DB(_h)})

    def test_parse_and_dict_flags(self):
        from rml_python.compiler import RMLReportCompiler
        comp = RMLReportCompiler(xml_text=self._XML)
        cols = comp.columns()
        self.assertTrue(cols[0].hide_blank)
        self.assertFalse(cols[0].hide_zero)
        self.assertTrue(cols[1].hide_zero)
        self.assertTrue(cols[3].hide_blank)
        d = cols[0].to_dict()
        self.assertTrue(d["hide_blank"] and d["hideBlank"])
        d2 = cols[1].to_dict()
        self.assertTrue(d2["hide_zero"] and d2["hideZero"])

    def test_hide_filters_synthesized(self):
        eng = self._eng()
        syn = eng._hide_empty_filters()
        self.assertEqual([(s["field"], s["op"]) for s in syn],
                         [("cn", "not_blank"), ("amt", "not_zero"),
                          ("tot", "not_zero"), ("amt_nb", "not_blank")])

    def test_preview_sql_hides_rows(self):
        eng = self._eng()
        sql = eng.preview_sql({"activeTable": "T1", "page": 1, "pageSize": 50})
        # blank on TEXT (sentinel guard, no '' literal — Oracle treats '' as NULL!):
        # NULL or spaces-only hidden, real values kept
        self.assertIn("CONCAT(LTRIM(RTRIM(CAST(", sql)
        self.assertIn("<> '|')", sql)
        # blank on NUMERIC: cheap IS NOT NULL (no functions, index-friendly)
        self.assertIn('("AMT" IS NOT NULL)', sql)
        # zero on NUMERIC field: exact numeric compare (covers fractions)
        self.assertIn("<> 0", sql)
        # aggregated column routes to the outer alias wrapper
        self.assertIn("SELECT * FROM (", sql)
        # regression guard: '' literal comparison hides EVERYTHING on Oracle
        self.assertNotIn("<> ''", sql)
        self.assertNotIn("NOT IN ('', '0')", sql)

    def test_wizard_markup_has_hide_options(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "..", "rml_python", "shared",
                                "rml_wizard_script.html"), encoding="utf-8") as f:
            js = f.read()
        for needle in ("hide_blank:!!(c.hide_blank||c.hideBlank)",
                       "hide_zero:!!(c.hide_zero||c.hideZero)",
                       "updateWizardColumn(${i},'hide_blank',this.checked)",
                       "updateWizardColumn(${i},'hide_zero',this.checked)",
                       "إخفاء السجل الفارغ", "إخفاء السجل الصفري"):
            self.assertIn(needle, js)

class ChartsTypesTests(SimpleTestCase):
    """10 chart families: designer schema + persistence round-trip."""

    _XML = """<rml><rpt_metadata name="t" displayName="t"/>
      <rml_chart>
        <chart id="1" type="stacked" title="s" category="branch" value="amt" series="mon" stacking="stacked"/>
        <chart id="2" type="gauge" title="g" value="rate" min="0" max="100" target="tgt" thresholds="50,80"/>
        <chart id="3" type="treemap" title="m" category="sec" parent="div" value="amt" color="grp"/>
        <chart id="4" type="bar" title="legacy" x="branch" y="amt" y2="mon"/>
      </rml_chart></rml>"""

    def test_parse_all_slots(self):
        from rml_python.compiler import RMLReportCompiler
        comp = RMLReportCompiler(xml_text=self._XML)
        ch = {c.id: c for c in comp.charts()}
        self.assertEqual((ch["1"].category, ch["1"].value, ch["1"].series, ch["1"].stacking),
                         ("branch", "amt", "mon", "stacked"))
        self.assertEqual((ch["2"].value, ch["2"].min, ch["2"].max, ch["2"].target, ch["2"].thresholds),
                         ("rate", "0", "100", "tgt", "50,80"))
        self.assertEqual((ch["3"].category, ch["3"].parent, ch["3"].value, ch["3"].color),
                         ("sec", "div", "amt", "grp"))
        # legacy x/y/y2 readable as category/value/series
        self.assertEqual((ch["4"].category, ch["4"].value, ch["4"].series),
                         ("branch", "amt", "mon"))
        d = ch["1"].to_dict()
        for k in ("category", "value", "series", "size", "color", "parent", "min",
                  "max", "target", "thresholds", "stacking", "totals"):
            self.assertIn(k, d)

    def test_save_roundtrip(self):
        from urs.views import _render_rml_xml
        from rml_python.compiler import RMLReportCompiler
        charts = [
            {"id": "1", "type": "stacked", "title": "s", "category": "branch",
             "value": "amt", "series": "mon", "stacking": "percent"},
            {"id": "2", "type": "waterfall", "title": "w", "category": "step",
             "value": "amt", "totals": "last"},
            {"id": "3", "type": "bubble", "title": "b", "x": "price",
             "y": "qty", "size": "sales", "series": "grp"},
        ]
        xml = _render_rml_xml("t", "t", "fa-chart-bar", "HR", "S", "", None,
                              [], [], [], charts, [], report_type="master")
        comp = RMLReportCompiler(xml_text=xml)
        got = {c.id: c for c in comp.charts()}
        self.assertEqual(len(got), 3)
        self.assertEqual((got["1"].stacking, got["2"].totals, got["3"].size),
                         ("percent", "last", "sales"))
        self.assertEqual(got["3"].x, "price")

    def test_designer_schema_markup(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "..", "rml_python", "shared",
                                "rml_wizard_script.html"), encoding="utf-8") as f:
            js = f.read()
        for needle in ("WB_CHART_TYPES", "WB_CHART_SCHEMA", "WB_CHART_KIND",
                       "WB_CHART_SELS", "function chSlot(", "function wbChartSummary(",
                       "wbChartTypeLabel", "function renderWizardCharts()",
                       "function wbChartCardHTML(", "function wbChartsRerender()",
                       "function wbChartSrcOpts()", "function wbChartFldInput(",
                       "onclick=\"addWizardChart()\""):
            self.assertIn(needle, js)
        # sources are imported table columns (fields), grouped per table —
        # not final report columns
        for needle in ("(wizardFields || [])", "optgroup", "table_source",
                       "اختر العمود المستورد", "(سابق)"):
            self.assertIn(needle, js)
        self.assertNotIn("function wbChartAliasOpts()", js)
        for t in ("'bar'", "'hbar'", "'line'", "'pie'", "'doughnut'", "'stacked'",
                  "'grouped'", "'scatter'", "'bubble'", "'area'", "'gauge'", "'kpi'",
                  "'waterfall'", "'treemap'"):
            self.assertIn(t, js)
        for slot in ("category',", "series',", "stacking',", "thresholds',", "totals',"):
            self.assertIn(slot, js)

    def test_chart_sources_are_imported_fields(self):
        # functional: dropdown lists imported table columns grouped per table
        # (qualified when duplicated), never final report columns.
        import shutil
        import subprocess
        import tempfile
        base = os.path.dirname(__file__)
        src = os.path.join(base, "..", "rml_python", "shared", "rml_wizard_script.html")
        if not shutil.which("node"):
            self.skipTest("node missing")
        with open(src, encoding="utf-8") as f:
            js = f.read()
        start = js.index("function wbChartSrcOpts()")
        brace = js.index("{", start)
        depth = 0
        end = brace
        for k in range(brace, len(js)):
            if js[k] == "{":
                depth += 1
            elif js[k] == "}":
                depth -= 1
                if not depth:
                    end = k + 1
                    break
        runner = (
            "var wizardFields=[{name:'BILL_NO',table_source:'T1',conn_id:'1'},"
            "{name:'AMT',table_source:'T1',conn_id:'1'},"
            "{name:'AMT',table_source:'T2',conn_id:'2'},"
            "{name:'CNAME',table_source:'T2',conn_id:'2'}];"
            "var wizardConnNames={'1':'L','2':'R'};"
            "function wbEsc(s){return String(s==null?'':s);}"
            + js[start:end] +
            "var html=wbChartSrcOpts();"
            "var assert=require('assert');"
            "assert(html.indexOf('<optgroup')>=0,'optgroups');"
            "assert(html.indexOf('T1.AMT')>=0 && html.indexOf('T2.AMT')>=0,'dup qualified');"
            "assert(html.indexOf('value=\"BILL_NO\"')>=0,'bare unique');"
            "assert(html.indexOf('wizardColumns')<0,'no final columns');"
            "console.log('SRC-OPTS OK');"
        )
        with tempfile.NamedTemporaryFile("w", suffix=".js", delete=False, encoding="utf-8") as t:
            t.write(runner)
            tmp = t.name
        try:
            r = subprocess.run(["node", tmp], capture_output=True, text=True)
        finally:
            try:
                os.unlink(tmp)
            except OSError:
                pass
        self.assertEqual(r.returncode, 0, (r.stderr or "")[-2000:])
        self.assertIn("SRC-OPTS OK", r.stdout)

    def test_overview_never_shows_none(self):
        # Django prints Python None as "None" — chart headers/JS must fall back
        # to category/value/series (new schema) then legacy x/y/y2, never None.
        from django.template import Context, Template
        hdr = ('{{ ch.category|default:ch.x|default:"—" }} ← '
               '{{ ch.value|default:ch.y|default:"—" }}'
               '{% if ch.series|default:ch.y2 %} + {{ ch.series|default:ch.y2 }}{% endif %}')
        out = Template(hdr).render(Context({"ch": {"category": None, "value": None,
                                                    "series": None, "x": None,
                                                    "y": None, "y2": None}}))
        self.assertEqual(out, "— ← —")
        self.assertNotIn("None", out)
        out2 = Template(hdr).render(Context({"ch": {"category": "BILL_DATE",
                                                     "value": "SALES_VAL",
                                                     "series": "I_CODE",
                                                     "x": None, "y": None,
                                                     "y2": None}}))
        self.assertEqual(out2, "BILL_DATE ← SALES_VAL + I_CODE")
        js = ('x:"{{ ch.category|default:ch.x|default:\'\'|escapejs }}"')
        out3 = Template(js).render(Context({"ch": {"category": None, "x": None}}))
        self.assertEqual(out3, 'x:""')
        self.assertNotIn("None", out3)

    def test_overview_templates_use_fallbacks(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "templates", "app_detail.html"), encoding="utf-8") as f:
            ad = f.read()
        self.assertIn("ch.category|default:ch.x", ad)
        self.assertIn("ch.value|default:ch.y", ad)
        self.assertIn("ch.series|default:ch.y2", ad)
        self.assertNotIn("{{ ch.x }} ← {{ ch.y }}", ad)
        with open(os.path.join(base, "templates", "my_reports.html"), encoding="utf-8") as f:
            mr = f.read()
        self.assertIn("ch.category || ch.x", mr)
        self.assertIn("column: _cx", mr)

    def test_charts_foot_and_sheet_buttons_visible(self):
        # root cause of the missing Add button: foot kept Tailwind `hidden`
        # while code only cleared inline style — assert the class toggle fix
        # plus an inline Add button inside the charts sheet itself.
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "..", "rml_python", "shared",
                                "rml_wizard_script.html"), encoding="utf-8") as f:
            js = f.read()
        self.assertIn("fbar.classList.toggle('hidden', !foot)", js)
        self.assertNotIn("fbar.style.display = foot ? '' : 'none'", js)
        start = js.index("else if(wbSheet==='charts'){")
        end = js.index("else if(wbSheet==='docs'){", start)
        body = js[start:end]
        self.assertIn("addWizardChart()", body)
        self.assertIn("wbChartCardHTML(ch, i, _chOpts)", body)

class FmlkGridSrcTests(SimpleTestCase):
    """مصدر الإكمال المشروط لأعمدة شبكة التفاصيل (جدول تعبيري + رقم/اسم)."""

    _XML = """<fml><fml_metadata name="t" displayName="t" table="T0"/>
      <details><detail table="LN" master="id" detail="pid" sub_key="id">
        <columns>
          <column name="item_no" alias="الصنف" src_expr="if(@doc_type='items',[IAS_ITM_MST],[SUP_MST])" key_col="NO" label_col="NAME" src_conn="5"/>
          <column name="qty" alias="الكمية"/>
        </columns>
      </detail></details></fml>"""

    def test_detail_col_parse(self):
        from fmlk_engine.compiler import FMLKFormCompiler
        comp = FMLKFormCompiler(xml_text=self._XML)
        cols = comp.details()[0].columns
        c = {x["name"]: x for x in cols}["item_no"]
        self.assertEqual(c["src_expr"], "if(@doc_type='items',[IAS_ITM_MST],[SUP_MST])")
        self.assertEqual(c["key_col"], "NO")
        self.assertEqual(c["label_col"], "NAME")
        self.assertEqual(c["src_conn"], "5")
        q = {x["name"]: x for x in cols}["qty"]
        self.assertEqual(q["src_expr"], "")

    def test_save_persists_grid_src_keys(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "views.py"), encoding="utf-8") as fh:
            vw = fh.read()
        for k in ('"src_expr"', '"key_col"', '"label_col"', '"src_conn"'):
            self.assertIn(k, vw)
        self.assertIn("def api_fmlk_grid_source", vw)
        self.assertIn("def _fmlk_grid_source_external", vw)

    def test_grid_source_validation(self):
        from django.test import RequestFactory
        from urs import views as _v
        rf = RequestFactory()
        r = _v.api_fmlk_grid_source(rf.get("/api/fmlk/grid-source"))
        self.assertEqual(r.status_code, 400)
        r = _v.api_fmlk_grid_source(rf.get("/api/fmlk/grid-source",
                                           {"table": "T;DROP", "key": "ID"}))
        self.assertEqual(r.status_code, 400)
        r = _v.api_fmlk_grid_source(rf.get("/api/fmlk/grid-source",
                                           {"table": "T", "key": "ID"}))
        self.assertEqual(r.status_code, 400)

    def test_designer_grid_src_markup(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "..", "fmlk_engine", "shared",
                                "forms_wizard_script.html"), encoding="utf-8") as fh:
            js = fh.read()
        for needle in ("fwBranchSrcDraft", "fwBranchSrcToggle", "fwBranchSrcSet",
                       "fwBranchSrcIns", "data-bsrc-expr", "مصدر الإكمال المشروط",
                       "الرقم: تلقائي", "_sxOf"):
            self.assertIn(needle, js)

    def test_player_grid_src_markup(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "templates", "forms_player.html"), encoding="utf-8") as fh:
            js = fh.read()
        for needle in ("fwDetailACSrc", "api/fmlk/grid-source", "a.sk", "skOv",
                       "الرقم المرجعي + الاسم"):
            self.assertIn(needle, js)

    def test_player_grid_ad_ws_pin(self):        # 'ad:forms@ws' → kind + &ws= (اتصال ثابت لشجرة مساحة محددة)
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "templates", "forms_player.html"), encoding="utf-8") as fh:
            js = fh.read()
        for needle in ("A-Za-z0-9_:@", "kind.indexOf('@')", "&ws=",
                       "ad:forms@workspace_1"):
            self.assertIn(needle, js)
        with open(os.path.join(base, "..", "fmlk_engine", "shared",
                                "forms_wizard_script.html"), encoding="utf-8") as fh:
            djs = fh.read()
        self.assertIn("ad:forms@workspace_1", djs)

    def test_player_download_tree_mode(self):
        # مودال الإنزال للأعمدة ad: يبحث في الشجرة (+@ws) لا الجدول الفارغ
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "templates", "forms_player.html"), encoding="utf-8") as fh:
            js = fh.read()
        for needle in ("fwDlState.ad", "api/ad/search/?kind=", "_ad_name",
                       "st.ad.ws", "st.ad.kind", "المتاحة وإضافتها",
                       "direct_tables", "(ad|src)"):
            self.assertIn(needle, js)

    def test_designer_dyn_tab_from_source(self):
        # زر التبويب الديناميكي: اختيار مصدر الشجرة + توليد الفرع (+default)
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "..", "fmlk_engine", "shared",
                                "forms_wizard_script.html"), encoding="utf-8") as fh:
            js = fh.read()
        for needle in ("fwDynTabOpen", "fwDynTabClose", "fwDynTabCreate",
                       "fwDynTabSrcChanged", "fwDynSrcList", "api/ad/sources/",
                       "تبويب ديناميكي", "ad_perm_'", "ref_key", "src_expr: s.lit",
                       "fwIsPermFormDs", "permissions",
                       "default: (old.default", "default: (c.default"):
            self.assertIn(needle, js)


class MigrateDirectDetailsTests(SimpleTestCase):
    """ترحيل النموذج ينشئ جداوله الفرعية المباشرة (كانت تُفترض موجودة)."""

    def _fmlk(self):
        import tempfile
        from pathlib import Path as _P
        td = _P(tempfile.mkdtemp(prefix="migdt_"))
        self.addCleanup(lambda: __import__("shutil").rmtree(str(td), ignore_errors=True))
        p = td / "perm.fmlk"
        p.write_text(
            '<?xml version="1.0" encoding="utf-8"?><fml>'
            '<fml_metadata name="perm" displayName="صلاحيات" table="perm_t" '
            'model_type="form" schema="public" connection="urs_local"/>'
            '<fields><field id="1" name="code" alias="الكود" inputType="text" dataType="VARCHAR"/>'
            '</fields>'
            '<details><detail table="perm_rows" alias="الصفوف" master="id" detail="perm_id" '
            'link_mode="direct" sub_key="id" visible="1">'
            '<columns>'
            '<column name="form_name" alias="النموذج" data_type="VARCHAR" input_type="text"/>'
            '<column name="can_view" alias="عرض" data_type="BOOLEAN" input_type="boolean"/>'
            '</columns></detail></details></fml>', encoding="utf-8")
        return p

    def _fake_pg(self):
        import sys as _sys
        import types as _t

        class _Cur:
            def __init__(self):
                self.executed = []
            def execute(self, sql, params=None):
                self.executed.append(str(sql))
            def fetchall(self):
                return []
            def fetchone(self):
                return None
            def close(self):
                pass

        class _Conn:
            def __init__(self):
                self.cur = _Cur()
                self.autocommit = True
            def cursor(self):
                return self.cur
            def close(self):
                pass
            def rollback(self):
                pass

        holder = {}
        mod = _t.ModuleType("psycopg2")

        def _connect(**kw):
            holder["conn"] = _Conn()
            return holder["conn"]
        mod.connect = _connect
        _old = _sys.modules.get("psycopg2")
        _sys.modules["psycopg2"] = mod
        self.addCleanup(lambda: (_sys.modules.pop("psycopg2", None),
                                 _sys.modules.__setitem__("psycopg2", _old) if _old else None))
        return holder

    def test_migrate_creates_direct_detail_table(self):
        import urs.models as _m
        p = self._fmlk()
        holder = self._fake_pg()
        rf = RequestFactory()
        with mock.patch.object(_v, "_find_fml_path", return_value=p), \
             mock.patch.object(_m.Connection.objects, "filter") as _f:
            _f.return_value.first.return_value = None
            req = rf.post("/api/apps/s/models/migrate/", data=json.dumps({"file": "perm.fmlk"}),
                          content_type="application/json")
            res = _v.api_models_migrate(req, app_name="s")
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        dts = body.get("direct_tables") or []
        self.assertEqual(len(dts), 1, body)
        self.assertEqual(dts[0]["detail"], "perm_rows")
        self.assertTrue(dts[0]["created"], dts)
        sql = "\n".join(holder["conn"].cur.executed)
        self.assertIn('CREATE TABLE IF NOT EXISTS "public"."perm_rows"', sql)
        self.assertIn('"id" SERIAL PRIMARY KEY', sql)
        self.assertIn('"perm_id" INTEGER', sql)
        self.assertIn('"form_name" VARCHAR(255)', sql)
        self.assertIn('"can_view" BOOLEAN', sql)

    def test_migrate_preview_lists_direct_tables(self):
        import urs.models as _m
        p = self._fmlk()
        rf = RequestFactory()
        with mock.patch.object(_v, "_find_fml_path", return_value=p), \
             mock.patch.object(_m.Connection.objects, "filter") as _f:
            _f.return_value.first.return_value = None
            req = rf.post("/api/apps/s/models/migrate/preview/",
                          data=json.dumps({"file": "perm.fmlk"}),
                          content_type="application/json")
            res = _v.api_models_migrate_preview(req, app_name="s")
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        dts = body.get("direct_tables") or []
        self.assertEqual(len(dts), 1, body)
        self.assertEqual(dts[0]["detail"], "perm_rows")
        self.assertIn("form_name", dts[0]["columns"])

    def test_migrate_direct_detail_default_clause(self):
        import tempfile
        from pathlib import Path as _P
        import urs.models as _m
        td = _P(tempfile.mkdtemp(prefix="migdf_"))
        self.addCleanup(lambda: __import__("shutil").rmtree(str(td), ignore_errors=True))
        p = td / "perm.fmlk"
        p.write_text(
            '<?xml version="1.0" encoding="utf-8"?><fml>'
            '<fml_metadata name="perm" displayName="صلاحيات" table="perm_t" '
            'model_type="form" schema="public" connection="urs_local"/>'
            '<fields><field id="1" name="code" alias="الكود" inputType="text" dataType="VARCHAR"/>'
            '</fields>'
            '<details><detail table="perm_rows" alias="الصفوف" master="id" detail="perm_id" '
            'link_mode="direct" sub_key="id" visible="1">'
            '<columns>'
            '<column name="source" alias="المصدر" data_type="VARCHAR" default="src:accounts"/>'
            '<column name="ref_key" alias="المرجع" data_type="VARCHAR"/>'
            '</columns></detail></details></fml>', encoding="utf-8")
        holder = self._fake_pg()
        rf = RequestFactory()
        with mock.patch.object(_v, "_find_fml_path", return_value=p), \
             mock.patch.object(_m.Connection.objects, "filter") as _f:
            _f.return_value.first.return_value = None
            req = rf.post("/api/apps/s/models/migrate/", data=json.dumps({"file": "perm.fmlk"}),
                          content_type="application/json")
            res = _v.api_models_migrate(req, app_name="s")
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        sql = "\n".join(holder["conn"].cur.executed)
        self.assertIn('"source" VARCHAR(255) DEFAULT \'src:accounts\'', sql)


class AdSystemTests(SimpleTestCase):
    """ad_system: tree scan, sources CRUD, search shaping, endpoints."""

    def _wsdir(self):
        import tempfile
        from pathlib import Path as _P
        td = _P(tempfile.mkdtemp(prefix="adws_"))
        self.addCleanup(lambda: __import__("shutil").rmtree(str(td), ignore_errors=True))
        app = td / "apps" / "app1"
        app.mkdir(parents=True)
        (app / "t.fmlk").write_text(
            '<?xml version="1.0" encoding="utf-8"?><fml>'
            '<fml_metadata name="t" displayName="نموذج ت" table="t_tbl"/>'
            '<fields>'
            '<field id="1" name="code" alias="الرمز" inputType="text" dataType="VARCHAR"/>'
            '<field id="2" name="qty" alias="الكمية" inputType="number" dataType="INTEGER" serial="true"/>'
            '</fields></fml>', encoding="utf-8")
        (app / "r.rml").write_text(
            '<?xml version="1.0" encoding="utf-8"?><rml>'
            '<rpt_metadata name="r" displayName="تقرير ر"/>'
            '<fields><field id="1" name="C1" table_source="T1" conn_id="1" type="TEXT"/></fields>'
            '<columns><column id="1" name="c1" alias="ع1" expr="[T1.C1]"/></columns>'
            '<rules><rule id="1" name="disc"><variables>'
            '<variable name="pct" display="P" type="number"/>'
            '</variables></rule></rules></rml>', encoding="utf-8")
        return td

    def test_tree_scan(self):
        from ad_system import engine as _ad
        t = _ad.tree(self._wsdir())
        self.assertEqual(len(t["forms"]), 1)
        self.assertEqual(t["forms"][0]["name_ar"], "نموذج ت")
        self.assertEqual(len(t["forms"][0]["inputs"]), 2)
        self.assertEqual(t["forms"][0]["serials"], ["qty"])
        self.assertEqual(len(t["reports"]), 1)
        self.assertEqual(t["reports"][0]["columns"], [{"name": "c1", "alias": "ع1"}])
        self.assertEqual(len(t["rules"]), 1)
        self.assertEqual(t["rules"][0]["variables"], ["pct"])
        self.assertEqual(len(t["sequences"]), 1)
        self.assertTrue(any(s["key"] == "forms" for s in t["sources"]))

    def test_sources_crud(self):
        from ad_system import engine as _ad
        wd = self._wsdir()
        self.assertEqual(_ad.read_sources(wd), [])
        s = {"key": "items", "ar": "الأصناف", "kind": "table", "conn": "5",
             "table": "ITM", "schema": "", "seq_col": "NO", "name_col": "NAME"}
        self.assertEqual(_ad.validate_source(s), "")
        bad = dict(s, key="9x")
        self.assertTrue(_ad.validate_source(bad))
        _ad.write_sources(wd, [s])
        got = _ad.read_sources(wd)
        self.assertEqual(len(got), 1)
        self.assertEqual(_ad.all_sources(wd)[-1]["key"], "items")

    def test_search_rows(self):
        from ad_system import engine as _ad
        t = _ad.tree(self._wsdir())
        rows = _ad.search_rows("forms", t, "نموذج")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["form_file"], "t.fmlk")
        self.assertEqual(rows[0]["ref_name"], "نموذج ت")
        self.assertEqual(_ad.search_rows("forms", t, "zzz"), [])
        rr = _ad.search_rows("rules", t, "disc")
        self.assertEqual(rr[0]["rule_name"], "disc")
        self.assertEqual(_ad.form_inputs_lookup(t, "t.fmlk")[0]["name"], "code")
        self.assertEqual(_ad.report_columns_lookup(t, "r.rml")[0]["name"], "c1")
        self.assertEqual(_ad.rule_variables_lookup(t, "r.rml", "disc"), ["pct"])

    def test_tree_endpoint(self):
        import urs.workspace as _wsm
        wd = self._wsdir()
        if not hasattr(self, "rf"):
            self.rf = RequestFactory()
        with mock.patch.object(_wsm, "workspace_dirs", return_value=[wd]):
            req = self.rf.get("/api/ad/tree/", {"ws": wd.name})
            res = _v.api_ad_tree(req)
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        self.assertEqual(len(body["forms"]), 1)
        self.assertEqual(body["ws"], wd.name)
        # وضع الدخول: الأسماء محجوبة والمفاتيح باقية
        self.assertEqual(body["forms"][0]["name_ar"], "")
        self.assertEqual(body["forms"][0]["file"], "t.fmlk")

    def test_tree_names_by_mode(self):
        import urs.workspace as _wsm
        wd = self._wsdir()
        if not hasattr(self, "rf"):
            self.rf = RequestFactory()
        with mock.patch.object(_wsm, "workspace_dirs", return_value=[wd]):
            req = self.rf.get("/api/ad/tree/", {"ws": wd.name})
            req.session = {"ws_mode": "edit", "master_ok": True}
            res = _v.api_ad_tree(req)
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        self.assertEqual(body["forms"][0]["name_ar"], "نموذج ت")
        self.assertEqual(body["reports"][0]["name_ar"], "تقرير ر")
        self.assertEqual(body["rules"][0]["rule_name"], "disc")

    def test_search_endpoint_builtin(self):
        import urs.workspace as _wsm
        wd = self._wsdir()
        if not hasattr(self, "rf"):
            self.rf = RequestFactory()
        with mock.patch.object(_wsm, "workspace_dirs", return_value=[wd]):
            req = self.rf.get("/api/ad/search/", {"ws": wd.name, "kind": "reports", "q": "تقرير"})
            res = _v.api_ad_search(req)
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        self.assertEqual(body["rows"][0]["report_file"], "r.rml")

    def test_search_names_hidden_in_view_mode(self):
        import urs.workspace as _wsm
        wd = self._wsdir()
        if not hasattr(self, "rf"):
            self.rf = RequestFactory()
        with mock.patch.object(_wsm, "workspace_dirs", return_value=[wd]):
            # دخول: الاسم محجوب والمفتاح باقٍ
            req = self.rf.get("/api/ad/search/", {"ws": wd.name, "kind": "forms", "q": "نموذج"})
            res = _v.api_ad_search(req)
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        self.assertEqual(body["rows"][0]["form_file"], "t.fmlk")
        self.assertEqual(body["rows"][0]["form_name"], "")
        self.assertEqual(body["rows"][0]["ref_name"], "")
        with mock.patch.object(_wsm, "workspace_dirs", return_value=[wd]):
            # تعديل: الاسم ظاهر
            req = self.rf.get("/api/ad/search/", {"ws": wd.name, "kind": "forms", "q": "نموذج"})
            req.session = {"ws_mode": "edit", "master_ok": True}
            res = _v.api_ad_search(req)
        body = json.loads(res.content.decode())
        self.assertEqual(body["rows"][0]["ref_name"], "نموذج ت")
        with mock.patch.object(_wsm, "workspace_dirs", return_value=[wd]):
            # المصادر المخصصة: أسماؤها للعمل اليومي — لا تُحجب حتى في الدخول
            req = self.rf.get("/api/ad/search/", {"ws": wd.name, "kind": "sources", "q": ""})
            res = _v.api_ad_search(req)
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        self.assertTrue(all(r["ref_name"] for r in body["rows"]), body)

    def test_search_endpoint_bad_kind(self):
        import urs.workspace as _wsm
        wd = self._wsdir()
        if not hasattr(self, "rf"):
            self.rf = RequestFactory()
        with mock.patch.object(_wsm, "workspace_dirs", return_value=[wd]):
            req = self.rf.get("/api/ad/search/", {"ws": wd.name, "kind": "nope"})
            res = _v.api_ad_search(req)
        self.assertEqual(res.status_code, 400)

    def test_sources_endpoint_crud(self):
        import urs.workspace as _wsm
        wd = self._wsdir()
        if not hasattr(self, "rf"):
            self.rf = RequestFactory()
        with mock.patch.object(_wsm, "workspace_dirs", return_value=[wd]):
            req = self.rf.get("/api/ad/sources/", {"ws": wd.name})
            res = _v.api_ad_sources(req)
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        self.assertTrue(any(s["key"] == "forms" for s in body["sources"]))
        with mock.patch.multiple(_v, _ws_master_ok=mock.Mock(return_value=False)):
            req = self.rf.post("/api/ad/sources/", data=json.dumps({"ws": wd.name, "key": "x"}),
                               content_type="application/json")
            res = _v.api_ad_sources(req)
        self.assertEqual(res.status_code, 403)
        with mock.patch.object(_wsm, "workspace_dirs", return_value=[wd]):
            with mock.patch.multiple(_v, _ws_master_ok=mock.Mock(return_value=True)):
                req = self.rf.post("/api/ad/sources/", data=json.dumps(
                    {"ws": wd.name, "key": "items", "ar": "الأصناف", "table": "ITM",
                     "seq_col": "NO", "name_col": "NAME"}), content_type="application/json")
                res = _v.api_ad_sources(req)
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        self.assertTrue(any(s["key"] == "items" for s in body["sources"]))

    def test_logs_endpoint_requires_master(self):
        if not hasattr(self, "rf"):
            self.rf = RequestFactory()
        with mock.patch.multiple(_v, _ws_master_ok=mock.Mock(return_value=False)):
            req = self.rf.get("/api/ad/logs/")
            res = _v.api_ad_logs(req)
        self.assertEqual(res.status_code, 403)

    def test_sidebar_and_pages(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "templates", "partials", "main_nav.html"), encoding="utf-8") as fh:
            nav = fh.read()
        self.assertIn("سجلات النظام", nav)
        self.assertIn("شجرة دليل النظام", nav)
        self.assertIn("/ad-logs/", nav)
        self.assertIn("/ad-tree/", nav)
        # رابط الشجرة وضع التعديل فقط (دخول edit + ماستر)
        self.assertIn("{% if ws_can_design %}", nav)
        self.assertLess(nav.index("{% if ws_can_design %}"), nav.index("/ad-tree/"))
        for tpl in ("ad_logs.html", "ad_tree.html"):
            with open(os.path.join(base, "templates", tpl), encoding="utf-8") as fh:
                self.assertTrue(len(fh.read()) > 500)

    def test_tree_page_unidbpicker_wiring(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "templates", "ad_tree.html"), encoding="utf-8") as fh:
            html = fh.read()
        for needle in ('{% include "UniDBPicker.js" %}', "adS_conn", "adConnsFetch",
                       "adSrcTblSearch", "adSrcColSearch", "UniDBPicker.pickTable",
                       "UniDBPicker.pickColumn", "adCsrf", "X-CSRFToken", "adReadJson"):
            self.assertIn(needle, html)
        # التضمين داخل <script> (خارجه يظهر الكود كنص في الصفحة)
        self.assertLess(html.index("<script>"), html.index('{% include "UniDBPicker.js" %}'))

    def test_catalog_scrub_view_vs_edit(self):
        fml = [{"file": "a.fmlk",
                "metadata": {"displayName": "نموذج أ", "description": "وصف", "category": "عام"}}]
        rml = [{"file": "r.rml",
                "metadata": {"displayName": "تقرير ر", "description": "وصف", "charts": []}}]
        rules = [{"file": "r.rml", "name": "disc", "display": "الخصم",
                  "description": "وصف", "policy_count": 0}]
        rf = RequestFactory()
        req = rf.get("/x/")
        req.session = {}
        f2, r2, ru2 = _v._scrub_catalog_names(req, fml, rml, rules)
        self.assertEqual(f2[0]["metadata"]["displayName"], "a.fmlk")
        self.assertEqual(f2[0]["metadata"]["description"], "")
        self.assertEqual(r2[0]["metadata"]["displayName"], "r.rml")
        self.assertEqual(ru2[0]["display"], "disc")
        self.assertEqual(ru2[0]["description"], "")
        # الأصل بلا مساس (نسخ)
        self.assertEqual(fml[0]["metadata"]["displayName"], "نموذج أ")
        self.assertEqual(rules[0]["display"], "الخصم")
        req2 = rf.get("/x/")
        req2.session = {"ws_mode": "edit", "master_ok": True}
        f3, r3, ru3 = _v._scrub_catalog_names(req2, fml, rml, rules)
        self.assertEqual(f3[0]["metadata"]["displayName"], "نموذج أ")
        self.assertEqual(ru3[0]["display"], "الخصم")

    def test_app_detail_view_mode_hides_names(self):
        rf = RequestFactory()
        req = rf.get("/app/settings/")
        req.session = {"workspace": "workspace_1"}
        res = _v.app_detail(req, "settings")
        html = res.content.decode()
        self.assertEqual(res.status_code, 200)
        self.assertIn("permissions.fmlk", html)
        self.assertNotIn("صلاحيات المستخدمين", html)

    def test_player_info_and_grid_hooks(self):
        base = os.path.dirname(__file__)
        with open(os.path.join(base, "templates", "forms_player.html"), encoding="utf-8") as fh:
            js = fh.read()
        for needle in ("paintRecordAudit", "api/ad/record-info", "سجل الرقابة",
                       "permMapEditor", "permMapSave", "permMapKind",
                       "ad_perm_forms", "fwDetailACAd", "api/ad/search"):
            self.assertIn(needle, js)


class AdWsResolutionTests(SimpleTestCase):
    """api/ad/* must honor the CURRENT login workspace, not primary.

    Regression: empty ws used to resolve primary_workspace() first, so the
    permissions-grid autocomplete showed another workspace's files."""

    def _mk(self, tag, form_ar):
        import tempfile
        from pathlib import Path as _P
        td = _P(tempfile.mkdtemp(prefix="adws_%s_" % tag))
        self.addCleanup(lambda: __import__("shutil").rmtree(str(td), ignore_errors=True))
        app = td / "apps" / "app1"
        app.mkdir(parents=True)
        (app / ("%s.fmlk" % tag)).write_text(
            '<?xml version="1.0" encoding="utf-8"?><fml>'
            '<fml_metadata name="%s" displayName="%s" table="%s_tbl"/>'
            '<fields><field id="1" name="code" alias="الرمز" inputType="text" '
            'dataType="VARCHAR"/></fields></fml>' % (tag, form_ar, tag), encoding="utf-8")
        return td

    def _req(self, path, params, session_ws="", edit=False):
        from django.test import RequestFactory as _RF
        req = _RF().get(path, params)
        req.session = {"workspace": session_ws} if session_ws else {}
        if edit:
            req.session["ws_mode"] = "edit"
            req.session["master_ok"] = True
        elif not session_ws:
            req.session = {}
        return req

    def test_search_prefers_session_over_primary(self):
        import json as _js
        from unittest import mock as _mock
        import urs.workspace as _wsm
        import urs.views as _v
        wsA, wsB = self._mk("a", "نموذج أ"), self._mk("b", "نموذج ب")
        with _mock.patch.object(_wsm, "workspace_dirs", return_value=[wsA, wsB]), \
             _mock.patch.object(_wsm, "primary_workspace", return_value=wsA):
            req = self._req("/api/ad/search/", {"kind": "forms", "q": ""}, wsB.name, edit=True)
            res = _v.api_ad_search(req)
        body = _js.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        self.assertTrue(body["rows"], body)
        self.assertEqual(body["rows"][0]["form_file"], "b.fmlk")
        self.assertEqual(body["rows"][0]["ref_name"], "نموذج ب")

    def test_tree_prefers_session_over_primary(self):
        import json as _js
        from unittest import mock as _mock
        import urs.workspace as _wsm
        import urs.views as _v
        wsA, wsB = self._mk("a", "نموذج أ"), self._mk("b", "نموذج ب")
        with _mock.patch.object(_wsm, "workspace_dirs", return_value=[wsA, wsB]), \
             _mock.patch.object(_wsm, "primary_workspace", return_value=wsA):
            req = self._req("/api/ad/tree/", {}, wsB.name)
            res = _v.api_ad_tree(req)
        body = _js.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        self.assertEqual(body["ws"], wsB.name)
        self.assertEqual(body["forms"][0]["file"], "b.fmlk")

    def test_explicit_ws_still_wins_over_session(self):
        import json as _js
        from unittest import mock as _mock
        import urs.workspace as _wsm
        import urs.views as _v
        wsA, wsB = self._mk("a", "نموذج أ"), self._mk("b", "نموذج ب")
        with _mock.patch.object(_wsm, "workspace_dirs", return_value=[wsA, wsB]), \
             _mock.patch.object(_wsm, "primary_workspace", return_value=wsA):
            req = self._req("/api/ad/search/", {"ws": wsA.name, "kind": "forms"}, wsB.name)
            res = _v.api_ad_search(req)
        body = _js.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        self.assertEqual(body["rows"][0]["form_file"], "a.fmlk")


class PermissionEngineTests(SimpleTestCase):
    """permission_engine core + enforcement wiring."""

    def _bundle(self):
        from urs import permission_engine as _pe
        return _pe.build_bundle(
            [{"id": 1, "perm_source": "local", "user_name": "u", "is_active": True}],
            {"forms": [{"form_file": "a.fmlk", "form_name": "نموذج أ", "can_view": 1,
                        "can_delete": 0, "fields_visibility": {"price": 0, "qty": 1}}],
             "reports": [{"report_file": "r.rml", "can_view": 0,
                          "columns_visibility": {"c1": 0}}],
             "rules": [{"rule_file": "r.rml", "rule_name": "disc", "can_edit": 0,
                        "variables": {"pct": 0}}]})

    def test_decide(self):
        from urs import permission_engine as _pe
        b = self._bundle()
        self.assertTrue(_pe.check(None, "forms", "a.fmlk", "view"))
        self.assertTrue(_pe.check(b, "forms", "a.fmlk", "view"))
        self.assertFalse(_pe.check(b, "forms", "a.fmlk", "delete"))
        self.assertFalse(_pe.check(b, "forms", "نموذج أ", "delete"))
        self.assertEqual(_pe.hidden_columns(b, "forms", "a.fmlk"), ["price"])
        self.assertFalse(_pe.check(b, "reports", "r.rml", "view"))
        self.assertEqual(_pe.hidden_columns(b, "reports", "r.rml"), ["c1"])
        self.assertEqual(_pe.rule_vars(b, "r.rml", "disc"), {"pct": False})
        self.assertFalse(_pe.check(b, "rules", ("r.rml", "disc"), "edit"))

    def test_external_visibility_only(self):
        from urs import permission_engine as _pe
        b = _pe.build_bundle(
            [{"id": 2, "perm_source": "خارجي", "user_name": "u", "is_active": True}],
            {"forms": [{"form_file": "a.fmlk", "can_delete": 0,
                        "fields_visibility": {"price": 0}}]})
        self.assertTrue(_pe.is_external_only(b))
        self.assertTrue(_pe.check(b, "forms", "a.fmlk", "delete"))
        self.assertEqual(_pe.hidden_columns(b, "forms", "a.fmlk"), ["price"])

    def test_inactive_deny_and_fail_open(self):
        from urs import permission_engine as _pe
        b = _pe.build_bundle(
            [{"id": 3, "perm_source": "local", "user_name": "u", "is_active": 0}], {"forms": []})
        self.assertFalse(_pe.check(b, "forms", "a.fmlk", "view"))
        self.assertIsNone(_pe.build_bundle([], {}))

    def test_permissions_form_shape(self):
        from fmlk_engine.compiler import FMLKFormCompiler
        import pathlib as _pl
        base = _pl.Path(__file__).resolve().parent.parent
        for ws in ("workspace_1", "workspace_2"):
            comp = FMLKFormCompiler(path=base / ws / "apps" / "settings" / "permissions.fmlk")
            fields = {f.name: f for f in comp.fields()}
            self.assertIn("perm_source", fields)
            self.assertIn("user_name", fields)
            self.assertEqual(fields["user_name"].input_type, "datamodal")
            self.assertTrue(str(fields["user_name"].config.get("src_expr") or "").strip())
            dets = {d.table: d for d in comp.details()}
            for t in ("ad_perm_forms", "ad_perm_reports", "ad_perm_rules", "ad_perm_objects"):
                self.assertIn(t, dets)
            fcols = {c["name"]: c for c in dets["ad_perm_forms"].columns}
            self.assertTrue(fcols["form_name"].get("src_expr"))
            for a in ("can_view", "can_delete"):
                self.assertIn(a, fcols)
            rcols = {c["name"]: c for c in dets["ad_perm_reports"].columns}
            self.assertIn("columns_visibility", rcols)
            rucols = {c["name"]: c for c in dets["ad_perm_rules"].columns}
            self.assertIn("variables", rucols)
            if ws == "workspace_2":
                # فرع الحسابات الديناميكي من مصدر الشجرة src:accounts
                self.assertIn("ad_perm_accounts", dets)
                acols = {c["name"]: c for c in dets["ad_perm_accounts"].columns}
                self.assertEqual(acols["source"].get("default"), "src:accounts")
                self.assertEqual(acols["ref_key"].get("src_expr"), "'src:accounts'")
                for a in ("ref_name", "can_view", "can_delete", "extra"):
                    self.assertIn(a, acols)

    def test_form_perm_endpoint(self):
        if not hasattr(self, "rf"):
            self.rf = RequestFactory()
        with mock.patch.multiple(
                _v,
                _perm_ws_for_file=mock.Mock(return_value="workspace_1"),
                _perm_bundle_for=mock.Mock(return_value=("u", None))):
            req = self.rf.get("/api/ad/form-perm/", {"fml": "a.fmlk", "app": "x"})
            res = _v.api_ad_form_perm(req)
        body = json.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        self.assertEqual(body["hidden"], [])

    def test_gate_helpers(self):
        from urs import permission_engine as _pe
        b = self._bundle()
        from types import SimpleNamespace as _NS
        # can_view=0 in fixture → gate blocks with 403
        eng = _NS(columns=[_NS(alias="Price", name="price"), _NS(alias="Qty", name="qty")])
        err = _v._perm_gate_report(eng, "r.rml", b)
        self.assertIsNotNone(err)
        self.assertEqual(err.status_code, 403)
        b2 = _pe.build_bundle(
            [{"id": 9, "perm_source": "local", "user_name": "u", "is_active": True}],
            {"reports": [{"report_file": "r.rml", "can_view": 1,
                           "columns_visibility": {"price": 0}}]})
        err = _v._perm_gate_report(eng, "r.rml", b2)
        self.assertIsNone(err)
        self.assertEqual([c.name for c in eng.columns], ["qty"])
        feng = _NS(fields=[_NS(name="price"), _NS(name="id"), _NS(name="qty")])
        err = _v._perm_gate_form_engine(feng, "a.fmlk", b)
        self.assertIsNone(err)
        self.assertEqual(sorted([f.name for f in feng.fields]), ["id", "qty"])

class AdLogStorageTests(TestCase):
    """Activity-log table: ensure + write + query + record info (isolated test DB)."""

    def test_log_roundtrip(self):
        from ad_system import engine as _ad
        self.assertTrue(_ad.ensure_log_table())
        self.assertTrue(_ad.log(ws_id="w1", username="u1", action="insert", entity="t.fmlk",
                                entity_id="7", ref_table="t_tbl", device="1.2.3.4 | pc",
                                details="x"))
        self.assertTrue(_ad.log(ws_id="w1", username="u1", action="update", entity="t.fmlk",
                                entity_id="7", ref_table="t_tbl", device="1.2.3.4 | pc"))
        q = _ad.query_log(ws_id="w1", username="u1")
        self.assertEqual(q["total"], 2)
        self.assertEqual(q["rows"][0]["action"], "update")
        q2 = _ad.query_log(action="insert")
        self.assertEqual(q2["total"], 1)
        info = _ad.record_info(ref_table="t_tbl", entity_id="7", ws_id="w1")
        self.assertEqual(info["count"], 2)
        self.assertEqual(info["created"]["username"], "u1")
        self.assertEqual(info["modified"]["action"], "update")
        self.assertEqual(_ad.record_info(ref_table="t_tbl", entity_id="999")["count"], 0)

class ExtRealCaseTests(SimpleTestCase):
    """Oracle folds unquoted idents to UPPER: quoted lowercase must resolve
    to the true case (probe variants), else ORA-00904 — the empty-users bug."""

    class _Cur:
        def __init__(self, desc=None, rows=None, fail=None):
            self._desc = desc or []
            self._rows = rows or []
            self._fail = fail

        @property
        def description(self):
            return [(c,) for c in self._desc]

        def fetchall(self):
            if self._fail:
                raise Exception(self._fail)
            return list(self._rows)

        def fetchone(self):
            if self._fail:
                raise Exception(self._fail)
            return self._rows[0] if self._rows else None

        def close(self):
            pass

    class _OraShim:
        """Quoted lowercase table/column on Oracle → ORA error (like the DB)."""

        def __init__(self, cols, rows):
            self.cols = cols
            self.rows = rows
            self.seen = []
            self.conn = None

        def _exec(self, sql, params):
            import re as _re
            self.seen.append(sql)
            for m in _re.finditer(r'"([^"]+)"', sql):
                nm = m.group(1)
                if nm in ("T", "v", "label"):
                    continue
                if nm != nm.upper() and nm.upper() in [c.upper() for c in self.cols + ["USER_R", "IAS20264"]]:
                    raise Exception('ORA-00904: "%s": invalid identifier' % nm)
            if "WHERE 1=0" in sql:
                return ExtRealCaseTests._Cur(desc=self.cols)
            if "COUNT(*)" in sql:
                return ExtRealCaseTests._Cur(rows=[(len(self.rows),)])
            rows = self.rows
            mi = _re.search(r"\bIN\s*\(([^)]*)\)", sql)
            if mi:
                want = set(_re.findall(r"'((?:[^']|'')*)'", mi.group(1)))
                want = {w.replace("''", "'") for w in want}
                rows = [r for r in rows if str(r[0]) in want]
            return ExtRealCaseTests._Cur(desc=self.cols, rows=rows)

    def _shim(self):
        return self._OraShim(["U_ID", "U_A_NAME"],
                             [("10", "Admin"), ("20", "User")])

    def test_options_resolves_true_case(self):
        shim = self._shim()
        with mock.patch.multiple(
                _v, _fmlk_external_exec=mock.Mock(return_value=(shim, "oracle"))):
            out = _v._fmlk_ext_options("c", "IAS20264", "user_r", "u_id", "u_a_name", "", 10)
        self.assertEqual(len(out["options"]), 2)
        self.assertEqual(out["options"][0], {"value": "10", "label": "Admin"})
        main = [s for s in shim.seen if "WHERE 1=0" not in s][0]
        self.assertIn('"T"."U_ID"', main)
        self.assertNotIn('"T"."u_id"', main)

    def test_options_upper_table_retry(self):
        shim = self._shim()
        with mock.patch.multiple(
                _v, _fmlk_external_exec=mock.Mock(return_value=(shim, "oracle"))):
            out = _v._fmlk_ext_options("c", "ias20264", "user_r", "u_id", "", "", 10)
        self.assertEqual(len(out["options"]), 2)

    def test_grid_source_resolves_true_case(self):
        shim = self._shim()
        with mock.patch.multiple(
                _v, _fmlk_external_exec=mock.Mock(return_value=(shim, "oracle"))):
            res = _v._fmlk_grid_source_external("c", "IAS20264", "user_r", "u_id", "u_a_name", "1", 8)
        import json as _js
        body = _js.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        self.assertEqual(body["sub_key"], "U_ID")
        self.assertEqual(body["key"], "U_ID")
        self.assertTrue(any(r.get("U_ID") == "10" for r in body["rows"]))

    def test_ext_map_resolves_true_case(self):
        shim = self._shim()
        with mock.patch.multiple(
                _v, _fmlk_external_exec=mock.Mock(return_value=(shim, "oracle"))):
            out = _v._fmlk_ext_map("c", "IAS20264", "user_r", "u_id", "u_a_name", ["10"])
        self.assertEqual(out, {"10": "Admin"})

    def test_detail_search_external_still_works(self):
        shim = self._shim()
        det = {"sub_key": "id", "columns": []}
        with mock.patch.multiple(
                _v, _fmlk_external_exec=mock.Mock(return_value=(shim, "oracle"))):
            res = _v._fmlk_detail_search_external(det, "c", "IAS20264", "user_r", "1", "", 1, 8)
        import json as _js
        body = _js.loads(res.content.decode())
        self.assertEqual(res.status_code, 200, body)
        self.assertTrue(len(body["rows"]) > 0)

    def test_import_wizard_syntax(self):
        import shutil
        import subprocess
        import tempfile
        base = os.path.dirname(__file__)
        src = os.path.join(base, "..", "rml_python", "shared", "rml_wizard_script.html")
        if not shutil.which("node"):
            self.skipTest("node missing")
        with open(src, encoding="utf-8") as f:
            code = f.read()
        with tempfile.NamedTemporaryFile("w", suffix=".js", delete=False, encoding="utf-8") as t:
            t.write(code)
            tmp = t.name
        try:
            r = subprocess.run(["node", "--check", tmp], capture_output=True, text=True)
        finally:
            try:
                os.unlink(tmp)
            except OSError:
                pass
        self.assertEqual(r.returncode, 0, r.stderr[-2000:] if r.stderr else "")